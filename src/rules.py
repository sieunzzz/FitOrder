"""
FitOrder 규칙 엔진 — 마스터 조회 / 금액 계산 / 검증

사용:
    from rules import Master, calc_erp, validate
    M = Master("../data/master/fitorder_master.xlsx")
    row = M.find_item("200", "원코드", "L자", "DU")
    issues = validate(order_dict, M)
"""
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import openpyxl

# ─────────────────────────────────────────────
# 거래처 (1차 구현 7개사)
# ─────────────────────────────────────────────
CLIENT_INFO = {
    # 상호: (관리코드, 단가컬럼, 기본배송, 내부표시, 입력형태)
    "DI":   ("대구",       "출고L가", "DI/SP",   "(N)", "excel"),
    "휴안": ("광주",       "출고I가", "택배",     "(K)", "excel"),
    "M":    ("대구",       "출고I가", None,      "(K)", "image"),
    "DU":   ("대구광역시", "출고I가", None,      "(K)", "image"),
    "RT":   ("경기도",     "출고I가", "판별필요", "(K)", "image"),
    "JO":   ("대구광역시", "출고I가", None,      "(K)", "image"),
    "JL":   ("대구",       "출고I가", None,      "(K)", "image"),
}

# 세로 규격에 " 표기를 쓰지 않는 거래처
NO_DITTO = {"휴안"}

# 창이 1개여도 설치장소를 기재사항에 합치지 않는 거래처
# (JO 는 기재사항 열에 주문번호만, 설치장소는 다음 열에 적는다)
NO_MERGE_PLACE = {"JO"}

EXCEPTION_WORDS = ["수리", "롤스크린", "홀딩도어", "BMIX", "B Mix",
                   "부속", "레일", "커버", "점보"]

# 포장비용 (휴안) — 경영박사 전표 맨 아래에 창당 1줄
PACKING_ITEM = {"품명": "포장비용(25mm)", "규격": "창당", "단가": 500}
PACKING_CLIENTS = {"휴안"}

MIN_WIDTH = {"원코드": 36, "투코드": 10, "셔터": 10}
CONFIDENCE_THRESHOLD = 0.8

# DI 는 색상명으로만 발주 -> 슬랫 번호 매핑
DI_COLOR = {
    "화이트": "102", "백아이보리": "105", "아이보리": "200", "밀크로즈": "220",
    "크림베이지": "230", "바닐라": "260", "더블믹스달콤코코아": "260",
    "밀크티": "270", "머쉬룸": "280", "베이지": "290", "연핑크": "330",
    "옐로우": "500", "머스타드": "590", "더블믹스스카이블루": "620",
    "딥블루": "690", "멜론": "740", "그린": "760", "딥그린": "790",
    "연그레이": "820", "그레이": "860", "스톤그레이": "870",
    "모던그레이": "880", "다크그레이": "890", "차콜그레이": "920",
    "블랙": "990", "민트": "023", "아쿠아": "026", "오렌지": "025",
    "글로시실버": "030", "카멜": "140",
    # 믹스 계열은 베이스 코드를 쓰되 확인 대상
    "포인트믹스밀크코코아": "200", "포인트믹스블랙야크": "102",
    "포인트믹스애쉬그레이": "102",
}
DI_COLOR_UNSURE = {"포인트믹스밀크코코아", "포인트믹스블랙야크",
                   "포인트믹스애쉬그레이"}


# ─────────────────────────────────────────────
# 마스터
# ─────────────────────────────────────────────
class Master:
    def __init__(self, path):
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["품목"]
        hdr = [c.value for c in ws[1]]
        self.tier_col = {h: i for i, h in enumerate(hdr) if str(h).startswith("출고")}
        self.items = {}          # 품명 -> dict
        self.by_code = {}        # 코드 -> 기본 품명 (접미사 없는 것)
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = row[0]
            if not name:
                continue
            self.items[name] = {
                "품명": name, "대분류": row[1], "규격": row[2] or "",
                "비고2": row[7] or "",
                **{h: (row[i] or 0) for h, i in self.tier_col.items()},
            }
            m = re.match(r"^B([0-9]{3}[A-Za-z]*)\(([^)]*)\)$", str(name))
            if m:
                code, inner = m.group(1), m.group(2)
                # '펄無' 같은 변형은 기본 품목으로 쓰지 않는다
                if code in self.by_code and ("/" in inner or "無" in inner):
                    continue
                self.by_code[code] = name

    # 손잡이실 색상 (DI 정렬용)
    def thread_color(self, code):
        base = self.by_code.get(code)
        if not base:
            return None
        m = re.search(r"실\s*:?\s*([가-힣A-Za-z]+)", self.items[base]["비고2"])
        return m.group(1) if m else None

    def build_name(self, code, kind, type_, client):
        """품명 조합: B200(IV)-L18/원코드 형태"""
        base = self.by_code.get(code)
        if not base:
            return None
        if type_ == "L자":                       # L자는 DI도 동일 단가
            return f"{base}-L18/{kind}"
        if client == "DI":
            return base + ("-DI" if kind == "투코드" else f"-{kind}/DI")
        return base if kind == "투코드" else f"{base}-{kind}"

    def find_item(self, code, kind, type_, client):
        name = self.build_name(code, kind, type_, client)
        if not name or name not in self.items:
            return None
        it = dict(self.items[name])
        tier = CLIENT_INFO.get(client, (None, "출고I가"))[1]
        it["단가"] = it.get(tier, 0)
        it["단가등급"] = tier
        return it


# ─────────────────────────────────────────────
# 계산
# ─────────────────────────────────────────────
def calc_erp(width_cm, height_cm, unit_price):
    """헤베(둘째자리 반올림) / 최소 1.5 / 금액·부가세 원단위 반올림"""
    raw = (Decimal(str(width_cm)) / 100) * (Decimal(str(height_cm)) / 100)
    hebe = raw.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    qty = max(hebe, Decimal("1.5"))
    amount = (Decimal(str(unit_price)) * qty).quantize(
        Decimal("1"), rounding=ROUND_HALF_UP)
    vat = (amount / 10).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    return {"헤베": hebe, "수량": qty, "단가": unit_price,
            "금액": amount, "부가세": vat}


def default_handle_length(kind, height):
    if height is None:
        return None
    if kind == "원코드":
        return 150 if height >= 210 else 130
    if kind == "투코드":
        return 150 if height >= 210 else 100
    if kind == "셔터":
        if height >= 210:
            return 150
        if height >= 100:
            return 100
        return int(height // 10) * 10 - 10
    return None


def normalize_handle(length):
    """10단위 내림. 보정 여부 반환"""
    if not length:
        return None, False
    r = (int(length) // 10) * 10
    return r, r != int(length)


def ledger_color(code, kind, type_):
    """장부 색상 열 문자열"""
    if not code:
        return None
    if type_ == "L자":
        return f"B L18-{kind} {code}"
    return f"B {code}" if kind == "투코드" else f"B {kind} {code}"


# 장부 색상 열 부분 서식 (D4 에서 CellRichText 로 적용)
COLOR_FONT = {"원코드": "008000", "셔터": "FF0000"}     # 초록 / 빨강
NOTE_FONT = {"틀안": "0000FF", "선불": "FF0000"}        # 파랑 / 빨강


# ─────────────────────────────────────────────
# 검증
# ─────────────────────────────────────────────
def validate(order, M):
    """order = extract_order() 또는 파서 결과. [(등급, 코드, 메시지, 행번호)]"""
    out = []
    client = order.get("거래처")
    if client not in CLIENT_INFO:
        out.append(("red", "거래처불명", "거래처를 판별하지 못했습니다", None))

    items = order.get("items") or []
    if not items:
        out.append(("red", "항목없음", "추출된 주문 항목이 없습니다", None))

    # 발주서 메모의 총 창 수와 대조
    memo = str(order.get("전체원문") or "") + str(order.get("전체기재사항") or "")
    m = re.search(r"총\s*(\d+)\s*창", memo)
    if m and len(items) != int(m.group(1)):
        out.append(("red", "창개수불일치",
                    f"발주서 메모는 {m.group(1)}창인데 {len(items)}건이 추출되었습니다", None))

    for i, it in enumerate(items, 1):
        w, h = it.get("가로"), it.get("세로")
        kind, type_ = it.get("종류"), it.get("타입")
        code = it.get("품목코드")

        if it.get("예외품목"):
            out.append(("info", "예외품목",
                        f"{it['예외품목']} — 장부 제외, 담당자 확인 필요", i))
            continue

        if w in (None, ""):
            out.append(("red", "가로누락", "가로 치수가 없습니다", i))
        if h in (None, ""):
            out.append(("red", "세로누락", "세로 치수가 없습니다", i))
        if w and h:
            if w > h and h < 100:
                out.append(("yellow", "치수순서의심",
                            f"가로({w}) > 세로({h}) — 순서 확인", i))
            lo = MIN_WIDTH.get(kind)
            if lo and w < lo:
                out.append(("red", "최소사이즈미달",
                            f"{kind} 최소 가로 {lo}cm 미만 ({w}cm)", i))
            if not (10 <= w <= 400) or not (10 <= h <= 400):
                out.append(("yellow", "치수범위이상",
                            f"비정상적인 치수 {w}x{h}", i))

        if not code:
            out.append(("red", "품목미등록", "색상 코드를 읽지 못했습니다", i))
        else:
            hit = M.find_item(code, kind, type_, client)
            if not hit:
                out.append(("red", "품목미등록",
                            f"마스터에 없는 조합: {code}/{kind}/{type_}", i))
            elif not hit["단가"]:
                out.append(("red", "단가없음",
                            f"{hit['품명']} — {hit['단가등급']} 단가가 0입니다", i))

        if it.get("손잡이방향") is None:
            out.append(("yellow", "손잡이미기재",
                        "방향 미기재 — 우측 기본 적용", i))
        _, adj = normalize_handle(it.get("손잡이길이"))
        if adj:
            out.append(("yellow", "손잡이길이보정",
                        f"{it['손잡이길이']} → 10단위 내림", i))

        # 투코드 + '봉' 표기는 M 이외 거래처에서 부품일 수 있음
        if kind == "투코드" and client != "M" and "봉" in str(it.get("원문") or ""):
            out.append(("yellow", "봉표기모호",
                        "'봉'이 손잡이인지 별도 부품인지 확인", i))

        conf = it.get("확신도") or {}
        low = [k for k, v in conf.items() if isinstance(v, (int, float))
               and v < CONFIDENCE_THRESHOLD]
        if low:
            out.append(("yellow", "확신도낮음",
                        "판독 확신도 낮음: " + ", ".join(low), i))

    if order.get("변경요청"):
        txt = (order.get("변경문구") or "").strip()
        out.append(("review", "변경요청",
                    "변경 요청으로 보입니다"
                    + (f" — \"{txt[:60]}\"" if txt else "")
                    + " · 대상 주문을 확인하세요", None))

    if client == "RT" and not (order.get("배송") or {}).get("방식"):
        out.append(("red", "배송판별불가", "RT — 화물/택배 구분 불가", None))

    return out


def worst(issues):
    for lv in ("red", "yellow", "info"):
        if any(i[0] == lv for i in issues):
            return lv
    return "ok"


# ─────────────────────────────────────────────
# 중복 감지 / 주문 변경 탐지
# ─────────────────────────────────────────────
DUP_LOOKBACK_DAYS = 7
CHANGE_LOOKBACK_DAYS = 2
SIZE_TOL = 1.0          # 치수 ±1cm 이내면 같은 창으로 의심


def item_key(client, it):
    """중복 판정 키"""
    return (client, it.get("품목코드"), it.get("종류"), it.get("타입"),
            it.get("가로"), it.get("세로"), it.get("수량"),
            it.get("손잡이방향"), it.get("손잡이길이"))


def _near(a, b):
    try:
        return abs(float(a) - float(b)) <= SIZE_TOL
    except (TypeError, ValueError):
        return a == b


def similar(client_a, a, client_b, b):
    """같은 창으로 의심되는지 (치수 ±1cm 허용)"""
    if client_a != client_b:
        return False
    if a.get("품목코드") != b.get("품목코드"):
        return False
    return _near(a.get("가로"), b.get("가로")) and _near(a.get("세로"), b.get("세로"))


def find_duplicates(orders, past=None):
    """배치 안 중복 + 과거(미출력) 주문과의 중복.
       past = [(라벨, 거래처, item), ...]
       반환: [(등급, 코드, 메시지, 전역행번호)]"""
    out, seen, rows = [], {}, []
    for oi, o in enumerate(orders):
        c = o.get("거래처")
        for it in o.get("items", []):
            if it.get("예외품목"):
                continue
            rows.append((oi, c, it))

    for n, (oi, c, it) in enumerate(rows, 1):
        k = item_key(c, it)
        if k in seen:
            out.append(("review", "중복", f"{seen[k]}행과 완전히 동일합니다", n))
        else:
            seen[k] = n
            for m, (_, c2, it2) in enumerate(rows[:n - 1], 1):
                if similar(c, it, c2, it2):
                    out.append(("review", "중복의심",
                                f"{m}행과 치수가 거의 같습니다 "
                                f"({it2.get('가로')}x{it2.get('세로')})", n))
                    break
        for label, pc, pit in (past or []):
            if item_key(pc, pit) == k:
                out.append(("review", "과거중복",
                            f"{label} 주문과 동일합니다", n))
                break
    return out


def change_candidate(client_a, a, client_b, b):
    """변경 대상 후보 판정 — 한쪽 치수만 바뀌는 경우가 많으므로
       거래처·품목이 같고 가로 또는 세로 하나가 비슷하면 후보로 본다."""
    if client_a != client_b:
        return False
    if a.get("품목코드") != b.get("품목코드"):
        return False
    return _near(a.get("가로"), b.get("가로")) or _near(a.get("세로"), b.get("세로"))


def find_changes(orders, past):
    """최근 주문 중 변경 대상 후보를 찾는다.
       past = [(라벨, 거래처, item, 출력여부), ...]"""
    out, n = [], 0
    for o in orders:
        c = o.get("거래처")
        for it in o.get("items", []):
            if it.get("예외품목"):
                continue
            n += 1
            for label, pc, pit, printed in past:
                if not change_candidate(c, it, pc, pit):
                    continue
                diff = []
                for f, nm in (("가로", "가로"), ("세로", "세로"),
                              ("손잡이길이", "손잡이"), ("품목코드", "색상")):
                    a, b = pit.get(f), it.get(f)
                    if a != b:
                        diff.append(f"{nm} {a} → {b}")
                if not diff:
                    continue
                msg = f"{label} 주문의 변경일 수 있습니다 · " + ", ".join(diff)
                if printed:
                    out.append(("red", "변경-출력됨",
                                msg + " · 장부를 다시 출력해야 합니다", n))
                else:
                    out.append(("review", "변경의심", msg, n))
                break
    return out
