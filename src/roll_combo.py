"""FitOrder 롤/콤비 전용 처리 엔진.

- 안산)보노: 하비창 Excel 주문서 전용 파서
- 그 외 업체: 카카오톡 이미지/PDF 비전 추출
- 롤/콤비/쉐이드 품목장 3개를 하나의 마스터로 매칭
- C는 장부 접두사 생략, R/S는 표시
- 연창은 자동 판정하지 않음
- 피카소는 가로/세로 반전
"""
from __future__ import annotations

import json
import re
import tempfile
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

import openpyxl

from parsers import _bono_freight_info, _direction, _sender_text
from master_loader import read_directory, resolve_master_dir

ROLL_CLIENT_INFO = {
    "안산)보노": {"erp": "안산)주식회사 보노", "price": "출고C가", "delivery": "택배", "input": "excel"},
    "미래가공": {"erp": "주식회사 미래가공", "price": "출고C가", "delivery": "내사", "input": "image"},
    "천안)채원": {"erp": "천안)채원홈데코", "price": "출고C가", "delivery": "택배", "input": "image"},
    "세일인테리어": {"erp": "세일인테리어(평리)", "price": "출고D가", "delivery": "택배", "input": "image"},
    "진천)충북혁신": {"erp": "진천)충북혁신커텐", "price": "출고D가", "delivery": "택배", "input": "image"},
    "음성)충북혁신": {"erp": "음성)충북혁신 홈데코", "price": "출고D가", "delivery": "택배", "input": "image"},
    "논공)현대종합장식": {"erp": "논공)현대종합장식", "price": "출고D가", "delivery": "택배", "input": "image"},
    "원주)창과방": {"erp": "원주)원주창과방커튼", "price": "출고C가", "delivery": "택배", "input": "image"},
    "피카소": {"erp": "피카소인테리어(구암)", "price": "출고D가", "delivery": "택배", "input": "image", "swap_wh": True},
    "제주)우림": {"erp": "제주)우림종합장식", "price": "출고C가", "delivery": "택배", "input": "image"},
    "제주)중문": {"erp": "제주)중문장식", "price": "출고C가", "delivery": "택배", "input": "image"},
}
ROLL_CLIENTS = list(ROLL_CLIENT_INFO)

ROLL_ALIASES = {
    "보노": "안산)보노", "이끌림": "안산)보노", "안산보노": "안산)보노",
    "미래": "미래가공", "미래가공": "미래가공",
    "채원": "천안)채원", "채원홈데코": "천안)채원",
    "세일": "세일인테리어", "세일인테리어": "세일인테리어",
    "진천점": "진천)충북혁신", "진천충북": "진천)충북혁신", "진천충북혁신": "진천)충북혁신",
    "음성점": "음성)충북혁신", "음성충북": "음성)충북혁신", "음성충북혁신": "음성)충북혁신",
    "현대종합장식": "논공)현대종합장식", "논공현대": "논공)현대종합장식",
    "창과방": "원주)창과방", "원주창과방": "원주)창과방",
    "피카소": "피카소",
    "우림": "제주)우림", "우림종합장식": "제주)우림",
    "중문": "제주)중문", "중문장식": "제주)중문",
}

COLOR_ABBR = {"화이트": "W", "흰색": "W", "백색": "W", "아이보리": "IV"}


ROLL_IGNORED_NOTE_KEYS = {
    "닫힌창", "닫힌사이즈", "닫힌창사이즈", "닫힌창크기",
}


def _roll_note_ignored(value: Any) -> bool:
    key = re.sub(r"[^0-9A-Za-z가-힣]", "", str(value or "")).casefold()
    return key in {x.casefold() for x in ROLL_IGNORED_NOTE_KEYS}


def _clean_roll_note_text(value: Any) -> str | None:
    parts = [x.strip() for x in re.split(r"[/\n]+", str(value or "")) if x.strip()]
    kept = [x for x in parts if not _roll_note_ignored(x)]
    return "/".join(kept) or None


def _norm(value: Any) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣]", "", str(value or "")).casefold()


def _num(value):
    if value in (None, ""):
        return None
    try:
        f = float(value)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        m = re.search(r"\d+(?:\.\d+)?", str(value))
        if not m:
            return None
        f = float(m.group(0))
        return int(f) if f.is_integer() else f


def canonical_client(value: Any, raw_text: str = "") -> str | None:
    text = re.sub(r"\s+", "", str(value or ""))
    whole = f"{text} {raw_text}"
    # 충북혁신은 지점명이 업체명을 결정한다.
    if "진천점" in whole:
        return "진천)충북혁신"
    if "음성점" in whole:
        return "음성)충북혁신"
    for client in ROLL_CLIENTS:
        if re.sub(r"\s+", "", client) in text:
            return client
    key = _norm(text)
    for alias, client in sorted(ROLL_ALIASES.items(), key=lambda x: len(x[0]), reverse=True):
        if _norm(alias) and _norm(alias) in key:
            return client
    # 전체 원문에서도 강한 별칭만 확인한다.
    whole_key = _norm(raw_text)
    for alias, client in sorted(ROLL_ALIASES.items(), key=lambda x: len(x[0]), reverse=True):
        if len(_norm(alias)) >= 3 and _norm(alias) in whole_key:
            return client
    return None


def simplify_place(value: Any) -> str | None:
    """시공위치를 원문 그대로 보존한다.

    보노/미래가공의 1-1, 1-2 같은 번호는 각 창을 구분하는 핵심 기재사항이므로
    임의로 줄이거나 삭제하지 않는다. 앞뒤 공백만 정리한다.
    """
    text = str(value or "").strip()
    return text or None


def color_display(value: Any) -> str:
    text = str(value or "").strip()
    return COLOR_ABBR.get(text, text)


def _strip_fire(text: str) -> str:
    text = re.sub(r"[（(]?\s*방염\s*[)）]?", "", str(text or ""), flags=re.I)
    return re.sub(r"\s+", " ", text).strip(" -_/()")


@dataclass
class RollMatch:
    data: dict
    score: float


def resolve_roll_master_dir() -> Path:
    """공통 품목장 로더를 사용해 롤/콤비 품목장 폴더를 찾는다."""
    return resolve_master_dir("roll_combo", extensions=(".xls", ".xlsx"))


class RollComboMaster:
    """롤/콤비/쉐이드 xls 3개를 하나의 품목장처럼 읽는다."""
    is_roll_combo_master = True

    def __init__(self, master_dir: str | Path):
        self.master_dir = Path(master_dir)
        self.rows: list[dict] = []
        self._by_code: dict[str, dict] = {}
        raw_rows = read_directory(
            self.master_dir, required_headers=("관리코드", "품명"),
            extensions=(".xls", ".xlsx")
        )
        for row in raw_rows:
            self._add_row(row)
        if not self.rows:
            raise FileNotFoundError(f"롤/콤비 품목장이 없습니다: {self.master_dir}")

    def _add_row(self, row: dict):
        code = str(row.get("관리코드") or "").strip()
        if not code:
            return
        category = code[0].upper() if code[:1].upper() in {"C", "R", "S"} else ""
        body = code[1:] if category else code
        if "-" in body:
            product_from_code, color_from_code = body.rsplit("-", 1)
        else:
            product_from_code, color_from_code = body, ""
        product = str(row.get("품명") or product_from_code).strip()
        product = re.sub(r"^[CRS]\s*", "", product, flags=re.I)
        if product_from_code and _norm(product_from_code) in _norm(product):
            product = product_from_code
        color = color_from_code
        fire = ("방염" in code or "방염" in str(row.get("품명") or "")
                or "필증" in str(row.get("규격") or ""))
        item = dict(row)
        item.update({
            "_source": row.get("_master_file") or "",
            "_category": category,
            "_product": _strip_fire(product),
            "_color": _strip_fire(color),
            "_fireproof": bool(fire),
        })
        self.rows.append(item)
        self._by_code[_norm(code)] = item

    def _unit_price(self, row: dict, client: str | None) -> int:
        tier = ROLL_CLIENT_INFO.get(client or "", {}).get("price", "출고C가")
        value = row.get(tier)
        try:
            return int(round(float(value)))
        except (TypeError, ValueError):
            return 0

    def match(self, product: Any, color: Any = None, fireproof: bool | None = None,
              client: str | None = None) -> RollMatch | None:
        p = str(product or "").strip()
        c = str(color or "").strip()
        raw = f"{p}-{c}".strip("-")
        direct = self._by_code.get(_norm(raw))
        if direct:
            out = dict(direct)
            out["단가"] = self._unit_price(out, client)
            return RollMatch(out, 1.0)

        pn, cn = _norm(_strip_fire(p)), _norm(_strip_fire(c))
        best = None
        best_score = 0.0
        for row in self.rows:
            rp, rc = _norm(row.get("_product")), _norm(row.get("_color"))
            # 제품명이 완전히 다르면 색상만으로 매칭하지 않는다.
            pscore = SequenceMatcher(None, pn, rp).ratio() if pn and rp else 0.0
            cscore = SequenceMatcher(None, cn, rc).ratio() if cn and rc else (1.0 if not cn else 0.0)
            score = 0.78 * pscore + 0.22 * cscore
            if pn == rp:
                score += 0.12
            if cn and cn == rc:
                score += 0.08
            if fireproof is True and row.get("_fireproof"):
                score += 0.08
            elif fireproof is True and not row.get("_fireproof"):
                score -= 0.15
            elif fireproof is False and row.get("_fireproof") and "방염" not in p:
                score -= 0.04
            if score > best_score:
                best, best_score = row, score
        if best is None or best_score < 0.58:
            return None
        out = dict(best)
        out["단가"] = self._unit_price(out, client)
        return RollMatch(out, min(1.0, best_score))

    def find_order_item(self, item: dict, client: str | None = None):
        cached = item.get("_roll_match")
        if isinstance(cached, dict) and cached.get("관리코드"):
            row = dict(cached)
            row["단가"] = self._unit_price(row, client)
            return row
        match = self.match(item.get("롤품명") or item.get("품명") or item.get("색상원문"),
                           item.get("롤색상") or item.get("원단명"),
                           item.get("_roll_fireproof"), client)
        return match.data if match else None

    # 기존 출력 코드의 조건부 호출과 호환용
    def thread_color(self, code):
        return None


def ledger_text_from_match(row: dict) -> str:
    prefix = str(row.get("_category") or "").upper()
    product = _strip_fire(row.get("_product") or row.get("품명") or "")
    color = color_display(row.get("_color"))
    lead = "" if prefix == "C" else (prefix + " " if prefix in {"R", "S"} else "")
    # 장부 정석: 방염은 제품명 바로 뒤에 붙인다.
    # 예: 리마암막방염 (그레이). 방염 제품에는 <필증>도 같은 색상 셀에 표시한다.
    if row.get("_fireproof"):
        product = f"{product}방염"
    text = f"{lead}{product}".strip()
    if color:
        text += f" ({color})"
    # 장부 품명은 접두사-제품명-색상 사이 공백을 정확히 한 칸으로 맞춘다.
    return re.sub(r"[ \t]+", " ", text).strip()


def attach_match(item: dict, master: RollComboMaster, client: str | None) -> dict:
    product = item.get("롤품명") or item.get("품명") or item.get("제품명")
    color = item.get("롤색상") or item.get("원단명") or item.get("색상")
    fire = bool(item.get("_roll_fireproof") or "방염" in " ".join(
        str(item.get(k) or "") for k in ("작동방식필증", "기재사항", "원문", "품명")))
    match = master.match(product, color, fire, client)
    item["_product_group"] = "roll_combo"
    item["제품군"] = "롤/콤비"
    item["연창"] = False
    item["좌개수"] = 1 if item.get("손잡이방향") == "좌" else 0
    item["우개수"] = 1 if item.get("손잡이방향") == "우" else 0
    item["창개수"] = int(item.get("창개수") or 1)
    if match:
        item["_roll_match"] = match.data
        item["_roll_match_score"] = match.score
        item["_roll_fireproof"] = bool(match.data.get("_fireproof") or fire)
        item["_ledger_text"] = ledger_text_from_match(match.data)
        item["품목코드"] = match.data.get("관리코드")
        item["롤품명"] = match.data.get("_product") or product
        item["롤색상"] = match.data.get("_color") or color
    else:
        item["_roll_match"] = None
        item["_roll_match_score"] = 0.0
        prefix = str(item.get("롤구분") or "").upper()
        prefix = prefix if prefix in {"C", "R", "S"} else ""
        base = _strip_fire(str(product or ""))
        if fire:
            base = f"{base}방염"
        ledger = (("" if prefix == "C" else f"{prefix} ") + base).strip()
        if color:
            ledger += f" ({color_display(color)})"
        item["_ledger_text"] = ledger or str(product or color or "미확인")
    return item


def normalize_roll_order(order: dict, master: RollComboMaster, client_hint: str | None = None) -> dict:
    raw = str(order.get("전체원문") or "")
    client = client_hint or canonical_client(order.get("거래처"), raw)
    if client:
        order["거래처"] = client
    order["_product_mode"] = "roll_combo"
    order.setdefault("배송", {})
    if client and not order["배송"].get("방식"):
        order["배송"]["방식"] = ROLL_CLIENT_INFO[client]["delivery"]
    # 닫힌창/닫힌사이즈는 제작 메모가 아니므로 롤/콤비 기재사항에서 제외한다.
    order["전체기재사항"] = _clean_roll_note_text(order.get("전체기재사항"))
    for item in order.get("items") or []:
        item["기재사항"] = _clean_roll_note_text(item.get("기재사항"))
        item["손잡이방향"] = _direction(item.get("손잡이방향"))
        item["가로"] = _num(item.get("가로"))
        item["세로"] = _num(item.get("세로"))
        item["손잡이길이"] = _num(item.get("손잡이길이"))
        item["연창"] = False
        if client == "피카소" and not item.get("_picasso_swapped"):
            item["가로"], item["세로"] = item.get("세로"), item.get("가로")
            item["_picasso_swapped"] = True
        attach_match(item, master, client)
        # 롤스크린(R)은 세로 길이와 무관하게 기본 손잡이 길이가 150이다.
        # 주문서에 별도 길이가 적힌 경우에만 그 값을 우선한다.
        matched = item.get("_roll_match") or {}
        category = str(matched.get("_category") or item.get("롤구분") or "").upper()
        if category == "R" and item.get("손잡이길이") in (None, ""):
            item["손잡이길이"] = 150
    # 업체 기본 배송이 택배인 거래처는 '업체로 배송'이 기본 의미다.
    # 별도 고객 주소가 없는 경우 거래처 등록 주소를 사용하는 것이므로 주소 입력을 요구하지 않는다.
    if client and ROLL_CLIENT_INFO.get(client, {}).get("delivery") == "택배":
        delivery = order.setdefault("배송", {})
        if str(delivery.get("방식") or "택배").strip() == "택배" and not delivery.get("주소"):
            order["_delivery_to_client"] = True
    return order


def _head_map(ws):
    for r in range(1, min(ws.max_row, 20) + 1):
        vals = {str(ws.cell(r, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
        if "오더명" in vals and "품명" in vals and "원단명" in vals:
            return r, vals
    raise ValueError("안산)보노 롤/콤비 주문서 헤더를 찾지 못했습니다")


def _find_col(cols: dict, *keywords):
    for name, c in cols.items():
        compact = re.sub(r"\s+", "", name)
        if all(k in compact for k in keywords):
            return c
    return None


def parse_bono_roll(path: str | Path, master: RollComboMaster) -> list[dict]:
    """하비창 안산)보노 Excel 주문서 -> 주문 묶음."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    hr, cols = _head_map(ws)
    c_ship = cols.get("출고일")
    c_order = cols["오더명"]
    c_freight = cols.get("거래처(화물지점)")
    c_product = cols["품명"]
    c_color = cols["원단명"]
    c_size = _find_col(cols, "발주사이즈")
    c_width = _find_col(cols, "발주사이즈", "가로") or _find_col(cols, "가로") or c_size
    c_height = _find_col(cols, "발주사이즈", "세로") or _find_col(cols, "세로") or ((c_size + 1) if c_size else None)
    c_dir = _find_col(cols, "손잡이방향")
    c_handle = _find_col(cols, "길이")
    c_op = _find_col(cols, "작동방식") or _find_col(cols, "필증")
    c_place = _find_col(cols, "시공위치")
    c_note = _find_col(cols, "비고")
    c_addr = _find_col(cols, "주소")

    orders, current = [], None
    for r in range(hr + 1, ws.max_row + 1):
        order_name = str(ws.cell(r, c_order).value or "").strip()
        product = str(ws.cell(r, c_product).value or "").strip()
        color = str(ws.cell(r, c_color).value or "").strip()
        if order_name:
            if current and current.get("items"):
                normalize_roll_order(current, master, "안산)보노")
                orders.append(current)
            freight = str(ws.cell(r, c_freight).value or "").strip() if c_freight else ""
            explicit_address = str(ws.cell(r, c_addr).value or "").strip() if c_addr else None
            method, address, pay, receiver, phone, branch, _note1, is_bono = _bono_freight_info(
                freight, explicit_address=explicit_address, self_names=("보노",))
            # 일부 하비창 주문은 전화번호가 `받는사람:` 앞에 오므로 전체 원문에서 보완한다.
            if not phone:
                pm = re.search(r"01\d[-.\d]{7,}", freight)
                phone = pm.group(0) if pm else None
            receiver = re.sub(r"/(?:선불|착불|후불).*?$", "", str(receiver or "")).strip(" /-") or None
            if branch and not explicit_address:
                address = "-".join(x for x in (branch, receiver) if x) or branch
            ship_value = ws.cell(r, c_ship).value if c_ship else None
            ship_iso = None
            if hasattr(ship_value, "date"):
                ship_iso = ship_value.date().isoformat()
            elif ship_value:
                try:
                    ship_iso = str(ship_value)[:10]
                except Exception:
                    ship_iso = None
            current = {
                "거래처": "안산)보노",
                "주문번호": order_name,
                "고객명": None,
                "전체기재사항": None,
                "전체원문": freight,
                "배송": {"방식": method, "주소": address, "화물지점": branch,
                         "수령인": receiver, "연락처": phone, "선불착불": pay,
                         "발신": _sender_text(None)},
                "items": [],
                "_bono_receiver_is_bono": bool(is_bono),
                "_bono_branch": branch,
                "_product_mode": "roll_combo",
                "_ship_date": ship_iso,
            }
        if current is None or not (product or color):
            continue
        w = _num(ws.cell(r, c_width).value) if c_width else None
        h = _num(ws.cell(r, c_height).value) if c_height else None
        direction = _direction(ws.cell(r, c_dir).value) if c_dir else None
        handle = _num(ws.cell(r, c_handle).value) if c_handle else None
        operation = str(ws.cell(r, c_op).value or "").strip() if c_op else ""
        place_raw = str(ws.cell(r, c_place).value or "").strip() if c_place else ""
        note = str(ws.cell(r, c_note).value or "").strip() if c_note else ""
        note = _clean_roll_note_text(note) or ""
        raw = " / ".join(x for x in (product, color, operation, place_raw, note) if x)
        item = {
            "제품군": "롤/콤비", "_product_group": "roll_combo",
            "롤품명": product, "롤색상": color,
            "가로": w, "세로": h, "수량": 1,
            "손잡이방향": direction, "손잡이길이": handle,
            "작동방식필증": operation or None,
            "설치장소": simplify_place(place_raw),
            "_roll_place_raw": place_raw or None,
            "기재사항": note or None,
            "원문": raw, "연창": False, "창개수": 1,
            "좌개수": 1 if direction == "좌" else 0,
            "우개수": 1 if direction == "우" else 0,
            "_roll_fireproof": "방염" in raw or "필증" in operation,
        }
        current["items"].append(item)
        if note:
            parts = [x.strip() for x in str(current.get("전체기재사항") or "").split("/") if x.strip()]
            if note not in parts:
                parts.append(note)
            current["전체기재사항"] = "/".join(parts)
    if current and current.get("items"):
        normalize_roll_order(current, master, "안산)보노")
        orders.append(current)
    wb.close()
    return orders


ROLL_RESPONSE_FORMAT = {
    "type": "json_schema",
    "json_schema": {
        "name": "fitorder_roll_combo",
        "strict": True,
        "schema": {
            "type": "object", "additionalProperties": False,
            "properties": {
                "거래처": {"type": ["string", "null"]},
                "주문번호": {"type": ["string", "null"]},
                "고객명": {"type": ["string", "null"]},
                "전체기재사항": {"type": ["string", "null"]},
                "전체원문": {"type": ["string", "null"]},
                "배송": {
                    "type": "object", "additionalProperties": False,
                    "properties": {
                        "방식": {"type": ["string", "null"]}, "주소": {"type": ["string", "null"]},
                        "화물지점": {"type": ["string", "null"]}, "수령인": {"type": ["string", "null"]},
                        "연락처": {"type": ["string", "null"]}, "선불착불": {"type": ["string", "null"]},
                        "전달사항": {"type": ["string", "null"]}, "발신": {"type": ["string", "null"]}
                    },
                    "required": ["방식", "주소", "화물지점", "수령인", "연락처", "선불착불", "전달사항", "발신"]
                },
                "items": {
                    "type": "array",
                    "items": {
                        "type": "object", "additionalProperties": False,
                        "properties": {
                            "롤구분": {"type": ["string", "null"]},
                            "품명": {"type": ["string", "null"]}, "색상": {"type": ["string", "null"]},
                            "가로": {"type": ["number", "null"]}, "세로": {"type": ["number", "null"]},
                            "수량": {"type": ["number", "string", "null"]},
                            "손잡이방향": {"type": ["string", "null"]},
                            "손잡이길이": {"type": ["number", "null"]},
                            "작동방식필증": {"type": ["string", "null"]},
                            "시공위치": {"type": ["string", "null"]},
                            "기재사항": {"type": ["string", "null"]}, "원문": {"type": ["string", "null"]}
                        },
                        "required": ["롤구분", "품명", "색상", "가로", "세로", "수량", "손잡이방향", "손잡이길이", "작동방식필증", "시공위치", "기재사항", "원문"]
                    }
                }
            },
            "required": ["거래처", "주문번호", "고객명", "전체기재사항", "전체원문", "배송", "items"]
        }
    }
}

ROLL_SYSTEM_PROMPT = """너는 블라인드 제조공장의 롤스크린/콤비/쉐이드 발주서 판독기다.
이미지에 실제로 보이는 값만 JSON으로 옮겨라. 모르는 값은 null이다.
품명과 색상(원단명)을 반드시 분리한다. 가로/세로는 cm 숫자다.
손잡이방향은 좌/우만, 길이는 숫자만 적는다. 방염/필증 표기는 작동방식필증에 보존한다.
연창 여부는 절대로 추정하지 않는다.
거래처 후보: 안산)보노, 미래가공, 천안)채원, 세일인테리어, 진천)충북혁신, 음성)충북혁신,
논공)현대종합장식, 원주)창과방, 피카소, 제주)우림, 제주)중문.
충북혁신 주문에서 진천점이면 진천)충북혁신, 음성점이면 음성)충북혁신으로 적는다.
피카소는 원본에 보이는 가로/세로를 그대로 JSON에 적고, 프로그램이 이후 뒤집는다.
배송 주소/화물지점/수령인/연락처/선불착불/전달사항도 보이는 경우 누락하지 않는다.
"""


def extract_roll_images(image_paths, master: RollComboMaster, client_hint: str | None = None, model=None) -> dict:
    from extract import MAX_RETRY, MODEL, _encode, client
    if client is None:
        raise RuntimeError("API 키가 없습니다. .env 파일을 확인해 주세요.")
    if isinstance(image_paths, (str, Path)):
        image_paths = [image_paths]
    content = [{"type": "text", "text": "롤/콤비 주문서를 표준 JSON으로 추출하세요."}]
    for p in image_paths:
        content.append({"type": "image_url", "image_url": {
            "url": f"data:image/png;base64,{_encode(str(p))}", "detail": "high"}})
    last = None
    for _attempt in range(MAX_RETRY):
        try:
            resp = client.chat.completions.create(
                model=model or MODEL,
                messages=[{"role": "system", "content": ROLL_SYSTEM_PROMPT},
                          {"role": "user", "content": content}],
                response_format=ROLL_RESPONSE_FORMAT, temperature=0)
            data = json.loads(resp.choices[0].message.content)
            if client_hint:
                data["거래처"] = client_hint
            # 표준 내부 키로 치환
            for it in data.get("items") or []:
                it["롤품명"] = it.pop("품명", None)
                it["롤색상"] = it.pop("색상", None)
                it["설치장소"] = it.pop("시공위치", None)
                it["_roll_fireproof"] = "방염" in " ".join(str(it.get(k) or "") for k in ("작동방식필증", "기재사항", "원문", "롤품명"))
                it["_product_group"] = "roll_combo"
                it["연창"] = False
            return normalize_roll_order(data, master, client_hint)
        except Exception as exc:
            last = exc
    raise last


def extract_roll_pdf(pdf_path: str | Path, master: RollComboMaster, client_hint: str | None = None, model=None) -> list[dict]:
    try:
        import pymupdf
    except ImportError as exc:
        raise RuntimeError("PDF 처리 모듈이 없습니다. PDF기능_설치.bat를 실행해 주세요.") from exc
    doc = pymupdf.open(str(pdf_path))
    try:
        if doc.needs_pass:
            raise ValueError("암호가 걸린 PDF는 읽을 수 없습니다")
        if doc.page_count > 50:
            raise ValueError("PDF는 한 번에 50페이지까지 처리할 수 있습니다")
        with tempfile.TemporaryDirectory(prefix="fitorder_roll_pdf_") as td:
            paths = []
            for i, page in enumerate(doc):
                pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                p = Path(td) / f"page_{i+1:03d}.png"
                pix.save(str(p)); paths.append(p)
            return [extract_roll_images(paths, master, client_hint=client_hint, model=model)]
    finally:
        doc.close()


def validate_roll_order(order: dict, master: RollComboMaster):
    issues = []
    client = order.get("거래처")
    if client not in ROLL_CLIENT_INFO:
        issues.append(("red", "거래처", "롤/콤비 거래처를 확인해 주세요.", None))
    delivery = order.get("배송") or {}
    default_mode = str((ROLL_CLIENT_INFO.get(client) or {}).get("delivery") or "")
    mode = str(delivery.get("방식") or default_mode or "")
    if mode in {"택배", "화물"}:
        has_destination = bool(str(
            (delivery.get("화물지점") or delivery.get("주소")) if mode == "화물"
            else delivery.get("주소") or ""
        ).strip())
        default_is_company = bool(order.get("_delivery_to_client")) or (
            not has_destination and mode == default_mode
        )
        if not has_destination and not default_is_company:
            issues.append(("yellow", "주소", f"별도 {mode} 배송 목적지를 확인해 주세요.", None))
    for i, it in enumerate(order.get("items") or [], 1):
        if not it.get("가로") or not it.get("세로"):
            issues.append(("red", "사이즈", "가로/세로를 확인해 주세요.", i))
        if not it.get("_roll_match"):
            issues.append(("red", "품목", f"품목장 매칭 실패: {it.get('롤품명') or ''} {it.get('롤색상') or ''}".strip(), i))
        else:
            score = float(it.get("_roll_match_score") or 0)
            if score < 0.78:
                issues.append(("yellow", "품목확인", f"품목장 유사 매칭 {score:.0%}: {it.get('_ledger_text')}", i))
            if not master.find_order_item(it, client).get("단가"):
                issues.append(("red", "단가", "업체 출고가가 0원입니다. 품목장을 확인해 주세요.", i))
    return issues


def client_label(client: str | None) -> str:
    if not client:
        return ""
    info = ROLL_CLIENT_INFO.get(client) or {}
    return f"{client} · {info.get('price','')} · 기본 {info.get('delivery','')}".strip(" ·")
