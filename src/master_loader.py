"""공통 품목장 로더.

블라인드/롤·콤비/홀딩도어가 서로 다른 파일 형식을 사용하더라도
이 모듈을 통해 ``list[dict]`` 형태로 통일해서 읽는다.

운영 원칙
- 외부 ``data/master`` 품목장이 있으면 항상 우선한다.
- SRC만 교체하는 환경을 위해 제품군별 ``src/*_master`` 폴백을 허용한다.
- 컬럼 순서가 바뀌어도 헤더명으로 읽는다.
- .xlsx/.xls 모두 지원한다.
"""
from __future__ import annotations

import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

import openpyxl
import xlrd


class MasterLoadError(RuntimeError):
    pass


@dataclass(frozen=True)
class LoadedTable:
    path: Path
    sheet_name: str
    header_row: int
    headers: tuple[str, ...]
    rows: tuple[dict, ...]


def source_dir() -> Path:
    return Path(__file__).resolve().parent


def app_root() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return source_dir().parent


def _unique_paths(paths: Iterable[Path]) -> list[Path]:
    result: list[Path] = []
    seen: set[str] = set()
    for p in paths:
        try:
            key = str(p.resolve())
        except Exception:
            key = str(p)
        if key not in seen:
            seen.add(key)
            result.append(p)
    return result


def master_dir_candidates(kind: str) -> list[Path]:
    """제품군별 품목장 폴더 후보를 운영 우선순위대로 반환한다."""
    kind = str(kind or "").strip().lower()
    root = app_root()
    here = source_dir()
    if kind == "blind":
        return _unique_paths([
            root / "data" / "master",
            Path.cwd() / "data" / "master",
            here / "blind_master",
        ])
    return _unique_paths([
        root / "data" / "master" / kind,
        Path.cwd() / "data" / "master" / kind,
        here / f"{kind}_master",
        here / "data" / "master" / kind,
    ])


def resolve_master_dir(kind: str, extensions: Sequence[str] = (".xlsx", ".xls")) -> Path:
    checked = []
    extset = {e.lower() for e in extensions}
    for directory in master_dir_candidates(kind):
        checked.append(directory)
        if not directory.is_dir():
            continue
        if any(p.is_file() and p.suffix.lower() in extset for p in directory.iterdir()):
            return directory
    raise FileNotFoundError(
        f"{kind} 품목장 폴더를 찾을 수 없습니다. 확인한 위치:\n - "
        + "\n - ".join(str(p) for p in checked)
    )


def resolve_master_file(kind: str, filenames: Sequence[str] | str) -> Path:
    """품목장 파일을 안전하게 찾는다.

    1) 알려진 파일명을 우선 사용한다.
    2) 압축 해제/이동 과정에서 한글 파일명이 변형되었더라도, 해당 제품군
       폴더에 스프레드시트가 정확히 하나뿐이면 그 파일을 자동 채택한다.

    따라서 배포 환경에서 파일명 인코딩이 달라져도 품목장 자체가 존재하면
    불필요한 FileNotFoundError가 발생하지 않는다.
    """
    if isinstance(filenames, str):
        filenames = (filenames,)
    filenames = tuple(str(x) for x in filenames)
    checked = []
    allowed_exts = {Path(name).suffix.lower() for name in filenames if Path(name).suffix}
    if not allowed_exts:
        allowed_exts = {".xlsx", ".xls"}

    for directory in master_dir_candidates(kind):
        for filename in filenames:
            path = directory / filename
            checked.append(path)
            if path.is_file():
                return path

        # 파일명은 달라졌지만 품목장 파일 자체가 하나만 있는 경우의 안전 폴백.
        if directory.is_dir():
            sheets = sorted(
                p for p in directory.iterdir()
                if p.is_file() and p.suffix.lower() in allowed_exts
            )
            if len(sheets) == 1:
                return sheets[0]

    raise FileNotFoundError(
        f"{kind} 품목장 파일을 찾을 수 없습니다. 확인한 위치:\n - "
        + "\n - ".join(str(p) for p in checked)
    )


def _clean_header(value) -> str:
    return str(value or "").replace("\n", " ").strip()


def _find_header(matrix: list[list], required_headers: Sequence[str], max_scan: int = 20):
    required = {str(x).strip() for x in required_headers if str(x).strip()}
    for idx, values in enumerate(matrix[:max_scan]):
        headers = [_clean_header(v) for v in values]
        if required.issubset(set(headers)):
            return idx, headers
    return None, None


def _rows_from_matrix(matrix: list[list], required_headers: Sequence[str], path: Path, sheet_name: str) -> LoadedTable | None:
    header_row, headers = _find_header(matrix, required_headers)
    if headers is None:
        return None
    rows = []
    for values in matrix[header_row + 1:]:
        row = {}
        nonempty = False
        for i, header in enumerate(headers):
            if not header:
                continue
            value = values[i] if i < len(values) else None
            row[header] = value
            if value not in (None, ""):
                nonempty = True
        if nonempty:
            rows.append(row)
    return LoadedTable(path, sheet_name, header_row + 1, tuple(headers), tuple(rows))


def read_tables(path: str | Path, required_headers: Sequence[str], preferred_sheet: str | None = None) -> list[LoadedTable]:
    """한 스프레드시트 파일의 모든 유효 시트를 헤더 기반으로 읽는다."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in {".xlsx", ".xls"}:
        raise MasterLoadError(f"지원하지 않는 품목장 형식입니다: {path.name}")
    tables: list[LoadedTable] = []

    if suffix == ".xlsx":
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        names = list(wb.sheetnames)
        if preferred_sheet in names:
            names.remove(preferred_sheet)
            names.insert(0, preferred_sheet)
        for name in names:
            ws = wb[name]
            matrix = [list(row) for row in ws.iter_rows(values_only=True)]
            table = _rows_from_matrix(matrix, required_headers, path, name)
            if table:
                tables.append(table)
        wb.close()
    else:
        book = xlrd.open_workbook(str(path))
        sheets = list(book.sheets())
        if preferred_sheet:
            sheets.sort(key=lambda sh: 0 if sh.name == preferred_sheet else 1)
        for sh in sheets:
            matrix = [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
            table = _rows_from_matrix(matrix, required_headers, path, sh.name)
            if table:
                tables.append(table)

    if not tables:
        required = ", ".join(required_headers)
        raise MasterLoadError(f"{path.name}에서 필수 헤더({required})가 있는 시트를 찾지 못했습니다.")
    return tables


def read_rows(path: str | Path, required_headers: Sequence[str], preferred_sheet: str | None = None) -> list[dict]:
    rows: list[dict] = []
    for table in read_tables(path, required_headers, preferred_sheet):
        for row in table.rows:
            item = dict(row)
            item.setdefault("_master_file", table.path.name)
            item.setdefault("_master_sheet", table.sheet_name)
            rows.append(item)
    return rows


def read_directory(directory: str | Path, required_headers: Sequence[str], extensions: Sequence[str] = (".xlsx", ".xls")) -> list[dict]:
    directory = Path(directory)
    extset = {e.lower() for e in extensions}
    files = sorted(p for p in directory.iterdir() if p.is_file() and p.suffix.lower() in extset)
    if not files:
        raise FileNotFoundError(f"품목장 파일이 없습니다: {directory}")
    rows: list[dict] = []
    errors: list[str] = []
    for path in files:
        try:
            rows.extend(read_rows(path, required_headers))
        except MasterLoadError as exc:
            errors.append(str(exc))
    if not rows:
        detail = "\n".join(errors)
        raise MasterLoadError(f"읽을 수 있는 품목이 없습니다: {directory}" + (f"\n{detail}" if detail else ""))
    return rows
