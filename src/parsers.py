"""
엑셀 발주 파서 (휴안 · DI · 유앤) — AI 를 거치지 않고 직접 읽는다.
출력 형태는 extract.extract_order() 와 동일하다.
"""
import re

import openpyxl

from rules import DI_COLOR, DI_COLOR_UNSURE

SIZE = re.compile(r"(\d+(?:\.\d+)?)\s*[*x×X]\s*(\d+(?:\.\d+)?)")


def _blank_item():
    return {"품목코드": None, "색상원문": None, "타입": "C자", "종류": "투코드",
            "가로": None, "세로": None, "수량": None, "손잡이방향": None,
            "손잡이길이": None, "연창": False, "설치장소": None,
            "기재사항": None, "예외품목": None, "원문": None,
            "창개수": 1, "좌개수": None, "우개수": None,
            "확신도": {"가로": 1.0, "세로": 1.0, "품목코드": 1.0, "손잡이": 1.0}}


def _shell(client):
    return {"거래처": client, "발주일": None, "고객명": None, "items": [],
            "배송": {"방식": None, "화물지점": None, "주소": None,
                     "수령인": None, "연락처": None, "선불착불": None,
                     "전달사항": None},
            "전체원문": None, "전체기재사항": None}


# ─────────────────────────────────────────────
# 휴안
# ─────────────────────────────────────────────
def _huan_product(s):
    """'YL500 L원코드' -> (500, L자, 원코드)"""
    s = str(s or "").strip()
    m = re.search(r"([A-Z]{1,2})?\s*(\d{3}[A-Za-z]*)", s)
    code = m.group(2) if m else None
    type_ = "L자" if re.search(r"\bL\s*(?=원|투|셔|18|타입|자)", s) or "L18" in s else "C자"
    if "셔터" in s:
        kind = "셔터"
    elif "원코드" in s or re.search(r"원\s*코드", s):
        kind = "원코드"
    else:
        kind = "투코드"
    return code, type_, kind


SIDO_SHORT = {"서울특별시": "서울시", "부산광역시": "부산시", "대구광역시": "대구시",
              "인천광역시": "인천시", "광주광역시": "광주시", "대전광역시": "대전시",
              "울산광역시": "울산시", "세종특별자치시": "세종시"}
DO = ("경기도", "강원도", "강원특별자치도", "충청북도", "충청남도", "전라북도",
      "전북특별자치도", "전라남도", "경상북도", "경상남도", "제주특별자치도")


def _short_addr(addr):
    """'경상남도 진주시 초전북로 151 (초전동) 인사인광고기획'
         -> '진주시 초전북로 151, 인사인광고기획'
       '서울특별시 영등포구 도신로 161 (도림동, 정성빌) 201호'
         -> '서울시 영등포구 도신로 161, 정성빌 201호'"""
    s = str(addr or "").strip()
    s = re.sub(r"^\s*\(\s*\d{5}\s*\)\s*", "", s)
    for k, v in SIDO_SHORT.items():          # 광역시/특별시는 줄여서 유지
        if s.startswith(k):
            s = v + s[len(k):]
            break
    else:
        for d in DO:                          # 도는 삭제
            if s.startswith(d):
                s = s[len(d):]
                break

    def keep_building(m):
        """괄호 안에서 '○○동' 은 버리고 건물명만 남긴다"""
        parts = [x.strip() for x in m.group(1).split(",")]
        keep = [x for x in parts if not re.search(r"[동읍면리]$", x)]
        return (", " + " ".join(keep)) if keep else ""

    s = re.sub(r"\s*\(([^)]*)\)", keep_building, s)
    s = re.sub(r"\s{2,}", " ", s).strip()
    s = re.sub(r"(\d+)동\s*(\d+)호\b", r"\1-\2", s)
    s = re.sub(r"아파트\b", "A", s)
    if "," not in s:                      # 끝의 상호/건물명 앞에 콤마
        head, _, tail = s.rpartition(" ")
        if head and tail and not re.search(r"\d", tail):
            s = f"{head}, {tail}"
    return s


def parse_huan(path):
    """휴안 발주 엑셀 -> 주문 리스트 (수취인 단위로 분리)"""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    orders, cur, last_prod = [], None, (None, "C자", "투코드", "")
    for r in range(2, ws.max_row + 1):
        g = lambda c: ws.cell(r, c).value
        size = g(9)
        if not size:
            continue
        name = g(3)
        if name:                                  # 새 주문 시작
            cur = _shell("휴안")
            d = g(1)
            cur["발주일"] = d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else None
            cur["고객명"] = str(name).strip()
            cur["배송"].update({
                "방식": "택배", "주소": _short_addr(g(6)),
                "수령인": str(name).strip(),
                "연락처": str(g(4) or "").strip(),
                "선불착불": "선불",
                "전달사항": (str(g(12)).strip() if g(12) else None),
            })
            cur["전체기재사항"] = "포장비용"       # 휴안은 창당 포장비
            orders.append(cur)
        if cur is None:
            continue
        m = SIZE.search(str(size))
        it = _blank_item()
        if m:
            it["가로"], it["세로"] = float(m.group(1)), float(m.group(2))
        if g(7):
            code, type_, kind = _huan_product(g(7))
            last_prod = (code, type_, kind, str(g(7)).strip())
        elif cur["items"]:                       # 상품명 공란 = 위 행과 동일
            code, type_, kind, _ = last_prod
        else:
            code, type_, kind = None, "C자", "투코드"
        it.update({"품목코드": code,
                   "색상원문": (str(g(7)).strip() if g(7) else last_prod[3]),
                   "타입": type_, "종류": kind,
                   "손잡이방향": (str(g(10)).strip() if g(10) else None),
                   "기재사항": cur["고객명"], "원문": f"{g(7)} {size} {g(10)}"})
        n = g(11)
        it["창개수"] = int(n) if isinstance(n, (int, float)) and n else 1
        cur["items"].append(it)
    return orders


# ─────────────────────────────────────────────
# DI
# ─────────────────────────────────────────────
DI_KIND = {"원": "원코드", "투": "투코드", "셔터": "셔터",
           "원코드": "원코드", "투코드": "투코드"}

# 헤더 이름 -> 내부 필드 (파일마다 열 위치가 달라 헤더로 찾는다)
DI_HEAD = {
    "수령인": "recv", "분류": "cls", "원단/색상": "color",
    "원/투": "kind", "커버/점보": "cover", "사이즈": "size",
    "줄": "line", "피스": "piece", "비고": "note",
}


def _di_header(ws):
    """헤더 행을 찾아 {필드: 열번호} 반환. 좌측(원본) 영역만 사용."""
    for r in range(1, 8):
        vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, 20)]
        if "사이즈" not in vals:
            continue
        col, seen = {}, set()
        for i, v in enumerate(vals, 1):
            f = DI_HEAD.get(v)
            if f and f not in seen:        # 같은 이름이 우측에 또 나오면 무시
                col[f] = i
                seen.add(f)
        if "size" in col and "color" in col:
            return r, col
    return None, {}


def parse_di(path, sheet="원본"):
    """DI 발주 엑셀(좌측 원본 영역) -> 주문 1건"""
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        sheet = "장부" if "장부" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet]
    hr, col = _di_header(ws)
    if not col:
        raise ValueError("DI 발주서 형식을 인식하지 못했습니다 (헤더를 찾을 수 없음)")

    g = lambda r, f: (ws.cell(r, col[f]).value if f in col else None)
    o = _shell("DI")
    for r in range(hr + 1, ws.max_row + 1):
        size = g(r, "size")
        if not size:
            continue
        m = SIZE.search(str(size))
        if not m:
            continue
        it = _blank_item()
        cname = str(g(r, "color") or "").strip()
        it["색상원문"] = cname
        it["품목코드"] = DI_COLOR.get(cname)
        if cname in DI_COLOR_UNSURE or (cname and not it["품목코드"]):
            it["확신도"]["품목코드"] = 0.4      # 미등록 색상명 -> 담당자 확인
        # 종류/부속 판정 — 헤더 이름이 파일마다 달라 값으로 판단한다.
        # '원'/'투'/'셔터' 이면 종류, 그 밖의 문구면 커버·점보 등 부속.
        ks = ""
        for f in ("kind", "cover"):
            v = str(g(r, f) or "").strip()
            if not v or v in ("0", "-"):
                continue
            if v in DI_KIND:
                ks = v
            else:
                it["예외품목"] = "부속"
        it["종류"] = DI_KIND.get(ks, "투코드")
        it["타입"] = "L자"                       # DI 는 L타입 발주
        it["가로"], it["세로"] = float(m.group(1)), float(m.group(2))
        d = str(g(r, "line") or "").strip()
        if d.startswith(("좌", "우")):
            it["손잡이방향"] = d[0]
            n = re.search(r"(\d+)", d)
            if n:
                it["손잡이길이"] = int(n.group(1))
        recv = str(g(r, "recv") or "").strip()
        note = str(g(r, "note") or "").strip()
        piece = str(g(r, "piece") or "").strip()
        # DI 장부의 기재사항에는 수령인만 표기한다.
        # 비고·피스는 원문에만 남겨 필요할 때 확인할 수 있게 한다.
        it["기재사항"] = recv if recv and recv != "None" else None
        it["원문"] = f"{cname} {ks} {size} {d} {note} {piece}".strip()
        o["items"].append(it)
    o["배송"]["방식"] = None                      # DI/SP 출고
    if not o["items"]:
        raise ValueError("DI 발주서에서 유효한 주문 행을 찾지 못했습니다")
    return o


# ─────────────────────────────────────────────
# 유앤아이티엔에스
# ─────────────────────────────────────────────
def _unit_directions(text, qty):
    """'좌2 우2', '좌/우', '오른쪽'을 창별 방향으로 펼친다."""
    s = str(text or "").strip()
    left = sum(int(x or 1) for x in re.findall(r"좌(?:측)?\s*(\d*)", s))
    right = sum(int(x or 1) for x in re.findall(r"우(?:른쪽|측)?\s*(\d*)", s))
    if "왼쪽" in s and not left:
        left = 1
    if "오른쪽" in s and not right:
        right = 1
    if not left and not right:
        return [None] * qty
    dirs = ["좌"] * left + ["우"] * right
    return (dirs + [None] * qty)[:qty]


def _unit_code(value):
    m = re.search(r"(?<!\d)(\d{3}[A-Za-z]*)(?!\d)", str(value or ""))
    return m.group(1) if m else None


def parse_unitns(path):
    """UNITNS 발주 엑셀 -> 고객명 블록별 주문."""
    wb = openpyxl.load_workbook(path, data_only=True)
    orders = []
    for ws in wb.worksheets:
        if "UNITNS" not in str(ws.cell(1, 1).value or "").upper():
            continue
        r = 1
        while r <= ws.max_row:
            if str(ws.cell(r, 1).value or "").strip() != "고객명":
                r += 1
                continue
            customer = str(ws.cell(r, 2).value or "").strip()
            address = str(ws.cell(r, 7).value or "").strip()
            phone = str(ws.cell(r + 1, 2).value or "").strip()
            delivery_note = str(ws.cell(r + 2, 2).value or "").strip()
            header = r + 3
            if str(ws.cell(header, 1).value or "").strip() != "NO":
                r += 1
                continue
            o = _shell("유앤")
            o["고객명"] = customer or None
            o["전체기재사항"] = "피스"
            o["_unit_accessories"] = []
            o["배송"].update({
                "방식": "택배", "주소": _short_addr(address),
                "수령인": customer or None, "연락처": phone or None,
                "선불착불": "선불", "전달사항": delivery_note or None,
            })
            rr = header + 1
            while rr <= ws.max_row:
                first = str(ws.cell(rr, 1).value or "").strip()
                if first == "고객명" or (rr > header + 1 and first == "NO"):
                    break
                code = _unit_code(ws.cell(rr, 3).value)
                w, h = ws.cell(rr, 4).value, ws.cell(rr, 5).value
                if code and isinstance(w, (int, float)) and isinstance(h, (int, float)):
                    raw_qty = ws.cell(rr, 7).value
                    qty = max(1, int(raw_qty)) if isinstance(raw_qty, (int, float)) else 1
                    dirs = _unit_directions(ws.cell(rr, 6).value, qty)
                    product = str(ws.cell(rr, 2).value or "")
                    item_note = str(ws.cell(rr, 8).value or "").strip()
                    screw = None
                    if "콘크리트" in item_note:
                        screw = "콘크리트"
                    elif "석고" in item_note:
                        screw = "석고"
                    if screw:
                        screw_count = 2 if float(w) < 150 else (3 if float(w) < 200 else 4)
                        o["_unit_accessories"].append({
                            "품명": f"{screw}({screw_count})", "세트": qty,
                        })
                    type_ = "L자" if re.search(r"L\s*(?:자|형|타입|18|21)", product, re.I) else "C자"
                    for direction in dirs:
                        it = _blank_item()
                        it.update({"품목코드": code, "색상원문": str(ws.cell(rr, 3).value),
                                   "타입": type_, "종류": "투코드",
                                   "가로": float(w), "세로": float(h),
                                   "손잡이방향": direction,
                                   "기재사항": None,
                                   "원문": " ".join(str(ws.cell(rr, c).value or "")
                                                      for c in range(2, 9)).strip()})
                        o["items"].append(it)
                rr += 1
            if o["items"]:
                orders.append(o)
            r = max(rr, r + 1)
    wb.close()
    if not orders:
        raise ValueError("유앤아이티엔에스 발주서에서 유효한 주문을 찾지 못했습니다")
    return orders


# ─────────────────────────────────────────────
# 보노 · 미래가공
# ─────────────────────────────────────────────
def _head_map(ws, required):
    for r in range(1, min(ws.max_row, 12) + 1):
        vals = {str(ws.cell(r, c).value or "").replace("\n", "").strip(): c
                for c in range(1, ws.max_column + 1)}
        if all(k in vals for k in required):
            return r, vals
    return None, {}


def _product_parts(product, color, operation):
    text = " ".join(str(x or "") for x in (product, color, operation))
    code = _unit_code(color) or _unit_code(text)
    type_ = "L자" if re.search(r"L\s*(?:자|형|타입|18|21)|알루미늄L", text, re.I) else "C자"
    if "셔터" in text:
        kind = "셔터"
    elif "원코드" in text:
        kind = "원코드"
    else:
        kind = "투코드"
    return code, type_, kind


def _direction(value):
    s = str(value or "").strip()
    if s.startswith(("좌", "왼")):
        return "좌"
    if s.startswith(("우", "오른")):
        return "우"
    return None


def _freight_info(value):
    """'안산초지점/착불/받는사람:보노(010...)'을 배송 정보로 분리."""
    s = str(value or "").strip()
    if not s:
        return "화물", None, None
    if "택배" in s:
        method = "택배"
    elif "배달" in s:
        method = "배달"
    else:
        method = "화물"
    first = s.split("/")[0].strip().replace("점", "")
    m = re.search(r"받는\s*사람\s*:\s*([^/(]+)?\s*\(?\s*(01\d[-\d]+)", s)
    receiver = (m.group(1).strip() if m and m.group(1) else "")
    phone = m.group(2).strip() if m else ""
    address = "-".join(x for x in (first, receiver) if x)
    if phone:
        address = f"{address} {phone}".strip()
    pay = "선불" if "선불" in s else ("착불" if "착불" in s else None)
    return method, address or None, pay


def _sender_text(value, default="보노 010-2478-9290"):
    s = str(value or "").strip()
    if not s:
        return default
    s = re.sub(r"^발신\s*:\s*", "", s)
    s = s.replace("(", " ").replace(")", "")
    return re.sub(r"\s+", " ", s).strip()


def _finish_excel_order(order):
    if not order:
        return None
    items = order.get("items") or []
    if len(items) > 1:
        for it in items:
            if it.get("종류") == "원코드":
                it["연창"] = True
    return order if items else None


def parse_bono(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    hr, col = _head_map(ws, ("오더명", "거래처(화물지점)", "품명", "원단명"))
    if not col:
        wb.close()
        raise ValueError("보노 발주서 헤더를 찾지 못했습니다")
    orders, cur = [], None
    for r in range(hr + 1, ws.max_row + 1):
        product = ws.cell(r, col["품명"]).value
        color = ws.cell(r, col["원단명"]).value
        size_col = col.get("발주사이즈")
        w = ws.cell(r, size_col).value if size_col else None
        h = ws.cell(r, size_col + 1).value if size_col else None
        if not (product and isinstance(w, (int, float)) and isinstance(h, (int, float))):
            continue
        order_no = str(ws.cell(r, col["오더명"]).value or "").strip()
        if order_no:
            done = _finish_excel_order(cur)
            if done:
                orders.append(done)
            cur = _shell("보노")
            cur["주문번호"] = order_no.replace("\n", "")
            d = ws.cell(r, col["출고일"]).value if "출고일" in col else None
            if hasattr(d, "date"):
                cur["_ship_date"] = d.date().isoformat()
            method, address, pay = _freight_info(
                ws.cell(r, col["거래처(화물지점)"]).value)
            cur["배송"].update({"방식": method, "주소": address,
                                "선불착불": pay, "발신": "보노 010-2478-9290"})
        if cur is None:
            continue
        operation = ws.cell(r, col.get("작동방식필증", 0)).value \
            if col.get("작동방식필증") else None
        code, type_, kind = _product_parts(product, color, operation)
        it = _blank_item()
        it.update({"품목코드": code, "색상원문": str(color or ""),
                   "타입": type_, "종류": kind, "가로": float(w),
                   "세로": float(h),
                   "손잡이방향": _direction(ws.cell(r, col.get("손잡이방향", 0)).value
                                             if col.get("손잡이방향") else None),
                   "손잡이길이": ws.cell(r, col.get("길이", 0)).value
                                  if col.get("길이") else None,
                   "설치장소": str(ws.cell(r, col.get("시공위치", 0)).value or "").strip() or None,
                   "원문": " ".join(str(ws.cell(r, c).value or "")
                                      for c in range(1, ws.max_column + 1)).strip()})
        cur["items"].append(it)
    done = _finish_excel_order(cur)
    if done:
        orders.append(done)
    wb.close()
    if not orders:
        raise ValueError("보노 발주서에서 유효한 주문을 찾지 못했습니다")
    return orders


def parse_future(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    hr, col = _head_map(ws, ("고객명", "태영품명", "색상", "작동방식"))
    if not col:
        wb.close()
        raise ValueError("미래가공 발주서 헤더를 찾지 못했습니다")
    cur = None
    for r in range(hr + 1, ws.max_row + 1):
        product = ws.cell(r, col["태영품명"]).value
        w = ws.cell(r, col.get("가로", 0)).value if col.get("가로") else None
        h = ws.cell(r, col.get("세로", 0)).value if col.get("세로") else None
        if product and isinstance(w, (int, float)) and isinstance(h, (int, float)):
            order_no = str(ws.cell(r, col["고객명"]).value or "").strip()
            if cur is None:
                cur = _shell("미래가공")
                cur["주문번호"] = order_no.replace("\n", "")
                d = ws.cell(r, col.get("출고일", 0)).value if col.get("출고일") else None
                if hasattr(d, "date"):
                    cur["_ship_date"] = d.date().isoformat()
            color = ws.cell(r, col["색상"]).value
            operation = ws.cell(r, col["작동방식"]).value
            code, type_, kind = _product_parts(product, color, operation)
            it = _blank_item()
            it.update({"품목코드": code, "색상원문": str(color or ""),
                       "타입": type_, "종류": kind, "가로": float(w), "세로": float(h),
                       "손잡이방향": _direction(ws.cell(r, col.get("방향", 0)).value
                                                 if col.get("방향") else None),
                       "손잡이길이": ws.cell(r, col.get("줄길이", 0)).value
                                      if col.get("줄길이") else None,
                       "설치장소": str(ws.cell(r, col.get("비고", 0)).value or "").strip() or None,
                       "원문": " ".join(str(ws.cell(r, c).value or "")
                                          for c in range(1, ws.max_column + 1)).strip()})
            cur["items"].append(it)
        elif cur and "발신" in str(ws.cell(r, col["고객명"]).value or ""):
            sender = _sender_text(ws.cell(r, col["고객명"]).value)
            method, address, pay = _freight_info(ws.cell(r, col["태영품명"]).value)
            cur["배송"].update({"방식": method, "주소": address,
                                "선불착불": pay, "발신": sender})
    wb.close()
    done = _finish_excel_order(cur)
    if not done:
        raise ValueError("미래가공 발주서에서 유효한 주문을 찾지 못했습니다")
    return [done]


# ─────────────────────────────────────────────
def detect_excel_client(path):
    """파일명이 아닌 실제 시트/헤더로 DI·휴안·유앤을 구분한다."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if any("UNITNS" in str(ws.cell(1, 1).value or "").upper()
               for ws in wb.worksheets):
            return "유앤"
        for ws in wb.worksheets:
            _, headers = _head_map(ws, ("오더명", "거래처(화물지점)", "품명", "원단명"))
            if headers:
                return "보노"
            _, headers = _head_map(ws, ("고객명", "태영품명", "색상", "작동방식"))
            if headers:
                return "미래가공"
        # DI는 '사이즈' + '원단/색상' 헤더가 있다.
        for ws in wb.worksheets:
            _, col = _di_header(ws)
            if col:
                return "DI"

        # 휴안은 고정 열 구조: C=수령인, G=상품명, I=규격.
        ws = wb.worksheets[0]
        for r in range(1, min(ws.max_row, 15) + 1):
            vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, 13)]
            joined = " ".join(vals)
            if "상품" in joined and any(x in joined for x in ("사이즈", "규격", "수령인")):
                return "휴안"
    finally:
        wb.close()
    raise ValueError("엑셀 발주서의 거래처를 자동 판별하지 못했습니다. 거래처를 직접 선택해 다시 시도해 주세요.")


def parse_excel(path, client=None):
    client = client or detect_excel_client(path)
    if client == "휴안":
        return parse_huan(path)
    if client == "DI":
        return [parse_di(path)]
    if client == "유앤":
        return parse_unitns(path)
    if client == "보노":
        return parse_bono(path)
    if client == "미래가공":
        return parse_future(path)
    raise ValueError(f"엑셀 파서 없음: {client}")


if __name__ == "__main__":
    import json, sys
    p = sys.argv[1]
    c = "DI" if "DI" in p.upper() else "휴안"
    for o in parse_excel(p, c):
        print(json.dumps(o, ensure_ascii=False, indent=2, default=str))
