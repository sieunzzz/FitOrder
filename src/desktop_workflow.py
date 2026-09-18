"""FitOrder PySide6 데스크톱 UI용 공통 워크플로.

Streamlit UI와 처리 엔진을 분리하기 위한 얇은 계층이다.
AI/파서/규칙/Excel 출력은 기존 모듈을 그대로 사용하고,
PySide6 화면은 이 모듈을 통해 주문을 읽고/표시하고/수정한다.
"""
from __future__ import annotations

import copy
import itertools
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

import openpyxl

from holding import (HOLDING_CLIENTS, apply_holding_feature_gate, holding_accessory_from_text,
                     holding_product_from_text, looks_like_holding_operation,
                     normalize_holding_operation)
from output import (
    DI_SHEET_GROUPS,
    _blind_counts,
    _di_paginate,
    _paginate_order_rows,
    apply_edit,
    build_erp,
    build_ledger,
    build_worksheet,
    expand_same_size_directions,
    synchronize_order_for_outputs,
    to_rows,
    _freight_ledger_text,
    _mix_group_key,
    split_room_parts,
    parse_freight_ledger_text,
)
from parsers import parse_excel
from rules import (BONO_LIKE_CLIENTS, BONO_STYLE_CLIENTS, CLIENT_INFO, NO_DELIVERY_LABEL_CLIENTS, Master,
                   clean_delivery_notice, validate)
from master_registry import get_master_for_mode, normalize_product_mode
from roll_combo import (ROLL_CLIENTS, ROLL_CLIENT_INFO, canonical_client as roll_canonical_client,
                        extract_roll_images, extract_roll_pdf,
                        normalize_roll_order, parse_bono_roll, validate_roll_order)

CLIENTS = list(CLIENT_INFO)
# 홀딩도어에서만 쓰는 거래처. 블라인드 거래처 목록/자동 판별에는 넣지 않는다.
# (미더스는 2026-09-13 사용자 요청으로 블라인드에서 제외. CLIENT_INFO 는 홀딩 단가 조회에 필요해 유지)
HOLDING_ONLY_CLIENTS = {"M", "구미)경남", "창문애", "한길"}
BLIND_CLIENTS = [c for c in CLIENTS if c not in HOLDING_ONLY_CLIENTS]


def _default_delivery_for_order(order: dict) -> str:
    client = order.get("거래처")
    if order.get("_product_mode") == "roll_combo":
        return str((ROLL_CLIENT_INFO.get(client) or {}).get("delivery") or "").strip()
    return str(CLIENT_INFO.get(client, (None, None, ""))[2] or "").strip()


def _has_explicit_destination(delivery: dict, mode: str) -> bool:
    if mode == "화물":
        return bool(str(delivery.get("화물지점") or delivery.get("주소") or "").strip())
    if mode == "택배":
        return bool(str(delivery.get("주소") or "").strip())
    return True


def _uses_registered_destination(order: dict, mode: str | None = None) -> bool:
    """별도 주소가 없고 실제 배송방식이 거래처 기본과 같으면 등록 주소1/화물지점을 사용한다."""
    delivery = order.get("배송") or {}
    default = _default_delivery_for_order(order)
    mode = str(mode or delivery.get("방식") or default or "").strip()
    if mode not in {"택배", "화물"}:
        return False
    if _has_explicit_destination(delivery, mode):
        return False
    return bool(default and mode == default)

FILE_CLIENT_ALIASES = {
    "두창블라인드": "DU", "두창": "DU",
    "대일산업": "DI", "대일": "DI",
    "루임트": "RT", "미더스": "M", "제이원": "JO",
    "대동산업": "DD", "대동": "DD",
    "유앤아이티엔에스": "유앤", "유앤아이": "유앤", "UNITNS": "유앤",
    "아지트": "아지트", "윈도우투모로우": "WT",
    "트루갤러리": "인천)트루", "인천)트루": "인천)트루",
    "미래가공": "미래가공", "이끌림": "보노", "보노": "보노",
    "미성텍스": "MS",
    # 홀딩도어 전용 거래처
    "구미)경남블라인드": "구미)경남", "구미)경남": "구미)경남", "경남블라인드": "구미)경남",
    "창문애아트블라인드": "창문애", "창문애": "창문애",
    "경산)한길산업": "한길", "한길산업": "한길", "한길": "한길",
}

LEDGER_COLUMNS = [
    "상호", "색상", "가로", "X", "세로", "수량", "방향", "길이",
    "특이", "기재사항1", "기재사항2", "출고일",
]

PRODUCT_EDIT_FIELD = {
    "상호": "상호",
    "색상": "색상",
    "가로": "가로",
    "세로": "세로",
    "수량": "수량",
    "방향": "방향",
    "길이": "길이",
    "특이": "특이",
    "기재사항1": "기재사항",
    "기재사항2": "기재사항2",
    "출고일": "출고일",
}


def client_label(client: str | None) -> str:
    if not client:
        return ""
    mark = CLIENT_INFO.get(client, (None, None, None, ""))[3]
    return f"{client} {mark}".strip()


def client_from_filename(path: Path) -> str | None:
    stem = path.stem.strip()
    upper = stem.upper()
    for name, code in FILE_CLIENT_ALIASES.items():
        if name in stem:
            return code
    for code in CLIENTS:
        if re.search(rf"(^|[^A-Z]){re.escape(code.upper())}([^A-Z]|$)", upper):
            return code
    return None


def default_ship() -> date:
    now = datetime.now()
    d = now.date() + timedelta(days=1 if now.hour < 15 else 2)
    if d.weekday() == 6:
        d += timedelta(days=1)
    return d


WD = "월화수목금토일"


def ship_label(d: date, mode: str | None = None) -> str:
    lab = WD[d.weekday()]
    return f"{lab}\n({mode})" if mode in ("배달", "택배", "화물", "내사") else lab


def _normalize_ingested_order(order: dict, path: Path, product_mode: str = "blind") -> dict:
    """Streamlit v72의 ingest 후처리와 동일한 핵심 보정을 적용한다."""
    product_mode = normalize_product_mode(product_mode)
    order["_product_mode"] = product_mode
    if product_mode == "blind":
        apply_holding_feature_gate(order)

    for it in order.get("items") or []:
        try:
            left_count = int(float(it.get("좌개수") or 0))
            right_count = int(float(it.get("우개수") or 0))
        except (TypeError, ValueError):
            left_count = right_count = 0
        if left_count > 0 and right_count == 0:
            it["손잡이방향"] = "좌"
        elif right_count > 0 and left_count == 0:
            it["손잡이방향"] = "우"

    if order.get("거래처") == "DD":
        for it in order.get("items") or []:
            for field in ("가로", "세로"):
                value = it.get(field)
                if isinstance(value, (int, float)) and value >= 500:
                    converted = value / 10
                    it[field] = int(converted) if converted.is_integer() else converted
            code = str(it.get("품목코드") or "").strip().upper()
            it["품목코드"] = re.sub(r"^(?:NA|GR)", "", code) or None

    if order.get("거래처") == "아지트":
        raw = " ".join(str(x or "") for x in (
            order.get("전체원문"), order.get("전체기재사항")))
        order["_az_place"] = "금빛커텐" if "에어캡+포장" in raw else "시온가공소"
        common = str(order.get("전체기재사항") or "")
        common = re.sub(r"(?:^|/)☆?(?:금빛커텐|시온가공소)(?=/|$)", "", common)
        order["전체기재사항"] = common.strip("/") or None
        order.setdefault("배송", {})["방식"] = "배달"

    expand_same_size_directions(order)
    synchronize_order_for_outputs(order)
    order["_file"] = path.name
    order["_source_path"] = str(path)
    return order




def _holding_operation_from_raw(item: dict) -> str:
    raw = " ".join(str(item.get(k) or "") for k in
                   ("홀딩방식", "수량", "원문", "기재사항", "색상원문"))
    direct = item.get("홀딩방식")
    if direct or looks_like_holding_operation(item.get("수량")):
        return normalize_holding_operation(direct or item.get("수량"))
    compact = re.sub(r"\s+", "", raw)
    # 긴 표현부터 확인한다.
    for pat, value in (
        (r"양개양자석", "양개 양자석"),
        (r"양자석", "양자석"),
        (r"양개", "양"),
        (r"편개", "편"),
        (r"(?:이등분|2등분)", "1/2"),
        (r"고정", "고정"),
    ):
        if re.search(pat, compact):
            return normalize_holding_operation(value)
    m = re.search(r"1/(10|[2-9])(?:양자석)?", compact)
    if m:
        token = m.group(0)
        return normalize_holding_operation(token)
    return "편"


def _normalize_holding_order(order: dict, path: Path) -> dict:
    """전용 홀딩도어 모드에서 H 본품과 홀딩 부속을 공식 품목으로 복원한다.

    한 문장에 본품과 '레일연결부속포함/라운드부속포함'이 함께 적힌 경우
    부속이 본품을 덮어쓰지 않도록 본품 행 + 부속 행으로 분리한다.
    """
    order["_product_mode"] = "holding"
    normalized = []
    for original in order.get("items") or []:
        item = original
        raw = " ".join(str(item.get(k) or "") for k in
                       ("색상원문", "품목코드", "원문", "기재사항", "홀딩부속"))
        accessory = holding_accessory_from_text(raw, item.get("홀딩부속색상") or "화이트")

        # 홀딩 전용 모드에서 비전이 B117처럼 블라인드 코드로 읽더라도
        # 숫자 코드를 H 품목장에 우선 대입한다.
        raw_code = str(item.get("품목코드") or "").strip().upper().replace(" ", "")
        holding_code_candidate = None
        mcode = re.fullmatch(r"[BH]?0*(\d{1,3}(?:-1)?)", raw_code)
        if mcode:
            num = mcode.group(1)
            if "-" in num:
                base, suffix = num.split("-", 1)
                holding_code_candidate = f"H{int(base):03d}-{suffix}"
            else:
                holding_code_candidate = f"H{int(num):03d}"

        # 부속 문구가 같이 있어도 본품 검색을 먼저 독립적으로 수행한다.
        product = (
            holding_product_from_text(holding_code_candidate)
            or holding_product_from_text(item.get("색상원문"))
            or holding_product_from_text(item.get("품목코드"))
            or holding_product_from_text(raw)
        )
        if accessory and product and str(product.get("품명") or "").startswith(("H레일연결부속", "H라운드부속")):
            # '레일연결부속 블랙'처럼 부속만 적힌 줄은 본품이 아니다.
            # 부속 품목명이 본품 검색에도 걸려 본품+부속으로 중복 생성되던 문제를 막는다.
            product = None

        def normalize_main(target: dict, hit):
            target["제품군"] = "홀딩도어"
            target["_product_group"] = "holding"
            target["_ledger_prefix"] = "H"
            target["예외품목"] = None
            target["연창"] = False
            target["손잡이방향"] = None
            target["손잡이길이"] = None
            target["타입"] = "C자"
            target["종류"] = "투코드"
            target.pop("_holding_accessory", None)
            if hit:
                target.pop("_holding_unmatched", None)
                target["_holding_product_name"] = hit["품명"]
                target["_holding_label"] = hit.get("장부표시")
                target["품목코드"] = hit.get("코드")
                # AI가 처음 읽은 품목코드 확신도는 공식 홀딩 품목장 매칭 전의 값이다.
                # 색상/겹수 또는 H코드가 공식 품목장 한 품목으로 확정된 뒤에는
                # 최종 품목코드에 과거의 낮은 확신도 경고를 남기지 않는다.
                conf = dict(target.get("확신도") or {})
                if conf:
                    target.setdefault("_원본확신도", dict(conf))
                conf["품목코드"] = 1.0
                target["확신도"] = conf
            else:
                target["_holding_unmatched"] = True
                target["_holding_raw"] = raw.strip()
                target["품목코드"] = None
                target["색상원문"] = "[홀딩 품목확인]"
            target["_holding_operation"] = _holding_operation_from_raw(target)
            # '레일연결부속포함'의 '레일' 글자를 본품 레일 종류로 오인하지 않는다.
            compact = re.sub(r"\s+", "", raw)
            rail_scan = compact.replace("레일연결부속포함", "").replace("레일연결부속", "")
            if "라운드" in rail_scan:
                target["_holding_rail"] = "라운드"
            elif "레일" in rail_scan:
                target["_holding_rail"] = "레일"
            else:
                target["_holding_rail"] = str(target.get("홀딩레일") or "").strip() or None
            target["_holding_upper_roller"] = bool(target.get("홀딩상하로라")) or "상하로라" in compact
            return target

        # 본품이 있으면 항상 본품을 먼저 남긴다.
        if product:
            normalized.append(normalize_main(item, product))
        elif not accessory:
            normalized.append(normalize_main(item, None))

        # 같은 문장에 부속이 있으면 본품을 대체하지 않고 별도 품목 행으로 추가한다.
        if accessory:
            acc = copy.deepcopy(original) if product else item
            acc["제품군"] = "홀딩도어"
            acc["_product_group"] = "holding"
            acc["_ledger_prefix"] = "H"
            acc["예외품목"] = None
            acc["연창"] = False
            acc["손잡이방향"] = None
            acc["손잡이길이"] = None
            acc["타입"] = "C자"
            acc["종류"] = "투코드"
            acc["_holding_accessory"] = True
            acc["_holding_product_name"] = accessory["품명"]
            acc["_holding_accessory_label"] = accessory.get("장부표시")
            acc["_holding_accessory_color"] = accessory.get("부속색상") or "화이트"
            acc["품목코드"] = accessory.get("코드")
            acc["색상원문"] = accessory.get("장부표시") or accessory.get("품명")
            acc_conf = dict(acc.get("확신도") or {})
            if acc_conf:
                acc.setdefault("_원본확신도", dict(acc_conf))
            acc_conf["품목코드"] = 1.0
            acc["확신도"] = acc_conf
            acc["가로"] = None
            acc["세로"] = None
            if product:
                # 본품 줄에 '레일연결부속포함'처럼 함께 적힌 부속은 1개로 본다.
                acc["창개수"] = 1
                acc["수량"] = 1
            else:
                # 부속만 적힌 줄은 발주서에 적힌 부속 개수를 그대로 유지한다.
                try:
                    acc_qty = int(float(str(acc.get("수량") or acc.get("창개수") or 1)
                                        .replace("X", "").replace("x", "")))
                except (TypeError, ValueError):
                    acc_qty = 1
                acc["창개수"] = 1
                acc["수량"] = max(1, acc_qty)
            acc["좌개수"] = 0
            acc["우개수"] = 0
            acc["_holding_operation"] = None
            acc["_holding_rail"] = None
            acc["_holding_upper_roller"] = False
            acc.pop("_holding_unmatched", None)
            normalized.append(acc)

    order["items"] = normalized
    return _normalize_ingested_order(order, path, product_mode="holding")

def ingest_file(path: str | Path, client_hint: str | None = None,
                product_mode: str = "blind") -> list[dict]:
    """이미지/PDF/거래처 Excel 파일 하나를 표준 주문 목록으로 읽는다."""
    product_mode = normalize_product_mode(product_mode)
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)

    suffix = path.suffix.lower()

    if product_mode == "roll_combo":
        M = get_master_for_mode("roll_combo")
        hint = client_hint if client_hint in ROLL_CLIENTS else None
        if hint is None:
            hint = roll_canonical_client(path.stem)
        if suffix in (".xlsx", ".xls"):
            # 안산)보노만 거래처 자체 Excel 주문서를 사용한다. 나머지 업체는
            # 카카오톡 이미지/PDF 주문으로 처리해 잘못된 범용 Excel 추정을 막는다.
            if hint not in (None, "안산)보노"):
                raise ValueError(f"{hint} 롤/콤비 주문은 카카오톡 이미지/PDF로 입력해 주세요.")
            got = parse_bono_roll(path, M)
        elif suffix == ".pdf":
            got = extract_roll_pdf(path, M, client_hint=hint)
        elif suffix in (".png", ".jpg", ".jpeg", ".webp", ".bmp"):
            got = [extract_roll_images([path], M, client_hint=hint)]
        else:
            raise ValueError(f"지원하지 않는 롤/콤비 파일 형식입니다: {suffix or path.name}")
        if not got or not any((o.get("items") or []) for o in got):
            raise ValueError("롤/콤비 발주서에서 유효한 품목을 읽지 못했습니다.")
        result = []
        for o in got:
            normalize_roll_order(o, M, hint)
            o["_file"] = path.name
            o["_source_path"] = str(path)
            # 롤/콤비는 연창을 프로그램이 자동 입력하지 않는다.
            for it in o.get("items") or []:
                it["연창"] = False
            synchronize_order_for_outputs(o)
            result.append(o)
        return result

    if product_mode == "holding":
        hint = client_hint or client_from_filename(path)
        if hint not in HOLDING_CLIENTS:
            hint = None
        if suffix in (".xlsx", ".xls"):
            got = parse_excel(path, hint)
        elif suffix == ".pdf":
            from extract import extract_pdf
            got = extract_pdf(path, client_hint=hint)
        elif suffix in (".png", ".jpg", ".jpeg", ".webp", ".bmp"):
            from extract import extract_order
            got = [extract_order([path], client_hint=hint)]
        else:
            raise ValueError(f"지원하지 않는 홀딩도어 파일 형식입니다: {suffix or path.name}")
        if not got or not any((o.get("items") or []) for o in got):
            raise ValueError("홀딩도어 발주서에서 유효한 품목을 읽지 못했습니다.")
        return [_normalize_holding_order(o, path) for o in got]

    hint = client_hint or client_from_filename(path)
    if hint not in BLIND_CLIENTS:
        hint = None

    if suffix in (".xlsx", ".xls"):
        got = parse_excel(path, hint)
    elif suffix == ".pdf":
        from extract import extract_pdf
        got = extract_pdf(path, client_hint=hint)
    elif suffix in (".png", ".jpg", ".jpeg", ".webp", ".bmp"):
        from extract import extract_order, prepare_sp_image
        # 스페이스 작업일지 사진은 PDF와 같은 보정을 거친다(눕힌 사진 세우기·대비).
        source = prepare_sp_image(path) if hint == "SP" else path
        got = [extract_order([source], client_hint=hint)]
    else:
        raise ValueError(f"지원하지 않는 파일 형식입니다: {suffix or path.name}")

    if not got or not any((o.get("items") or []) for o in got):
        raise ValueError(
            "발주서에서 유효한 품목을 읽지 못했습니다. 거래처 선택과 발주서 형식을 확인해 주세요."
        )

    return [_normalize_ingested_order(o, path, product_mode=product_mode) for o in got]


def ledger_orders_for_precheck(orders: list[dict], source: str | Path, product_mode: str,
                               ship: date) -> list[dict]:
    """직접 작성한 기존 장부(read_ledger)에서 읽은 주문을 사전점검 주문으로 준비한다.

    발주서 분석 결과와 같은 정규화를 거치므로 사전점검 확인·수정 → [파일 생성]에서
    장부·작업지시서·경영박사 EDI가 같은 데이터로 만들어진다.
    """
    product_mode = normalize_product_mode(product_mode)
    path = Path(source)
    ready: list[dict] = []
    for order in orders:
        if not order.get("items"):
            continue
        if product_mode == "holding":
            for item in order["items"]:
                _restore_holding_fields_from_ledger(item)
            _restore_holding_common_note(order)
            order = _normalize_holding_order(order, path)
            for item in order.get("items") or []:
                if item.get("_holding_accessory"):
                    m = re.fullmatch(r"[Xx×]?\s*(\d+)", str(item.get("_ledger_qty_text") or "").strip())
                    if m:
                        item["수량"] = int(m.group(1))
        else:
            order = _normalize_ingested_order(order, path, product_mode=product_mode)
        order["_ship_date"] = ship.isoformat()
        order["_from_ledger"] = True
        ready.append(order)
    return ready


def _restore_holding_common_note(order: dict) -> None:
    """홀딩 장부의 기재사항1(주문 공통값: 고객명 등)과 기재사항2(창별 메모)를 되돌린다.

    여러 창의 기재사항1이 모두 같고 모든 창에 기재사항2가 있으면(공통/개인 두 칸 구조)
    공통값은 고객명으로, 기재사항2는 창별 기재사항으로 둔다. 휴안처럼 한 칸만 쓰는 장부는 건드리지 않는다.
    """
    products = [it for it in order.get("items") or [] if not it.get("예외품목")]
    two_column = len(products) >= 2 and all(str(it.get("설치장소") or "").strip() for it in products)
    if two_column:
        lists = [[p.strip() for p in str(it.get("기재사항") or "").split("/") if p.strip()] for it in products]
        shared = [p for p in lists[0] if all(p in parts for parts in lists[1:])]
        if shared and not order.get("고객명"):
            order["고객명"] = "/".join(shared)
        customer_parts = [p for p in str(order.get("고객명") or "").split("/") if p.strip()]
        for it, parts in zip(products, lists):
            rest = [p for p in parts if p not in shared and p not in customer_parts]
            it["기재사항"] = "/".join(rest + [str(it.get("설치장소")).strip()])
            it["설치장소"] = None
    # 창별 메모 한 칸(`현장B/베란다`)에서 방 이름은 설치장소로 되돌린다.
    for it in order.get("items") or []:
        if it.get("설치장소"):
            continue
        notes, rooms = split_room_parts(it.get("기재사항"))
        if rooms:
            it["기재사항"] = "/".join(notes) or None
            it["설치장소"] = "/".join(rooms)


def _restore_holding_fields_from_ledger(item: dict) -> None:
    """홀딩 장부는 G열=작동방식, H열=개수, I열=레일이다. 블라인드 형식으로 읽힌 값을 홀딩 필드로 되돌린다."""
    color = str(item.get("_ledger_color_text") or "").strip()
    if color and not item.get("색상원문"):
        item["색상원문"] = color
    qty = str(item.get("_ledger_qty_text") or "").strip()
    count = str(item.get("_ledger_direction_text") or "").strip()
    rail = str(item.get("_ledger_length_text") or "").strip()
    if qty and looks_like_holding_operation(qty):
        item["홀딩방식"] = qty
        item["수량"] = None
    if re.fullmatch(r"\d+", count):
        item["창개수"] = int(count)
    if "레일" in rail or "라운드" in rail:
        item["홀딩레일"] = rail


def _order_ship_date(order: dict, fallback: date) -> date:
    """주문별 사전점검 출고일을 우선하고, 없을 때만 좌측 선택일을 사용한다."""
    raw = order.get("_ship_date")
    if isinstance(raw, date):
        return raw
    if raw:
        try:
            return date.fromisoformat(str(raw)[:10])
        except ValueError:
            pass
    return fallback


_ROW_ID_SEQ = itertools.count(1)


def ensure_row_ids(orders: list[dict]) -> None:
    """주문/품목마다 고유 ID를 붙인다.

    사전점검의 수동 병합·해제는 이 ID로 행을 기억한다. 그래서 행을 추가·이동·삭제해
    순서가 바뀌어도 같은 행을 다시 찾는다. 복사로 ID가 겹치면 뒤쪽 행에 새 ID를 준다.
    장부·작업지시서·EDI 출력 값에는 영향이 없다.
    """
    seen_orders = set()
    for order in orders:
        if not order.get("_order_uid") or order["_order_uid"] in seen_orders:
            order["_order_uid"] = f"#o{next(_ROW_ID_SEQ)}"
        seen_orders.add(order["_order_uid"])
        seen_items = set()
        for item in order.get("items") or []:
            if not item.get("_uid") or item["_uid"] in seen_items:
                item["_uid"] = f"#i{next(_ROW_ID_SEQ)}"
            seen_items.add(item["_uid"])


def build_output_rows(orders: list[dict], ship: date, M: Master | None = None,
                      sort_di: bool = False) -> list[dict]:
    """현재 주문 데이터에서 실제 장부 출력 행을 재생성한다."""
    # 창 단위 분할(synchronize)을 먼저 끝낸 뒤 ID를 붙여야 분할된 행도 서로 다른 ID를 갖는다.
    for order in orders:
        synchronize_order_for_outputs(order)
    ensure_row_ids(orders)
    rows: list[dict] = []
    sp_notices: list[str] = []
    for order in orders:
        if order.get("거래처") != "SP":
            continue
        for part in str(order.get("전체기재사항") or "").split("/"):
            part = part.strip()
            if "공지" in part and part not in sp_notices:
                sp_notices.append(part)
    sp_notice_written = False

    for oi, order in enumerate(orders):
        client = order.get("거래처")
        actual = (order.get("배송") or {}).get("방식")
        is_roll = order.get("_product_mode") == "roll_combo"
        if is_roll:
            default = (ROLL_CLIENT_INFO.get(client) or {}).get("delivery")
            manual_mode = bool(order.get("_delivery_mode_manual"))
            mode = None if manual_mode and not actual else (actual or default)
            if mode not in ("배달", "택배", "화물", "내사"):
                mode = None
            # 거래처 기본 배송지가 곧 택배 목적지인 경우에는 장부 출고일에
            # `(택배)`를 적지 않는다. 별도 주소/화물지점이 있는 택배만 표시한다.
            delivery = order.get("배송") or {}
            uses_registered_address = (
                mode == "택배" and default == "택배"
                and not str(delivery.get("주소") or "").strip()
                and not str(delivery.get("화물지점") or "").strip()
            )
            if uses_registered_address:
                mode = None
        else:
            default = CLIENT_INFO.get(client, (None, None, None))[2]
            mode = actual if client == "RT" or (
                actual in ("택배", "화물", "배달", "내사") and actual != default
            ) else None
        # 휴안/보노/미래가공은 별도 배송지로 보내더라도 장부 출고일 옆에
        # (택배)/(화물)을 적지 않는 업체 고정 규칙이다. 배송정보 자체는 유지한다.
        if client in NO_DELIVERY_LABEL_CLIENTS:
            mode = None
        order_ship = _order_ship_date(order, ship)
        # 출고일은 달력에서 날짜로 고르고, 장부·작업지시서·사전점검에는 요일로 적는다.
        # (블라인드·홀딩·롤콤비 공통. 경영박사 EDI는 실제 날짜 _ship_date 를 쓴다.)
        got = to_rows(order, ship_label(order_ship, mode), M, sort_di=sort_di)

        if client == "SP":
            first_product = None
            for row in got:
                if row.get("_특수") is not None:
                    continue
                if first_product is None:
                    first_product = row
                parts = [x.strip() for x in str(row.get("기재사항") or "").split("/")
                         if x.strip() and "공지" not in x]
                row["기재사항"] = "/".join(parts) or None
            if first_product is not None and sp_notices and not sp_notice_written:
                old = first_product.get("기재사항")
                first_product["기재사항"] = "/".join(sp_notices + ([old] if old else []))
                sp_notice_written = True

        visible_i = 0
        special_counts = {}
        for row in got:
            row["_oi"] = oi
            if row.get("_특수") is None:
                source_i = row.get("_source_item_index")
                row["_ii"] = source_i if isinstance(source_i, int) else visible_i
                row["_출고일날짜"] = order_ship.isoformat() if visible_i == 0 else None
                order_items = order.get("items") or []
                uid = (order_items[row["_ii"]].get("_uid")
                       if 0 <= row["_ii"] < len(order_items) else None)
                row["_row_key"] = ("product", order.get("_order_uid"), uid or f"ii{row['_ii']}")
                visible_i += 1
            else:
                kind = _infer_special_kind(order, row)
                ordinal = special_counts.get(kind, 0)
                special_counts[kind] = ordinal + 1
                row["_row_key"] = ("special", order.get("_order_uid"), kind, ordinal)
            row["_output_index"] = len(rows)
            rows.append(row)
    return rows


def merge_sp_page_orders(orders: list[dict]) -> list[dict]:
    """스페이스 작업일지 여러 장을 주문번호 기준으로 합친다(2026-09-18).

    PDF는 쪽마다 판독한 뒤 합치지만, 사진은 파일 하나가 주문 하나로 들어온다.
    한 주문이 두 장에 걸쳐 있으면 주문번호가 같으므로 한 건으로 합친다.
    번호를 못 읽은 장은 합치지 않고 그대로 둔다.
    """
    from extract import _merge_sp_orders, _sp_order_key
    mergeable = [o for o in orders
                 if o.get("거래처") == "SP" and _sp_order_key(o.get("주문번호"))]
    if len(mergeable) < 2:
        return orders
    merged_by_key = {_sp_order_key(m.get("주문번호")): m
                     for m in _merge_sp_orders([dict(o) for o in mergeable])}
    mergeable_ids = {id(o) for o in mergeable}
    out, seen = [], set()
    for order in orders:
        if id(order) not in mergeable_ids:
            out.append(order)
            continue
        key = _sp_order_key(order.get("주문번호"))
        if key in seen:
            continue
        seen.add(key)
        out.append(merged_by_key.get(key, order))
    return out


def count_windows(orders: list[dict]) -> int:
    """사전점검 마지막 행의 총 창수(블라인드·홀딩·롤콤비 공통, 부속·예외품목 제외)."""
    total = 0
    for order in orders:
        for it in order.get("items") or []:
            if it.get("예외품목") or it.get("_holding_accessory"):
                continue
            try:
                total += _blind_counts(it)[2]
            except Exception:
                total += 1
    return total


def _infer_special_kind(order: dict, row: dict) -> str:
    marker = str(row.get("_특수") if row.get("_특수") is not None else "")
    text = str(row.get("문구") or "").strip()
    delivery = order.get("배송") or {}

    if marker == "믹스":
        return "mix_part"
    if marker == "부속":
        if text and text == str(order.get("_추가부속") or "").strip():
            return "manual_accessory"
        return "structured_accessory"
    if marker == "#" and re.sub(r"\s+", "", text) == "포장비용":
        return "packing"
    if marker == "☆":
        if text == str(delivery.get("주소") or "").strip():
            return "address"
        if order.get("거래처") == "아지트" and text == str(order.get("_az_place") or "").strip():
            return "az_place"
        return "address"
    if marker == "":
        if re.match(r"^전달\s*[:：]", text):
            return "delivery_notice"
        if re.match(r"^발신\s*[:：]", text):
            return "sender"
        return "recipient_contact"
    return "readonly"


def build_preview_rows(orders: list[dict], ship: date, M: Master | None = None,
                       sort_di: bool = False) -> tuple[list[dict], list[dict]]:
    """PySide6 사전점검용 행과 실제 출력행을 함께 반환한다.

    실제 장부 특수행은 그대로 보여주며, 택배/화물 주소가 비어 있을 때만
    빈 ☆ 주소행을 UI 전용으로 추가해서 클릭 없이 즉시 수정할 수 있게 한다.
    """
    output_rows = build_output_rows(orders, ship, M=M, sort_di=sort_di)
    preview: list[dict] = []
    seen_order_has_address_row: dict[int, bool] = {}

    for row in output_rows:
        ui = dict(row)
        oi = int(ui.get("_oi", -1))
        order = orders[oi] if 0 <= oi < len(orders) else {}
        if ui.get("_특수") is not None:
            ui["_detail_kind"] = _infer_special_kind(order, ui)
            if ui["_detail_kind"] == "address":
                seen_order_has_address_row[oi] = True
                ui["_delivery_mode"] = str((order.get("배송") or {}).get("방식") or "").strip()
        else:
            ui["_detail_kind"] = "product"
        ui["_preview_id"] = ui.get("_row_key") or ("out", ui.get("_output_index"))
        preview.append(ui)

    # 주소 누락 경고가 난 주문도 클릭 없이 바로 수정할 수 있게 빈 주소행을 노출한다.
    insert_blocks: list[tuple[int, list[dict]]] = []
    for oi, order in enumerate(orders):
        delivery = order.get("배송") or {}
        default_mode = _default_delivery_for_order(order)
        mode = str(delivery.get("방식") or default_mode or "")
        needs_address = ("택배" in mode or "화물" in mode)
        # 별도 배송 목적지는 거래처 기본 등록지와 다른 곳으로 보내는 경우에만 필요하다.
        # 기본 택배/화물 방식과 동일하고 별도 주소가 없으면 주소1/등록 화물지점을 사용한다.
        if _uses_registered_destination(order, mode):
            continue
        if not needs_address or seen_order_has_address_row.get(oi):
            continue
        # 해당 주문의 마지막 preview 행 뒤에 삽입한다.
        positions = [i for i, r in enumerate(preview) if r.get("_oi") == oi]
        pos = (positions[-1] + 1) if positions else len(preview)
        address_text = (_freight_ledger_text(delivery, " " if order.get("거래처") in BONO_LIKE_CLIENTS else " / ")
                        if "화물" in mode
                        else str(delivery.get("주소") or ""))
        missing_rows = [{
            "_oi": oi, "_ii": None, "_특수": "☆",
            "문구": address_text,
            "_detail_kind": "address", "_output_index": None,
            "_preview_id": ("ui", oi, "address"), "_ui_only": True,
            "_delivery_mode": str(delivery.get("방식") or "").strip(),
        }]
        who = " ".join(x for x in (delivery.get("수령인"), delivery.get("연락처")) if x)
        if who:
            missing_rows.append({
                "_oi": oi, "_ii": None, "_특수": "", "문구": who,
                "_detail_kind": "recipient_contact", "_output_index": None,
                "_preview_id": ("ui", oi, "recipient_contact"), "_ui_only": True,
                "선불": "선불" if str(delivery.get("선불착불") or "").strip() == "선불" else None,
            })
        notice = clean_delivery_notice(delivery.get("전달사항")) or ""
        if notice:
            missing_rows.append({
                "_oi": oi, "_ii": None, "_특수": "", "문구": f"전달 : {notice}",
                "_detail_kind": "delivery_notice", "_output_index": None,
                "_preview_id": ("ui", oi, "delivery_notice"), "_ui_only": True,
            })
        if str(delivery.get("발신") or "").strip():
            missing_rows.append({
                "_oi": oi, "_ii": None, "_특수": "", "문구": f"발신 : {delivery['발신']}",
                "_detail_kind": "sender", "_output_index": None,
                "_preview_id": ("ui", oi, "sender"), "_ui_only": True,
            })
        insert_blocks.append((pos, missing_rows))

    # 뒤쪽 주문부터 블록 단위로 넣어 기존 주문 행 위치가 밀리지 않게 한다.
    for pos, block in sorted(insert_blocks, key=lambda x: x[0], reverse=True):
        preview[pos:pos] = block

    if orders:
        # 마지막 행: 총 창수 (화면 전용, 장부/EDI에는 나가지 않음)
        total = count_windows(orders)
        preview.append({
            "_oi": None, "_ii": None, "_특수": "", "문구": f"{total}창", "_총창수": total,
            "_detail_kind": "total", "_output_index": None,
            "_preview_id": ("ui", "total"), "_ui_only": True,
        })
    return preview, output_rows


def row_display_values(row: dict, row_number: int) -> dict[str, Any]:
    """preview row -> 화면에 보일 셀 값."""
    out = {c: "" for c in LEDGER_COLUMNS}
    if row.get("_특수") is None:
        vendor = row.get("거래처") or ""
        mark = row.get("내부표시") or ""
        out.update({
            "상호": f"{vendor} {mark}".strip(),
            "색상": row.get("색상") or "",
            "가로": row.get("가로") or "",
            "X": "X",
            "세로": row.get("세로") or "",
            "수량": row.get("수량") or "",
            "방향": row.get("모형1") or "",
            "길이": row.get("모형2") or "",
            "특이": row.get("특이") or "",
            "기재사항1": row.get("기재사항") or "",
            "기재사항2": row.get("기재사항2") or "",
            "출고일": row.get("출고일") or "",
        })
        return out

    marker = row.get("_특수")
    if row.get("_detail_kind") == "total":
        out["색상"] = "총 창수"
        out["가로"] = row.get("문구") or ""
        return out
    if marker == "믹스":
        length = row.get("길이")
        out["색상"] = row.get("조합") or ""
        out["가로"] = row.get("코드") or ""
        out["수량"] = f") {length:g}cm" if isinstance(length, (int, float)) else ")"
        return out
    if marker == "부속":
        out["색상"] = row.get("문구") or ""
        out["수량"] = row.get("수량") or ""
    else:
        out["색상"] = marker or ""
        out["가로"] = row.get("문구") or ""
        if row.get("_detail_kind") == "address":
            out["특이"] = row.get("_delivery_mode") or ""
        if marker == "#" and re.sub(r"\s+", "", str(row.get("문구") or "")) == "포장비용":
            out["수량"] = row.get("수량") or ""
        if row.get("선불"):
            out["기재사항2"] = row.get("선불")
    return out


def editable_columns_for_row(row: dict) -> set[str]:
    if row.get("_특수") is None:
        # 상호/출고일은 자유입력이 아니라 클릭 선택창으로만 변경한다.
        return set(PRODUCT_EDIT_FIELD) - {"X", "출고일", "상호"}
    kind = row.get("_detail_kind")
    if kind == "address":
        # 주소행의 특이 칸에서 택배/화물/배달/내사를 직접 수정할 수 있다.
        return {"가로", "특이"}
    if kind in {"az_place", "recipient_contact", "delivery_notice", "sender"}:
        return {"가로"}
    if kind == "mix_part":
        return {"가로", "수량"}   # 섞인 코드 / 길이
    if kind == "manual_accessory":
        return {"색상"}
    if kind == "packing":
        return {"수량"}
    return set()


def _coerce_packing_qty(value: Any) -> int | None:
    text = str(value or "").strip()
    if not text:
        return None
    m = re.search(r"\d+", text)
    if not m:
        return None
    return max(0, int(m.group(0)))


def _edit_mix_part(order: dict, row: dict, column: str, text: str) -> None:
    """MIX 구성행 수정: 가로 칸=섞인 코드, 수량 칸=길이. 같은 MIX 묶음의 모든 창에 반영하고 자동 계산을 멈춘다."""
    items = order.get("items") or []
    idx, k = row.get("_mix_item_index"), row.get("_mix_part_index")
    if not isinstance(idx, int) or not (0 <= idx < len(items)) or not isinstance(k, int):
        return
    target = items[idx]
    parts = [dict(p) for p in target.get("_mix_parts") or []]
    if not (0 <= k < len(parts)):
        return
    key = _mix_group_key(target)
    group = [idx]
    j = idx - 1
    while j >= 0 and _mix_group_key(items[j]) == key:
        group.insert(0, j)
        j -= 1
    if column == "가로":
        m = re.search(r"\d{3,4}(?:FP|P)?", text, re.I)
        if not m:
            return
        parts[k]["코드"] = m.group(0).upper()
    else:
        m = re.search(r"\d+(?:\.\d+)?", text)
        parts[k]["길이"] = float(m.group(0)) if m else None
    combo = "+".join(p["코드"] for p in parts)
    for gi in group:
        items[gi].update({"_mix_parts": [dict(p) for p in parts], "_mix_parts_manual": True,
                          "_mix_combo": combo, "_mix_codes": combo})
        items[gi].pop("_mix_source", None)


PHONE_RX = re.compile(r"(?:01[016789]|0\d{1,2})[-.\s]?\d{3,4}[-.\s]?\d{4}")


def apply_preview_cell_edit(orders: list[dict], row: dict, column: str,
                            value: Any, color_prefix: str = "B") -> None:
    """화면 셀 수정값을 표준 주문 데이터에 반영한다."""
    oi = row.get("_oi")
    if not isinstance(oi, int) or not (0 <= oi < len(orders)):
        return
    order = orders[oi]

    if row.get("_특수") is None:
        ii = row.get("_ii")
        field = PRODUCT_EDIT_FIELD.get(column)
        if isinstance(ii, int) and field and 0 <= ii < len(order.get("items") or []):
            # 병합되어 보이는 공통 기재사항은 실제 주문 데이터도 묶음 전체에 동일 적용한다.
            if order.get("거래처") == "JL" and column == "기재사항1":
                text = str(value or "").strip()
                for target in order.get("items") or []:
                    target["_manual_note1"] = text or None
                    target["기재사항"] = text or None
            elif order.get("거래처") == "JO" and column == "기재사항1":
                # JO의 기재사항1은 주문번호 그 자체다. 한 셀 수정 시 주문번호를 수정한다.
                text = str(value or "").strip()
                order["주문번호"] = text or None
                for target in order.get("items") or []:
                    target.pop("_manual_note1", None)
            elif order.get("_product_mode") == "roll_combo":
                target = order["items"][ii]
                if column == "색상":
                    target["_ledger_text_manual"] = str(value or "").strip() or None
                elif column == "가로":
                    try: target["가로"] = float(value)
                    except (TypeError, ValueError): target["가로"] = value
                elif column == "세로":
                    try: target["세로"] = float(value)
                    except (TypeError, ValueError): target["세로"] = value
                elif column == "수량":
                    target["수량"] = value
                elif column == "방향":
                    target["손잡이방향"] = str(value or "").strip() or None
                elif column == "길이":
                    target["_manual_handle_length"] = value
                    target["손잡이길이"] = value
                elif column == "특이":
                    target["_수동특이"] = str(value or "").strip() or None
                elif column == "기재사항1":
                    target["_manual_note1"] = str(value or "").strip() or None
                elif column == "기재사항2":
                    target["_manual_note2"] = str(value or "").strip() or None
            elif order.get("_product_mode") == "holding":
                target = order["items"][ii]
                if column == "색상":
                    text = str(value or "").strip()
                    target["색상원문"] = text or None
                    target["품목코드"] = text or None
                elif column == "가로":
                    try: target["가로"] = float(value)
                    except (TypeError, ValueError): target["가로"] = value
                elif column == "세로":
                    try: target["세로"] = float(value)
                    except (TypeError, ValueError): target["세로"] = value
                elif column == "수량":
                    target["수량"] = value
                elif column == "방향":
                    target["홀딩방식"] = str(value or "").strip() or None
                    target["_holding_operation"] = normalize_holding_operation(value)
                elif column == "길이":
                    rail = str(value or "").strip() or None
                    target["홀딩레일"] = rail
                    target["_holding_rail"] = rail
                elif column == "특이":
                    target["_수동특이"] = str(value or "").strip() or None
                elif column == "기재사항1":
                    target["_manual_note1"] = str(value or "").strip() or None
                    target["기재사항"] = target["_manual_note1"]
                elif column == "기재사항2":
                    target["_manual_note2"] = str(value or "").strip() or None
                # 품명/코드 수정 시 공식 H 품목 메타를 즉시 다시 매칭한다.
                if column == "색상":
                    # _normalize_holding_order는 order를 제자리에서 갱신하고 같은 dict를 반환한다.
                    # 반환값을 같은 객체에 clear/update하면 주문 전체가 지워지므로 절대 clear하지 않는다.
                    _normalize_holding_order(
                        order, Path(order.get("_source_path") or order.get("_file") or "manual")
                    )
            else:
                apply_edit(order, ii, field, value, color_prefix)
            synchronize_order_for_outputs(order)
        return

    kind = row.get("_detail_kind")
    delivery = order.setdefault("배송", {})
    text = str(value or "").strip()

    if kind == "address" and column == "가로":
        mode = str(delivery.get("방식") or "").strip()
        if "화물" in mode:
            branch, receiver, phone = parse_freight_ledger_text(text)
            delivery["화물지점"] = branch
            delivery["주소"] = branch
            delivery["수령인"] = receiver
            delivery["연락처"] = phone
            if order.get("거래처") in BONO_STYLE_CLIENTS:
                order["_bono_branch"] = branch
                compact = re.sub(r"\s+", "", str(receiver or ""))
                self_names = (("보노",) if order.get("거래처") == "보노"
                              else ("미래가공", "미래", "보노"))
                order["_bono_receiver_is_bono"] = any(name in compact for name in self_names)
        else:
            delivery["주소"] = text or None
    elif kind == "address" and column == "특이":
        raw = re.sub(r"\s+", "", text)
        normalized = None
        for name in ("택배", "화물", "배달", "내사"):
            if name in raw:
                normalized = name
                break
        if normalized:
            delivery["방식"] = normalized
    elif kind == "mix_part" and column in {"가로", "수량"}:
        _edit_mix_part(order, row, column, text)
    elif kind == "az_place" and column == "가로":
        order["_az_place"] = text or None
    elif kind == "manual_accessory" and column == "색상":
        order["_추가부속"] = text or None
    elif kind == "recipient_contact" and column == "가로":
        m = PHONE_RX.search(text)
        if m:
            delivery["연락처"] = m.group(0).strip() or None
            name = (text[:m.start()] + " " + text[m.end():]).strip(" /,·")
            delivery["수령인"] = name or None
        else:
            # 숫자 위주의 문자열이면 연락처, 그 외는 수령인으로 본다.
            compact = re.sub(r"[^0-9]", "", text)
            if len(compact) >= 9:
                delivery["연락처"] = text or None
                delivery["수령인"] = None
            else:
                delivery["수령인"] = text or None
                delivery["연락처"] = None
    elif kind == "delivery_notice" and column == "가로":
        text = re.sub(r"^전달\s*[:：]\s*", "", text)
        delivery["전달사항"] = clean_delivery_notice(text)
    elif kind == "sender" and column == "가로":
        text = re.sub(r"^발신\s*[:：]\s*", "", text)
        delivery["발신"] = text or None
    elif kind == "packing" and column == "수량":
        qty = _coerce_packing_qty(value)
        if qty is None or qty <= 0:
            order["_manual_packing_enabled"] = False
            order["_manual_packing_qty"] = None
        else:
            order["_manual_packing_enabled"] = True
            order["_manual_packing_qty"] = qty
    synchronize_order_for_outputs(order)


SEVERITY_RANK = {"red": 4, "review": 3, "yellow": 2, "info": 1, "ok": 0}
STATUS_LABEL = {"red": "오류", "review": "중복·변경", "yellow": "확인", "info": "참고", "ok": "정상"}


def preview_statuses(preview_rows: list[dict], orders: list[dict], M: Master) -> list[tuple[str, str, str]]:
    """각 preview 행의 (등급, 상태표시, tooltip) 반환."""
    issues_by_order: dict[int, list[tuple]] = {}
    for oi, order in enumerate(orders):
        try:
            if order.get("_product_mode") == "roll_combo" or getattr(M, "is_roll_combo_master", False):
                issues_by_order[oi] = list(validate_roll_order(order, M))
            else:
                issues_by_order[oi] = list(validate(order, M))
        except Exception as exc:
            issues_by_order[oi] = [("red", "검증실패", f"검증 중 오류: {exc}", None)]

    result = []
    for row in preview_rows:
        oi = row.get("_oi")
        issues = issues_by_order.get(oi, []) if isinstance(oi, int) else []
        if row.get("_특수") is None and isinstance(row.get("_ii"), int):
            item_no = row["_ii"] + 1
            applicable = [x for x in issues if x[3] is None or x[3] == item_no]
        else:
            # 특수행에는 주문 전체 이슈만 표시. 주소누락처럼 row=None 검증이 여기 걸린다.
            applicable = [x for x in issues if x[3] is None]
        if not applicable:
            result.append(("ok", "정상", "원본 보기"))
            continue
        worst = max((x[0] for x in applicable), key=lambda lv: SEVERITY_RANK.get(lv, 0))
        tooltip = "\n".join(f"[{code}] {msg}" for _, code, msg, _ in applicable)
        result.append((worst, STATUS_LABEL.get(worst, worst), tooltip))
    return result


def generate_outputs(orders: list[dict], ship: date, M: Master, out_dir: str | Path,
                     tag: str | None = None) -> tuple[dict[str, Path], list[dict]]:
    """현재 확정 주문으로 장부/작업지시서/EDI를 생성한다.

    제품군별 작업 화면의 데이터가 서로 섞여 출력되지 않도록 한 배치에는
    하나의 product_mode만 허용한다.
    """
    modes = {normalize_product_mode(o.get("_product_mode") or "blind") for o in orders}
    if len(modes) > 1:
        raise ValueError(f"서로 다른 제품군 주문은 한 번에 출력할 수 없습니다: {sorted(modes)}")
    ship_iso = ship.isoformat()
    for order in orders:
        # 사전점검에서 주문별로 바꾼 날짜는 보존한다.
        # 아직 날짜가 없는 주문에만 좌측 선택일을 기본값으로 넣는다.
        if not order.get("_ship_date"):
            order["_ship_date"] = ship_iso
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    tag = tag or datetime.now().strftime("%m%d_%H%M%S")
    first = str((orders[0].get("거래처") if orders else None) or "미확인")
    first = re.sub(r"[\\/:*?\"<>|\s]+", "", first) or "미확인"
    stem = f"{tag}-{first}"

    output_rows = build_output_rows(orders, ship, M=M, sort_di=True)
    ledger = out_dir / f"장부_{stem}.xlsx"
    worksheet = out_dir / f"작업지시서_{stem}.xlsx"
    erp = out_dir / f"경영박사EDI_{stem}.xls"

    build_ledger(output_rows, ledger, header_date=date.today())
    build_worksheet(output_rows, worksheet, header_date=date.today())
    for order in orders:
        synchronize_order_for_outputs(order)
    # 각 주문의 확정 _ship_date를 사용하도록 전역 날짜 인자를 넘기지 않는다.
    build_erp(orders, M, erp, None)
    return {"장부": ledger, "작업지시서": worksheet, "경영박사": erp}, output_rows


def _row_location_map(output_rows: list[dict]) -> dict[int, tuple[str, int]]:
    """build_ledger와 같은 페이지 분할 규칙으로 global output index -> (sheet, excel row)."""
    mapping: dict[int, tuple[str, int]] = {}
    di_rows = [r for r in output_rows if r.get("_거래처코드") == "DI"]
    general_rows = [r for r in output_rows if r.get("_거래처코드") != "DI"]

    for group_name, group_key in DI_SHEET_GROUPS:
        group_rows = [r for r in di_rows if r.get("_DI유형") == group_key]
        pages = _di_paginate(group_rows) if group_rows else []
        for page_no, chunk in enumerate(pages, 1):
            title = group_name if page_no == 1 else f"{group_name} ({page_no})"
            for local_i, row in enumerate(chunk):
                idx = row.get("_output_index")
                if isinstance(idx, int):
                    mapping[idx] = (title, 3 + local_i)

    known = {key for _, key in DI_SHEET_GROUPS}
    unknown_di = [r for r in di_rows if r.get("_DI유형") not in known]
    for local_i, row in enumerate(unknown_di):
        idx = row.get("_output_index")
        if isinstance(idx, int):
            mapping[idx] = ("DI 기타", 3 + local_i)

    pages = _paginate_order_rows(general_rows) if general_rows else []
    for page_no, chunk in enumerate(pages, 1):
        title = "장부" if page_no == 1 else f"장부 ({page_no})"
        for local_i, row in enumerate(chunk):
            idx = row.get("_output_index")
            if isinstance(idx, int):
                mapping[idx] = (title, 3 + local_i)
    return mapping


def resolve_stable_merge_specs(output_rows: list[dict], specs: Iterable[tuple]) -> tuple[list[tuple[int, int, int, int]], list[str]]:
    """row_key 기반 수동 병합 스펙을 현재 output index 스펙으로 변환한다.

    stable spec = (tuple(row_keys), start_col, end_col)
    주문 행이 앞쪽에 추가/삭제되어 전역 인덱스가 밀려도 같은 행을 다시 찾는다.
    """
    key_to_index = {r.get("_row_key"): i for i, r in enumerate(output_rows) if r.get("_row_key") is not None}
    resolved = []
    notes = []
    for spec in specs:
        if len(spec) != 3:
            notes.append(f"잘못된 병합 스펙: {spec}")
            continue
        keys, c1, c2 = spec
        indices = [key_to_index.get(tuple(k) if isinstance(k, list) else k) for k in keys]
        if not indices or any(i is None for i in indices):
            notes.append(f"수정된 행 구조에서 병합 대상을 찾지 못했습니다: {spec}")
            continue
        if indices != list(range(min(indices), max(indices) + 1)):
            notes.append(f"병합 대상 행이 더 이상 연속하지 않습니다: {spec}")
            continue
        resolved.append((min(indices), max(indices), int(c1), int(c2)))
    return resolved, notes


def _overlap(r1: int, r2: int, c1: int, c2: int,
             rr1: int, rr2: int, cc1: int, cc2: int) -> bool:
    """두 Excel 직사각형 범위가 한 셀이라도 겹치는지 판정한다."""
    return not (r2 < rr1 or rr2 < r1 or c2 < cc1 or cc2 < c1)


def apply_excel_merge_overrides(ledger_path: str | Path, output_rows: list[dict],
                                merges: Iterable[tuple[int, int, int, int]],
                                unmerges: Iterable[tuple[int, int, int, int]]) -> list[str]:
    """PySide6에서 수동 병합/해제한 범위를 최종 장부.xlsx에도 반영한다.

    spec = (output_start_index, output_end_index, excel_start_col, excel_end_col)
    페이지를 가로지르는 범위는 안전하게 건너뛴다.
    """
    ledger_path = Path(ledger_path)
    if not ledger_path.exists():
        return []

    # build_ledger()가 만든 색상 열은 openpyxl RichText를 사용한다.
    # 수동 병합이 하나도 없는데도 파일을 일반 load_workbook()으로 다시 저장하면
    # RichText가 평문으로 변환되어 `원코드` 초록 / P·FP 파랑 등의 부분색상이 사라진다.
    # 병합 변경이 없으면 파일을 아예 다시 열지 않고, 변경이 있을 때도 rich_text=True로
    # 읽어서 기존 부분 서식을 그대로 보존한다.
    merges = list(merges)
    unmerges = list(unmerges)
    if not merges and not unmerges:
        return []
    wb = openpyxl.load_workbook(ledger_path, rich_text=True)
    loc = _row_location_map(output_rows)
    notes: list[str] = []

    def resolve(spec):
        a, b, c1, c2 = spec
        row_ids = list(range(min(a, b), max(a, b) + 1))
        places = [loc.get(i) for i in row_ids]
        if not places or any(p is None for p in places):
            return None
        sheets = {p[0] for p in places if p}
        if len(sheets) != 1:
            return None
        sheet = places[0][0]
        excel_rows = [p[1] for p in places]
        if excel_rows != list(range(min(excel_rows), max(excel_rows) + 1)):
            return None
        return sheet, min(excel_rows), max(excel_rows), min(c1, c2), max(c1, c2)

    # 먼저 수동 병합 해제를 적용한다.
    for spec in unmerges:
        target = resolve(spec)
        if not target:
            notes.append(f"병합 해제 범위를 출력 장부에 연결하지 못했습니다: {spec}")
            continue
        sheet, r1, r2, c1, c2 = target
        ws = wb[sheet]
        for rng in list(ws.merged_cells.ranges):
            if _overlap(r1, r2, c1, c2,
                        rng.min_row, rng.max_row, rng.min_col, rng.max_col):
                ws.unmerge_cells(str(rng))

    # 수동 병합은 기존 자동 병합과 겹치면 해당 자동 병합을 먼저 풀고 적용한다.
    for spec in merges:
        target = resolve(spec)
        if not target:
            notes.append(f"병합 범위를 출력 장부에 연결하지 못했습니다: {spec}")
            continue
        sheet, r1, r2, c1, c2 = target
        ws = wb[sheet]
        for rng in list(ws.merged_cells.ranges):
            if _overlap(r1, r2, c1, c2,
                        rng.min_row, rng.max_row, rng.min_col, rng.max_col):
                ws.unmerge_cells(str(rng))
        if r1 != r2 or c1 != c2:
            ws.merge_cells(start_row=r1, end_row=r2, start_column=c1, end_column=c2)

    wb.save(ledger_path)
    return notes
