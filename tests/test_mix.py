"""비대일 블라인드 MIX 규칙 (2026-09-13 사용자 확정, 정답 장부: JL 믹스.xls + 4색 배합 예시)."""
from pathlib import Path

import openpyxl
import pytest
import xlrd

import _cases
import desktop_workflow as dw
import output
from master_registry import get_master_for_mode
from rules import parse_mix_parts, round_half, validate

M = get_master_for_mode("blind")


def _order(client, items, **extra):
    order = {"거래처": client, "주문번호": "2", "고객명": "배수현", "전체기재사항": "피스",
             "배송": {}, "items": items, **extra}
    order = dw._normalize_ingested_order(order, Path("t.png"), "blind")
    order["_ship_date"] = _cases.SHIP.isoformat()
    return order


def _mix_item(text, w, h, kind="투코드", type_="C자", **extra):
    return {"품목코드": None, "색상원문": text, "종류": kind, "타입": type_, "가로": w, "세로": h, **extra}


def _rows(orders):
    return dw.build_output_rows(orders, _cases.SHIP, M=M, sort_di=True)


def _parts(rows):
    return [(r["코드"], r["길이"], r.get("조합")) for r in rows if r.get("_특수") == "믹스"]


# ───────────────────────── 길이 계산 ─────────────────────────
def test_round_half_and_parse_rules():
    assert round_half(120.6) == 120.5 and round_half(13.4) == 13.5 and round_half(49.5) == 49.5
    assert parse_mix_parts("MIX 200(90%)+125(10%)", 120) == [{"코드": "200", "길이": 108.0},
                                                             {"코드": "125", "길이": 12.0}]
    assert parse_mix_parts("200(108cm)+125(12cm)", 120) == [{"코드": "200", "길이": 108.0},
                                                            {"코드": "125", "길이": 12.0}]
    assert parse_mix_parts("MIX 102+500", 55) == [{"코드": "102", "길이": 27.5},
                                                  {"코드": "500", "길이": 27.5}]
    assert parse_mix_parts("MIX 200(90%)+125(10%)", 134) == [{"코드": "200", "길이": 120.5},
                                                             {"코드": "125", "길이": 13.5}]
    assert parse_mix_parts("MIX 102", 100) is None


# ───────────────────────── 장부 ─────────────────────────
def test_ledger_rows_two_color_mix_like_jl_file():
    rows = _rows([_order("JL", [
        _mix_item("MIX 200(90%)+125(10%)", 125, 120),
        _mix_item("MIX 200(90%)+125(10%)", 125, 120, 손잡이방향="좌"),
        _mix_item("MIX 200(90%)+125(10%)", 83, 55),
    ])])
    products = [r for r in rows if r.get("_특수") is None]
    assert products[0]["색상"] == " B Mix 200+125"
    assert products[1]["색상"] is None and str(products[1]["세로"]).strip() == '"'
    assert products[2]["색상"] == " B Mix 200+125"          # 구성행 다음은 품목명을 다시 적는다
    kinds = [("P" if r.get("_특수") is None else r.get("_특수")) for r in rows]
    assert kinds[:7] == ["P", "P", "믹스", "믹스", "P", "믹스", "믹스"]
    assert _parts(rows)[:4] == [("200", 108.0, None), ("125", 12.0, None),
                                ("200", 49.5, None), ("125", 5.5, None)]


def test_ledger_rows_four_color_mix():
    rows = _rows([_order("JL", [_mix_item("MIX 280(92cm)+160(23cm)+280(92cm)+160(23cm)", 190, 230,
                                          kind="원코드")])])
    product = next(r for r in rows if r.get("_특수") is None)
    assert product["색상"] == " B 원코드 Mix"
    assert _parts(rows) == [("280", 92.0, "280+160+280+160"), ("160", 23.0, None),
                            ("280", 92.0, None), ("160", 23.0, None)]


def test_ledger_excel_component_rows_have_thick_code_box_and_merged_notes(tmp_path):
    orders = [_order("JL", [_mix_item("MIX 200(90%)+125(10%)", 83, 55)])]
    path = tmp_path / "장부.xlsx"
    output.build_ledger(_rows(orders), path, header_date=_cases.HEADER)
    ws = openpyxl.load_workbook(path).worksheets[0]
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert ws["C3"].value == " B Mix 200+125"
    assert (ws["D4"].value, ws["G4"].value) == ("200", ") 49.5cm")
    assert (ws["D5"].value, ws["G5"].value) == ("125", ") 5.5cm")
    assert {"D4:F4", "G4:H4", "D5:F5", "G5:H5"} <= merged
    assert ws["D4"].border.left.style == "thick" and ws["D4"].border.top.style == "thick"
    assert ws["F4"].border.right.style == "thick"
    assert any(r.startswith("K3:") and r.endswith("5") for r in merged), merged   # 기재사항이 구성행까지 병합
    assert "M3:M5" in merged                                                       # 출고일도 구성행까지


def test_worksheet_customer_merge_continues_through_mix_rows(tmp_path):
    orders = [_order("DU", [{"품목코드": "102", "종류": "투코드", "타입": "C자", "가로": 70, "세로": 100},
                            _mix_item("MIX 200(90%)+125(10%)", 83, 55),
                            {"품목코드": "102", "종류": "투코드", "타입": "C자", "가로": 75, "세로": 100}])]
    path = tmp_path / "작업지시서.xlsx"
    output.build_worksheet(_rows(orders), path, header_date=_cases.HEADER)
    ws = openpyxl.load_workbook(path).worksheets[0]
    merged = [r for r in ws.merged_cells.ranges if r.min_col == 2 and r.max_col == 2]
    assert any(r.max_row - r.min_row >= 4 for r in merged), [str(r) for r in merged]


def test_mix_status_is_confirm():
    order = _order("DU", [_mix_item("MIX 102(50%)+500(50%)", 100, 100)])
    output.synchronize_order_for_outputs(order)
    issues = [x for x in validate(order, M) if x[1] == "MIX확인"]
    assert issues and issues[0][0] == "yellow" and "102(50)+500(50)" in issues[0][2]


def test_di_keeps_existing_mix_format():
    order = _order("DI", [_mix_item("MIX 200+125", 125, 120, 기재사항="최수령")])
    rows = _rows([order])
    assert not _parts(rows)
    assert not order["items"][0].get("_blind_mix")


# ───────────────────────── 경영박사 ─────────────────────────
def _edi(orders, tmp_path):
    path = tmp_path / "edi.xls"
    for o in orders:
        output.synchronize_order_for_outputs(o)
    output.build_erp(orders, M, path, None)
    sheet = xlrd.open_workbook(str(path)).sheet_by_index(0)
    return [sheet.row_values(r) for r in range(sheet.nrows)]


@pytest.mark.parametrize("kind,type_,code,extra", [
    ("투코드", "C자", "B-MIX-투코드/25mm", "B추가비용-MIX"),
    ("원코드", "C자", "B-MIX-원코드/25mm", "B추가비용-MIX"),
    ("원코드", "L자", "B-MIX-L18/원코드/25mm", "B추가비용-MIX/Ltype"),
])
def test_edi_mix_code_memo_and_zero_cost_row(tmp_path, kind, type_, code, extra):
    orders = [_order("DU", [_mix_item("MIX 200(90%)+125(10%)", 125, 120, kind=kind, type_=type_),
                            _mix_item("MIX 200(90%)+125(10%)", 125, 120, kind=kind, type_=type_)])]
    rows = _edi(orders, tmp_path)
    products = [r for r in rows if r[8] == code]
    assert len(products) == 2
    assert all(" 200(108)+125(12)" in r[15] for r in products)
    assert r"200(108)+125(12)" == products[0][15].split(" ", 1)[1].split("/")[0]
    extras = [r for r in rows if r[8] == extra]
    assert len(extras) == 2 and all(r[11] == 1.0 and r[12] == 0 and r[13] == 0 for r in extras)


def test_l18_shutter_mix_is_registered_in_master():
    order = _order("DU", [_mix_item("MIX 200+125", 100, 150, kind="셔터", type_="L자")])
    output.synchronize_order_for_outputs(order)
    hit = M.find_order_item(order["items"][0], "DU")
    assert hit and hit["품명"] == "B-MIX-L18/셔터/25mm"
    assert hit["관리코드"] == "B-MIX-L18/원코드/25mm"      # 경영박사 화면 관리코드
    assert not [x for x in validate(order, M) if x[1] == "품목미등록"]


# ───────────────────────── 사전점검 / 장부 불러오기 ─────────────────────────
def test_precheck_edit_mix_length_applies_to_group():
    orders = [_order("DU", [_mix_item("MIX 200(90%)+125(10%)", 125, 120),
                            _mix_item("MIX 200(90%)+125(10%)", 125, 120)])]
    preview, _ = dw.build_preview_rows(orders, _cases.SHIP, M)
    part_rows = [r for r in preview if r.get("_detail_kind") == "mix_part"]
    assert [dw.row_display_values(r, 0)["수량"] for r in part_rows] == [") 108cm", ") 12cm"]
    assert dw.editable_columns_for_row(part_rows[0]) == {"가로", "수량"}
    dw.apply_preview_cell_edit(orders, part_rows[0], "수량", ") 100cm")
    dw.apply_preview_cell_edit(orders, part_rows[1], "수량", "20")
    preview, _ = dw.build_preview_rows(orders, _cases.SHIP, M)
    assert _parts(preview) == [("200", 100.0, None), ("125", 20.0, None)]
    assert all(it["_mix_parts"] == [{"코드": "200", "길이": 100.0}, {"코드": "125", "길이": 20.0}]
               for it in orders[0]["items"])


def test_mix_ledger_roundtrip_keeps_component_lengths(tmp_path):
    source = [_order("JL", [_mix_item("MIX 200(108cm)+125(12cm)", 125, 120),
                            _mix_item("MIX 200(108cm)+125(12cm)", 125, 120),
                            _mix_item("MIX 990(18cm)+102(72cm)", 120, 90)])]
    path = tmp_path / "장부.xlsx"
    output.build_ledger(_rows(source), path, header_date=_cases.HEADER)
    _, orders, _ = output.read_ledger(str(path), M, path.name)
    ready = dw.ledger_orders_for_precheck(orders, path, "blind", _cases.SHIP)
    assert _parts(_rows(ready)) == _parts(_rows(source))
