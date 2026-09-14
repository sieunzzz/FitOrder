"""2026-09-14 업체별 2차 수정 (사용자 확정 규칙)."""
import copy
from collections import Counter
from pathlib import Path

import openpyxl
import xlrd

import _cases
import desktop_workflow as dw
import output
import parsers
from master_registry import get_master_for_mode
from roll_combo import normalize_roll_order

MB = get_master_for_mode("blind")
MR = get_master_for_mode("roll_combo")
NOTE1 = 11                     # 장부 세부 양식 기재사항1 열
TAEAN = "충청남도 태안군 근흥면 근흥로 687-1, 서울한약방"


def _blind(client, items, delivery=None, **extra):
    order = {"거래처": client, "주문번호": "7", "고객명": "현장", "배송": delivery or {},
             "items": items}
    order.update(extra)
    order = dw._normalize_ingested_order(order, Path("t.png"), "blind")
    order["_ship_date"] = _cases.SHIP.isoformat()
    return order


def _item(code="102", w=100, h=150, kind="투코드", **extra):
    return {"품목코드": code, "종류": kind, "타입": "C자", "가로": w, "세로": h, **extra}


def _rows(orders, M=MB):
    return dw.build_output_rows(orders, _cases.SHIP, M=M, sort_di=True)


def _products(rows):
    return [r for r in rows if r.get("_특수") is None]


def _edi(orders, M, path):
    for o in orders:
        output.synchronize_order_for_outputs(o)
    output.build_erp(orders, M, path, None)
    sheet = xlrd.open_workbook(str(path)).sheet_by_index(0)
    return [sheet.row_values(r) for r in range(sheet.nrows)]


def _delivery(name, address=TAEAN):
    return {"방식": "택배", "주소": address, "수령인": name, "연락처": "010-0000-0000",
            "선불착불": "선불"}


# ───────────────────────── 유앤 ─────────────────────────
def _unit_order():
    # 90cm 2창(창당 피스 2) + 180cm 2창(창당 피스 3) = 피스 10EA
    return _blind("유앤", [_item(w=90, h=210, 창개수=2, _수동특이="석고"),
                          _item(w=180, h=210, 창개수=2, _수동특이="석고")],
                  _delivery("최군식"), 고객명="최군식", 전체기재사항="피스",
                  _unit_accessories=[{"품명": "피스(석고앙카)", "세트": 2, "피스수": 2},
                                     {"품명": "피스(석고앙카)", "세트": 2, "피스수": 3}])


def test_unit_piece_total_ea_in_ledger_and_edi(tmp_path):
    order = _unit_order()
    rows = _rows([order])
    acc = [(r["문구"], r["수량"]) for r in rows if r.get("_특수") == "부속"]
    assert acc == [("피스(석고앙카)", "10EA")]
    assert all(r["기재사항"] == "피스/최군식" for r in _products(rows))
    edi = _edi([order], MB, tmp_path / "e.xls")
    piece = [r for r in edi if r[8] == "피스(석고앙카)"]
    assert len(piece) == 1 and piece[0][11] == 0 and piece[0][15] == "10EA/최군식"
    assert all(r[15].endswith("석고/피스/최군식") for r in edi if str(r[8]).startswith("B102"))
    pack = next(r for r in edi if r[8] == "포장비용(25mm)")
    assert pack[15] == "최군식"


def test_unit_accessories_survive_ledger_roundtrip(tmp_path):
    orders = [_unit_order()]
    direct = _edi(copy.deepcopy(orders), MB, tmp_path / "a.xls")
    rows = _rows(copy.deepcopy(orders))
    ledger = tmp_path / "장부.xlsx"
    output.build_ledger(rows, ledger, header_date=_cases.HEADER)
    _, read, _ = output.read_ledger(str(ledger), MB, ledger.name)
    ready = dw.ledger_orders_for_precheck(read, ledger, "blind", _cases.SHIP)
    via = _edi(ready, MB, tmp_path / "b.xls")
    key = lambda rows: Counter(tuple(r[8:18]) for r in rows)
    assert key(via) == key(direct)


def test_unit_gypsum_names(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1, "UNITNS 발주서")
    ws.cell(2, 1, "고객명"); ws.cell(2, 2, "최군식"); ws.cell(2, 7, TAEAN)
    ws.cell(3, 2, "010-0000-0000")
    ws.cell(5, 1, "NO")
    for r, (w, note) in enumerate(((90, "석고"), (180, "석고칼블럭-날개")), 6):
        for c, v in {1: r - 5, 2: "알루미늄블라인드", 3: "WH102", 4: w, 5: 210, 6: "좌", 7: 1, 8: note}.items():
            ws.cell(r, c, v)
    path = tmp_path / "unit.xlsx"
    wb.save(path)
    order = parsers.parse_unitns(path)[0]
    assert [a["품명"] for a in order["_unit_accessories"]] == ["피스(석고앙카)", "피스(석고날개)"]


# ───────────────────────── 휴안 ─────────────────────────
def _huan_order():
    order = {"거래처": "휴안", "고객명": "박시은", "전체기재사항": "포장비용",
             "배송": _delivery("박시은"),
             "_huan_messages": ["★이지픽스메탈(3) 3 ★석고용앙카나사 3 부재시 문앞에 놓아주세요"],
             "items": [_item(w=90, h=115), _item(w=90, h=115), _item(w=153, h=200)]}
    parsers._apply_huan_message_notes(order)
    order = dw._normalize_ingested_order(order, Path("t.png"), "blind")
    order["_ship_date"] = _cases.SHIP.isoformat()
    return order


def test_huan_accessory_before_name_and_set(tmp_path):
    order = _huan_order()
    rows = _rows([order])
    assert [(r["문구"], r["수량"]) for r in rows if r.get("_특수") == "부속"] == [("노피스(3)", "3set")]
    tall = next(r for r in _products(rows) if r["세로"] == 200)
    assert tall["기재사항"] == "석고앙카3/박시은"
    edi = _edi([order], MB, tmp_path / "e.xls")
    memos = [r[15] for r in edi if str(r[8]).startswith("B102")]
    assert sum(m.endswith(" 앙카3/박시은") for m in memos) == 1
    assert next(r for r in edi if r[8] == "노피스브라켓(3EA)")[15] == "3set/박시은"
    assert next(r for r in edi if r[8] == "포장비용(25mm)")[15] == "박시은"


def test_erp_address_detail_on_new_line_without_comma(tmp_path):
    edi = _edi([_huan_order()], MB, tmp_path / "e.xls")
    i = next(k for k, r in enumerate(edi) if str(r[8]).startswith("#택배"))
    assert [edi[i][15], edi[i + 1][15], edi[i + 2][15]] == ["태안군 근흥면", "근흥로 687-1", "서울한약방"]
    assert not any("687-1," in str(r[15]) for r in edi)
    address = next(r for r in _rows([_huan_order()]) if r.get("_특수") == "☆")
    assert address["문구"] == "태안군 근흥면 근흥로 687-1, 서울한약방"     # 장부는 쉼표 유지


def test_worksheet_has_no_address_phone_or_packing(tmp_path):
    rows = _rows([_huan_order()])
    output.build_worksheet(rows, tmp_path / "w.xlsx", header_date=_cases.HEADER)
    text = " ".join(str(c.value) for row in openpyxl.load_workbook(tmp_path / "w.xlsx").active.iter_rows()
                    for c in row if c.value is not None)
    for gone in ("포장비용", "근흥로", "010-0000-0000", "선불"):
        assert gone not in text
    assert "노피스(3)" in text and "부재시 문앞에 놓아주세요" in text
    output.build_ledger(rows, tmp_path / "l.xlsx", header_date=_cases.HEADER)
    ledger = " ".join(str(c.value) for row in openpyxl.load_workbook(tmp_path / "l.xlsx").active.iter_rows()
                      for c in row if c.value is not None)
    assert "포장비용" in ledger and "010-0000-0000" in ledger


# ───────────────────────── 전체 ─────────────────────────
def test_piece_before_name_for_common_note():
    order = _blind("DD", [_item(w=100), _item(w=110)], _delivery("김수령"), 전체기재사항="피스")
    assert _products(_rows([order]))[0]["기재사항"] == "피스/김수령/현장"


def test_ditto_mark_left_bottom(tmp_path):
    rows = _rows([_blind("DD", [_item(w=90, h=115), _item(w=90, h=115)])])
    output.build_ledger(rows, tmp_path / "l.xlsx", header_date=_cases.HEADER)
    ws = openpyxl.load_workbook(tmp_path / "l.xlsx").active
    cell = next(ws.cell(r, 6) for r in range(3, 20) if str(ws.cell(r, 6).value or "").strip() == '"')
    assert (cell.alignment.horizontal, cell.alignment.vertical) == ("left", "bottom")


def test_huan_message_kalblock_is_gypsum_wing():
    specials, rest = parsers._huan_message_info("석고칼블럭-날개 문앞")
    assert specials == ["석고날개"] and rest == "문앞"


def test_roll_handle_rounds_down_to_5(tmp_path):
    def roll(handle):
        order = {"거래처": "천안)채원", "배송": {}, "_product_mode": "roll_combo",
                 "items": [{"롤품명": "달리", "롤색상": "그레이", "가로": 120, "세로": 180,
                            "손잡이길이": handle}]}
        order = normalize_roll_order(order, MR, "천안)채원")
        order["_ship_date"] = _cases.SHIP.isoformat()
        return order

    assert _products(_rows([roll(127)], MR))[0]["모형2"] == 125
    assert _products(_rows([roll(153)], MR))[0]["모형2"] is None      # 150 = 기본값
    edi = _edi([roll(127)], MR, tmp_path / "r.xls")
    assert any("손125" in str(r[15]) for r in edi)
