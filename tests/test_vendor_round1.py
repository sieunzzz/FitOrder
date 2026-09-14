"""2026-09-13 업체별 1차 수정 (사용자 확정 규칙)."""
from pathlib import Path

import openpyxl
import xlrd

import _cases
import desktop_workflow as dw
import output
import parsers
from master_registry import get_master_for_mode
from roll_combo import normalize_roll_order

MB = get_master_for_mode("blind")
MR = get_master_for_mode("roll_combo")
NOTE1, NOTE2 = 11, 12          # 장부 세부 양식 기재사항1/2 열


def _blind(client, items, delivery=None, **extra):
    order = {"거래처": client, "주문번호": "7", "고객명": "현장", "배송": delivery or {},
             "items": items}
    order.update(extra)
    order = dw._normalize_ingested_order(order, Path("t.png"), "blind")
    order["_ship_date"] = _cases.SHIP.isoformat()
    return order


def _item(code="102", w=100, h=150, kind="투코드", **extra):
    return {"품목코드": code, "종류": kind, "타입": "C자", "가로": w, "세로": h, **extra}


def _rows(orders, M=MB):
    return dw.build_output_rows(orders, _cases.SHIP, M=M, sort_di=True)


def _products(rows):
    return [i for i, r in enumerate(rows) if r.get("_특수") is None]


def _edi(orders, M, tmp_path):
    for o in orders:
        output.synchronize_order_for_outputs(o)
    path = tmp_path / "edi.xls"
    output.build_erp(orders, M, path, None)
    sheet = xlrd.open_workbook(str(path)).sheet_by_index(0)
    return [sheet.row_values(r) for r in range(sheet.nrows)]


# ───────────────────────── 전체 ─────────────────────────
def test_simple_means_shutter():
    assert parsers._huan_product("YL500 심플")[2] == "셔터"
    assert parsers._product_parts("알루미늄블라인드", "WH102", "심플")[2] == "셔터"
    assert parsers.DI_KIND["심플"] == "셔터"


def test_note1_merge_covers_empty_note2():
    order = _blind("DD", [_item(w=100), _item(w=110)], 전체기재사항="피스")
    rows = _rows([order])
    a, b = _products(rows)[:2]
    assert rows[a]["기재사항"] and rows[a]["기재사항"] == rows[b]["기재사항"]
    assert not rows[a]["기재사항2"] and not rows[b]["기재사항2"]
    assert (a, b, NOTE1, NOTE2) in output.ledger_merge_ranges(rows)


def test_same_note1_and_note2_merge_left_right():
    order = _blind("RT", [_item(w=100, 기재사항="안방", 설치장소="안방"),
                          _item(w=110, 기재사항="거실", 설치장소="거실")])
    rows = _rows([order])
    b = _products(rows)[1]
    assert rows[b]["기재사항"] == "거실" and rows[b]["기재사항2"] is None
    assert any(m[0] <= b <= m[1] and m[2] == NOTE1 and m[3] == NOTE2
               for m in output.ledger_merge_ranges(rows))


def test_jl_note2_merges_vertically():
    order = _blind("JL", [_item(w=100, 설치장소="거실"), _item(w=110, 설치장소="거실")],
                   전체기재사항="피스")
    rows = _rows([order])
    a, b = _products(rows)[:2]
    assert rows[a]["기재사항2"] == rows[b]["기재사항2"] == "거실"
    assert (a, b, NOTE2, NOTE2) in output.ledger_merge_ranges(rows)


def test_total_window_row_all_modes():
    blind = _blind("DD", [_item(창개수=2), _item(w=90)])
    preview, _ = dw.build_preview_rows([blind], _cases.SHIP, MB)
    last = preview[-1]
    assert last["_detail_kind"] == "total"
    assert dw.row_display_values(last, len(preview))["가로"] == "3창"
    assert dw.editable_columns_for_row(last) == set()
    assert dw.preview_statuses(preview, [blind], MB)[-1][0] == "ok"
    for mode, client in (("holding", "휴안"), ("roll_combo", "천안)채원")):
        orders = _cases.MODES[mode][1](client)
        preview, _ = dw.build_preview_rows(orders, _cases.SHIP, get_master_for_mode(mode))
        assert preview[-1]["_detail_kind"] == "total" and preview[-1]["문구"].endswith("창")


def test_joint_windows_get_hash_first_in_edi(tmp_path):
    order = _blind("DD", [_item(kind="원코드", 연창=True, 손잡이길이=120),
                          _item(kind="원코드", w=90, 연창=True)])
    rows = _rows([order])
    assert rows[_products(rows)[0]]["특이"] == "#"
    memos = [r[15] for r in _edi([order], MB, tmp_path) if r[8].startswith("B102")]
    assert len(memos) == 2 and all(m.split(" ", 1)[1].startswith("#") for m in memos)


def test_freight_row_qty_one_and_sender_code(tmp_path):
    order = _blind("DD", [_item()], {"방식": "택배", "주소": "서울특별시 테스트구 테스트로 1",
                                      "수령인": "김수령", "연락처": "010-0000-0000",
                                      "발신": "테스트 010-1111-1111"})
    edi = _edi([order], MB, tmp_path)
    ship = next(r for r in edi if str(r[8]).startswith("#택배"))
    assert ship[11] == 1
    assert all(r[11] == 0 for r in edi if r[8] == "**")
    assert any(r[8] == r[9] == "└ 발 신:" for r in edi)


def test_ledger_roundtrip_keeps_vertically_merged_notes(tmp_path):
    order = _blind("JL", [_item(w=100, 설치장소="거실"), _item(w=110, 설치장소="거실"),
                          _item(w=120, 설치장소="거실")], 전체기재사항="피스", 고객명="임지애")
    rows = _rows([order])
    path = tmp_path / "장부.xlsx"
    output.build_ledger(rows, path, header_date=_cases.HEADER)
    _, orders, _ = output.read_ledger(str(path), MB, path.name)
    items = [it for o in orders for it in o.get("items") or []]
    assert len(items) == 3 and all(it.get("설치장소") == "거실" or "거실" in str(it.get("기재사항"))
                                   for it in items)


# ───────────────────────── 대일 ─────────────────────────
def test_di_worksheet_has_no_empty_sheets(tmp_path):
    order = _blind("DI", [_item(기재사항="최수령")])
    rows = _rows([order])
    output.build_ledger(rows, tmp_path / "l.xlsx", header_date=_cases.HEADER)
    output.build_worksheet(rows, tmp_path / "w.xlsx", header_date=_cases.HEADER)
    ledger = openpyxl.load_workbook(tmp_path / "l.xlsx").sheetnames
    worksheet = openpyxl.load_workbook(tmp_path / "w.xlsx").sheetnames
    assert worksheet == ledger and len(worksheet) == 1


# ───────────────────────── 휴안 ─────────────────────────
def _huan_order():
    order = {"거래처": "휴안", "고객명": "김고객", "전체기재사항": "포장비용",
             "배송": {"방식": "택배", "주소": "서울특별시 테스트구 테스트로 1", "수령인": "김고객",
                      "연락처": "010-0000-0000", "선불착불": "선불", "발신": "휴안 010-2222-2222"},
             "_huan_messages": ["★이지픽스메탈(3) 2 ★석고용앙카나사 3"],
             "items": [_item(w=150, h=120), _item(w=90, h=200), _item(w=80, h=100)]}
    parsers._apply_huan_message_notes(order)
    order = dw._normalize_ingested_order(order, Path("t.png"), "blind")
    order["_ship_date"] = _cases.SHIP.isoformat()
    return order


def test_huan_accessory_row_prepaid_and_anchor(tmp_path):
    order = _huan_order()
    assert order["_huan_accessories"] == [{"표시": "노피스(3)", "수량": "2set"}]
    rows = _rows([order])
    acc = [r for r in rows if r.get("_특수") == "부속"]
    assert [(r["문구"], r["수량"]) for r in acc] == [("노피스(3)", "2set")]
    notes = [r["기재사항"] or "" for r in rows if r.get("_특수") is None]
    assert sum("석고앙카3" in n for n in notes) == 1 and "노피스" not in "".join(notes)
    tall = next(r for r in rows if r.get("_특수") is None and r.get("세로") == 200)
    assert "석고앙카3" in str(tall["기재사항"]).split("/")
    address = next(r for r in rows if r.get("_특수") == "☆")
    who = next(r for r in rows if r.get("_특수") == "" and "010-0000-0000" in str(r.get("문구")))
    assert address.get("선불") is None and who.get("선불") == "선불"
    edi = _edi([order], MB, tmp_path)
    memos = [r[15] for r in edi if str(r[8]).startswith("B102")]
    assert sum("앙카3" in m for m in memos) == 1 and not any("석고앙카" in m for m in memos)
    nopiece = [r for r in edi if r[8] == "노피스브라켓(3EA)"]
    assert len(nopiece) == 1 and nopiece[0][9] == "노피스브라켓(3EA)" and nopiece[0][11] == 2


def test_huan_etc_column(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in {1: "주문일", 3: "수취인", 4: "연락처", 6: "주소", 7: "상품명", 9: "사이즈",
                 10: "손잡이", 11: "수량", 12: "배송메시지", 13: "기타"}.items():
        ws.cell(1, c, h)
    for c, v in {3: "김고객", 4: "010-1234-5678", 6: "서울특별시 강남구 테스트로 1",
                 7: "YL500 원코드", 9: "100*150", 10: "좌", 11: 1,
                 13: "★이지픽스메탈(3) 2/손120/높게 달아주세요"}.items():
        ws.cell(2, c, v)
    path = tmp_path / "huan.xlsx"
    wb.save(path)
    order = parsers.parse_huan(path)[0]
    it = order["items"][0]
    assert it["손잡이길이"] == 120 and "높게 달아주세요" in str(it["기재사항"])
    assert order["_huan_accessories"] == [{"표시": "노피스(3)", "수량": "2set"}]


# ───────────────────────── 보노 · 미래가공 (블라인드 + 롤콤비) ─────────────────────────
BONO_DELIVERY = {"방식": "화물", "화물지점": "안양박달", "주소": "안양박달-보노", "수령인": "보노",
                 "연락처": "010-2478-9290", "선불착불": "착불", "발신": "보노 010-2478-9290"}


def _check_bono_edi(edi, product_prefix):
    product = next(r for r in edi if r[15].startswith(product_prefix))
    assert product[15].split(" ", 1)[1].startswith("손120/박달/거실1-1/부천67327013")
    ship = next(i for i, r in enumerate(edi) if str(r[8]).startswith("#화물"))
    assert edi[ship][11] == 1 and edi[ship][15] == "안양박달 - 보노"
    assert edi[ship + 1][8] == "**" and edi[ship + 1][15] == "010-2478-9290"
    assert not any("/010" in str(r[15]) or "/ 010" in str(r[15]) for r in edi)
    assert any(r[8] == "포장비용(B/R/C/S/HC)" for r in edi)


def test_bono_blind(tmp_path):
    order = _blind("보노", [_item(손잡이길이=120, 설치장소="거실1-1", 작동방식필증="방염필증")],
                   dict(BONO_DELIVERY), 주문번호="부천67327013",
                   _bono_receiver_is_bono=True, _bono_branch="안양박달")
    rows = _rows([order])
    assert "방염필증" in str(rows[_products(rows)[0]]["특이"])
    address = next(r for r in rows if r.get("_특수") == "☆")
    assert address["문구"] == "안양박달 - 보노 010-2478-9290"
    _check_bono_edi(_edi([order], MB, tmp_path), "100.0*150.0/1EA")


def test_bono_roll(tmp_path):
    order = {"거래처": "안산)보노", "주문번호": "부천67327013", "배송": dict(BONO_DELIVERY),
             "_bono_receiver_is_bono": True, "_bono_branch": "안양박달", "_product_mode": "roll_combo",
             "items": [{"롤품명": "달리", "롤색상": "그레이", "가로": 120, "세로": 180,
                        "손잡이길이": 120, "설치장소": "거실1-1", "작동방식필증": "방염필증"}]}
    order = normalize_roll_order(order, MR, "안산)보노")
    order["_ship_date"] = _cases.SHIP.isoformat()
    rows = _rows([order], MR)
    assert "방염필증" in str(rows[_products(rows)[0]]["특이"])
    assert any(r.get("_특수") == "#" and r.get("문구") == "포장비용" for r in rows)
    _check_bono_edi(_edi([order], MR, tmp_path), "120.0*180.0/1EA")


# ───────────────────────── JL / 루임트 ─────────────────────────
def test_jl_multiple_colors_in_one_cell_is_mix():
    order = _blind("JL", [_item(code="102", h=100, 색상원문="102 500")])
    rows = _rows([order])
    assert rows[_products(rows)[0]]["색상"] == " B Mix 102+500"
    assert [(r["코드"], r["길이"]) for r in rows if r.get("_특수") == "믹스"] == [("102", 50.0), ("500", 50.0)]


def test_rt_full_receiver_in_address_and_short_in_note():
    import extract
    data = extract._postprocess_rt({"배송": {"방식": "화물", "화물지점": "오산삼미", "수령인": "더",
                                             "연락처": "010-5391-1057", "선불착불": "착불"},
                                    "items": [{"기재사항": "더커튼/동탄푸르1502"}]})
    assert data["배송"]["수령인"] == "더커튼"
    assert data["items"][0]["기재사항"] == "더/동탄푸르1502"
    order = _blind("RT", [_item(기재사항="더/동탄푸르1502")], data["배송"])
    rows = _rows([order])
    address = [r for r in rows if r.get("_특수") == "☆"]
    assert address and address[0]["문구"] == "오산삼미 - 더커튼 / 010-5391-1057"
    assert "더커튼" not in str(rows[_products(rows)[0]]["기재사항"])
