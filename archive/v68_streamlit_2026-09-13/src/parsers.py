"""
엑셀 발주 파서 (휴안 · DI · 유앤) — AI 를 거치지 않고 직접 읽는다.
출력 형태는 extract.extract_order() 와 동일하다.
"""
import re

import openpyxl

from rules import (DI_CODE_ORDER, DI_COLOR, DI_COLOR_UNSURE, clean_delivery_notice,
                   di_mix_from_codes, di_mix_info, normalize_di_color_name)
from holding import (HOLDING_FEATURE_ENABLED, holding_product_from_text, looks_like_holding_operation,
                     normalize_holding_operation)

SIZE = re.compile(r"(\d+(?:\.\d+)?)\s*[*x×X]\s*(\d+(?:\.\d+)?)")


def _blank_item():
    return {"제품군": "블라인드", "홀딩방식": None, "홀딩레일": None,
            "홀딩상하로라": False, "홀딩부속": None, "홀딩부속색상": None,
            "품목코드": None, "색상원문": None, "타입": "C자", "종류": "투코드",
            "가로": None, "세로": None, "수량": None, "손잡이방향": None,
            "손잡이길이": None, "연창": False, "설치장소": None,
            "기재사항": None, "예외품목": None, "원문": None,
            "창개수": 1, "좌개수": None, "우개수": None,
            "확신도": {"가로": 1.0, "세로": 1.0, "품목코드": 1.0, "손잡이": 1.0}}


def _shell(client):
    return {"거래처": client, "발주일": None, "고객명": None, "items": [],
            "배송": {"방식": None, "화물지점": None, "주소": None,
                     "수령인": None, "연락처": None, "선불착불": None,
                     "전달사항": None},
            "전체원문": None, "전체기재사항": None}


# ─────────────────────────────────────────────
# 휴안
# ─────────────────────────────────────────────
def _huan_product(s):
    """'YL500 L원코드' -> (500, L자, 원코드)"""
    s = str(s or "").strip().upper()
    m = re.search(r"([A-Z]{1,2})?\s*(\d{3,4}[A-Za-z]*)", s)
    code = m.group(2) if m else None
    # 휴안은 L투코드를 `L코드`, `L투`, `L 투코드` 등으로 줄여 쓰기도 한다.
    # L18/L21도 같은 L타입으로 읽되 실제 출력 품명은 공통 규칙상 L18이다.
    type_ = ("L자" if re.search(
        r"(?:^|[\s/_-])L\s*(?=원|투|셔|코드|18|21|타입|자|$)", s)
        else "C자")
    if "셔터" in s:
        kind = "셔터"
    elif "원코드" in s or re.search(r"원\s*코드", s):
        kind = "원코드"
    else:
        kind = "투코드"
    return code, type_, kind


SIDO_SHORT = {"서울특별시": "서울시", "부산광역시": "부산시", "대구광역시": "대구시",
              "인천광역시": "인천시", "광주광역시": "광주시", "대전광역시": "대전시",
              "울산광역시": "울산시", "세종특별자치시": "세종시"}
DO = ("경기도", "강원도", "강원특별자치도", "충청북도", "충청남도", "전라북도",
      "전북특별자치도", "전라남도", "경상북도", "경상남도", "제주특별자치도")


def _short_addr(addr):
    """'경상남도 진주시 초전북로 151 (초전동) 인사인광고기획'
         -> '진주시 초전북로 151, 인사인광고기획'
       '서울특별시 영등포구 도신로 161 (도림동, 정성빌) 201호'
         -> '서울시 영등포구 도신로 161, 정성빌 201호'"""
    s = str(addr or "").strip()
    s = re.sub(r"^\s*\(\s*\d{5}\s*\)\s*", "", s)
    for k, v in SIDO_SHORT.items():          # 광역시/특별시는 줄여서 유지
        if s.startswith(k):
            s = v + s[len(k):]
            break
    else:
        for d in DO:                          # 도는 삭제
            if s.startswith(d):
                s = s[len(d):]
                break

    def keep_building(m):
        """괄호 안에서 '○○동' 은 버리고 건물명만 남긴다"""
        parts = [x.strip() for x in m.group(1).split(",")]
        keep = [x for x in parts if not re.search(r"[동읍면리]$", x)]
        return (", " + " ".join(keep)) if keep else ""

    s = re.sub(r"\s*\(([^)]*)\)", keep_building, s)
    s = re.sub(r"\s{2,}", " ", s).strip()
    s = re.sub(r"(\d+)동\s*(\d+)호\b", r"\1-\2", s)
    s = re.sub(r"아파트\b", "A", s)
    if "," not in s:                      # 끝의 상호/건물명 앞에 콤마
        head, _, tail = s.rpartition(" ")
        if head and tail and not re.search(r"\d", tail):
            s = f"{head}, {tail}"
    return s


def _huan_message_info(value):
    """휴안 배송메시지에서 제작 기재사항(노피스/석고앙카)과 나머지를 분리한다.

    예: ★이지픽스메탈(3) 1 -> 노피스(3)1
        ★석고용앙카나사 1 -> 석고앙카1
    """
    text = str(value or "").strip()
    if not text:
        return [], None
    specials = []

    def add(value):
        if value and value not in specials:
            specials.append(value)

    # 이미 줄여 적은 메시지도 그대로 인식한다.
    for m in re.finditer(r"노피스\s*\((\d+)\)\s*(\d+)", text):
        add(f"노피스({m.group(1)}){m.group(2)}")
    for m in re.finditer(r"석고앙카\s*(\d*)", text):
        add(f"석고앙카{m.group(1)}" if m.group(1) else "석고앙카")
    for word in ("콘크리트", "석고날개"):
        if word in text:
            add(word)

    # 원 발주 메시지의 제품명을 공장 장부용 짧은 표기로 변환한다.
    for m in re.finditer(r"[★☆*]*\s*이지픽스메탈\s*\((\d+)\)\s*(\d+)", text):
        add(f"노피스({m.group(1)}){m.group(2)}")
    for m in re.finditer(r"[★☆*]*\s*석고용앙카나사\s*(\d+)", text):
        add(f"석고앙카{m.group(1)}")

    remainder = text
    remainder = re.sub(r"[★☆*]*\s*이지픽스메탈\s*\(\d+\)\s*\d+", " ", remainder)
    remainder = re.sub(r"[★☆*]*\s*석고용앙카나사\s*\d+", " ", remainder)
    remainder = re.sub(r"노피스\s*\(\d+\)\s*\d+", " ", remainder)
    remainder = re.sub(r"석고앙카\s*\d*", " ", remainder)
    remainder = remainder.replace("콘크리트", " ").replace("석고날개", " ")
    remainder = re.sub(r"\s*[-/|,]+\s*", "/", remainder)
    remainder = re.sub(r"/{2,}", "/", remainder).strip(" /-★☆*")
    return specials, (remainder or None)


def _apply_huan_message_notes(order):
    """노피스/석고앙카는 한 사람 주문 중 가로가 가장 긴 한 행 맨 앞에 1회 기록."""
    raw_messages = order.pop("_huan_messages", [])
    raw_text = "/".join(dict.fromkeys(
        str(x).strip() for x in raw_messages if str(x or "").strip()))
    specials, remainder = _huan_message_info(raw_text)
    order.setdefault("배송", {})["전달사항"] = clean_delivery_notice(remainder)
    if specials:
        order["_추가부속"] = "/".join(specials)
    if not specials or not order.get("items"):
        return
    candidates = [(i, it) for i, it in enumerate(order["items"])
                  if not it.get("예외품목") and it.get("가로") not in (None, "")]
    if not candidates:
        return
    _, target = max(candidates, key=lambda pair: (float(pair[1].get("가로") or 0), -pair[0]))
    existing = [x.strip() for x in str(target.get("기재사항") or "").split("/") if x.strip()]
    target["기재사항"] = "/".join(specials + [x for x in existing if x not in specials]) or None


def parse_huan(path):
    """휴안 발주 엑셀 -> 주문 리스트 (수취인 단위로 분리)"""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    orders, cur, last_prod = [], None, (None, "C자", "투코드", "")
    for r in range(2, ws.max_row + 1):
        g = lambda c: ws.cell(r, c).value
        size = g(9)
        if not size:
            continue
        name = g(3)
        if name:
            name_text = str(name).strip()
            phone_text = str(g(4) or "").strip()
            address_text = _short_addr(g(6))
            same_person = bool(
                cur and cur.get("고객명") == name_text
                and (not phone_text or phone_text == cur["배송"].get("연락처"))
                and (not address_text or address_text == cur["배송"].get("주소"))
            )
        else:
            same_person = False
        if name and not same_person:               # 새 수령인 주문 시작
            cur = _shell("휴안")
            d = g(1)
            cur["발주일"] = d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else None
            cur["고객명"] = name_text
            cur["배송"].update({
                "방식": "택배", "주소": address_text,
                "수령인": name_text,
                "연락처": phone_text,
                "선불착불": "선불",
                "전달사항": None,
            })
            cur["_huan_messages"] = []
            cur["전체기재사항"] = "포장비용"       # 휴안은 창당 포장비
            orders.append(cur)
        if cur is None:
            continue
        if g(12):
            message = str(g(12)).strip()
            if message and message not in cur.setdefault("_huan_messages", []):
                cur["_huan_messages"].append(message)
        m = SIZE.search(str(size))
        it = _blank_item()
        if m:
            it["가로"], it["세로"] = float(m.group(1)), float(m.group(2))
        if g(7):
            code, type_, kind = _huan_product(g(7))
            last_prod = (code, type_, kind, str(g(7)).strip())
        elif cur["items"]:                       # 상품명 공란 = 위 행과 동일
            code, type_, kind, _ = last_prod
        else:
            code, type_, kind = None, "C자", "투코드"
        raw_handle = str(g(10) or "").strip()
        dir_match = re.search(r"[좌우]", raw_handle)
        len_match = re.search(r"(?:손잡이(?:길이)?|손|줄|봉|각도봉)?\s*[:=/-]?\s*(\d{2,3})(?!\d)", raw_handle)
        it.update({"품목코드": code,
                   "색상원문": (str(g(7)).strip() if g(7) else last_prod[3]),
                   "타입": type_, "종류": kind,
                   "손잡이방향": (dir_match.group(0) if dir_match else None),
                   "손잡이길이": (int(len_match.group(1)) if len_match else None),
                   # 수령인은 주문 공통값에 이미 있으므로 상품 기재사항에
                   # 다시 넣지 않는다. (박남희/박남희 중복 방지)
                   "기재사항": None, "원문": f"{g(7)} {size} {g(10)}"})
        hold = holding_product_from_text(it.get("색상원문")) if HOLDING_FEATURE_ENABLED else None
        if hold:
            it.update({"제품군": "홀딩도어", "_product_group": "holding",
                       "_ledger_prefix": "H", "품목코드": hold.get("코드"),
                       "_holding_product_name": hold["품명"],
                       "_holding_label": hold.get("장부표시"),
                       "_holding_operation": normalize_holding_operation(g(10)),
                       "_holding_upper_roller": "+상하로라" in str(it.get("색상원문") or "").replace(" ", ""),
                       "손잡이방향": None, "손잡이길이": None,
                       "타입": "C자", "종류": "투코드"})
        n = g(11)
        try:
            count = max(1, int(float(n))) if n not in (None, "") else 1
        except (TypeError, ValueError):
            count = 1
        direction_text = str(g(10) or "").strip()
        directions = re.findall(r"[좌우]", direction_text)
        if len(directions) == 1 and count > 1:
            directions *= count
        elif len(directions) != count:
            directions = ([directions[0]] * count if directions else [None] * count)

        # 수량 2 + 좌우 -> 좌 1개, 우 1개
        # 수량 2 + 우우 -> 우 2개처럼 각 창을 개별 품목으로 펼친다.
        for direction in directions:
            expanded = dict(it)
            expanded["손잡이방향"] = direction
            # 한 창씩 펼친 뒤 원 수량을 다시 보존하면 3창×3=9창이 되므로
            # 펼쳐진 행은 반드시 수량/창개수 모두 1로 고정한다.
            expanded["수량"] = 1
            expanded["창개수"] = 1
            expanded["좌개수"] = 1 if direction == "좌" else 0
            expanded["우개수"] = 1 if direction == "우" else 0
            cur["items"].append(expanded)
    for order in orders:
        _apply_huan_message_notes(order)
    return orders


# ─────────────────────────────────────────────
# DI
# ─────────────────────────────────────────────
DI_KIND = {"원": "원코드", "투": "투코드", "셔터": "셔터",
           "원코드": "원코드", "투코드": "투코드"}


def _di_direct_code(value):
    """DI 색상 칸의 숫자를 마스터 코드로 정규화한다.

    Excel 숫자 27은 표시 형식의 앞자리 0이 사라지므로 027로 복원한다.
    P/FP는 품목을 구분하는 코드이므로 버리지 않는다.
    `NA029FP`, `WH 102 P`, `B027(SV)`처럼 색상 약자·공백·괄호가
    함께 적힌 표기도 허용한다.
    """
    text = re.sub(r"\s+", "", str(value or "").strip().upper())
    # 끝의 숫자 코드만 읽어 앞의 WH/NA/GR 등 색상 약자는 제외한다.
    # FP를 P보다 먼저 매칭해 029FP가 029P로 잘리지 않게 한다.
    m = re.search(r"(\d{1,4})(FP|P)?(?:\([^)]*\))?$", text)
    if not m:
        return None
    d = m.group(1)
    suffix = m.group(2) or ""
    return (d.zfill(3) if len(d) <= 3 else d) + suffix


def _di_mix_codes(value):
    """`330(ㅇㄴ)-102(SJ)`처럼 메모가 섞인 한 셀에서 2개 이상 슬랫 코드를 뽑는다."""
    found = re.findall(r"(?<!\d)(\d{3})(?!\d)", str(value or ""))
    if len(found) < 2:
        return None
    # 같은 코드가 두 번 적힌 것은 그대로 2색 MIX로 보지 않는다.
    unique = list(dict.fromkeys(found))
    if len(unique) < 2:
        return None
    unique.sort(key=lambda x: (DI_CODE_ORDER.get(x, 10_000), int(x)))
    return "+".join(unique)

# 헤더 이름 -> 내부 필드 (파일마다 열 위치가 달라 헤더로 찾는다)
DI_HEAD = {
    "수령인": "recv", "분류": "cls", "원단/색상": "color",
    "원/투": "kind", "L/C": "type", "커버/점보": "cover", "사이즈": "size",
    "줄": "line", "피스": "piece", "비고": "note",
}


def _di_note_fields(value):
    """DI 비고에서 손잡이 길이와 나머지 특이사항을 분리한다."""
    text = str(value or "").strip()
    if not text or text in {"0", "-", "None"}:
        return None, None
    m = re.search(r"(?:손잡이(?:길이)?|손|줄|봉)\s*[:=]?\s*(\d{2,3})", text)
    handle = int(m.group(1)) if m else None
    if m:
        text = (text[:m.start()] + " " + text[m.end():]).strip()
    text = re.sub(r"^[★☆*]+\s*", "", text)
    text = re.sub(r"[\s,/]+", " ", text).strip(" -–—_=~.·ㆍ")
    # `긴급 ---`, `긴급 ——`처럼 긴급 뒤에 구분선만 붙은 표기는
    # 장부/EDI에 불필요한 기호를 남기지 않고 `긴급`만 보존한다.
    text = re.sub(r"(?<=긴급)\s*[-–—_=~.·ㆍ]+\s*$", "", text).strip()
    return handle, (text or None)


def _di_header(ws):
    """헤더 행을 찾아 {필드: 열번호} 반환. 좌측(원본) 영역만 사용."""
    for r in range(1, 8):
        vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, 20)]
        if "사이즈" not in vals:
            continue
        col, seen = {}, set()
        for i, v in enumerate(vals, 1):
            f = DI_HEAD.get(v)
            # 비고가 두 칸이면 앞쪽 비고는 사용하지 않고 뒤쪽 비고만 사용한다.
            if f == "note":
                col["note"] = i
                seen.add(f)
                continue
            # 우측 변환 영역의 두 번째 '원단/색상' 열은
            # `실버 SV 027`, `라임 B720(GN)`처럼 코드가 계산된 결과다.
            if f == "color" and f in seen and "resolved_color" not in col:
                col["resolved_color"] = i
                continue
            if f and f not in seen:        # 그 밖의 중복 헤더는 첫 열을 사용
                col[f] = i
                seen.add(f)
        if "size" in col and "color" in col:
            return r, col
    return None, {}


def _propagate_di_urgent(order):
    """같은 DI 수령인 중 한 행이라도 긴급이면 그 사람의 모든 창에 긴급을 표시한다."""
    groups = {}
    for it in order.get("items") or []:
        recipient = str(it.get("_DIrecipient_context") or it.get("기재사항") or "").strip()
        groups.setdefault(recipient, []).append(it)
    for recipient, items in groups.items():
        urgent = any("긴급" in " ".join(str(it.get(k) or "") for k in
                         ("_수동특이", "기재사항", "원문")) for it in items)
        if not urgent:
            continue
        for it in items:
            parts = [x.strip() for x in str(it.get("_수동특이") or "").split("/") if x.strip()]
            if "긴급" not in parts:
                parts.append("긴급")
            it["_수동특이"] = "/".join(parts) or None
    return order


def parse_di(path, sheet="원본"):
    """DI 발주 엑셀(좌측 원본 영역) -> 주문 1건"""
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        sheet = "장부" if "장부" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet]
    hr, col = _di_header(ws)
    if not col:
        raise ValueError("DI 발주서 형식을 인식하지 못했습니다 (헤더를 찾을 수 없음)")

    g = lambda r, f: (ws.cell(r, col[f]).value if f in col else None)
    o = _shell("DI")
    for r in range(hr + 1, ws.max_row + 1):
        size = g(r, "size")
        if not size:
            continue
        m = SIZE.search(str(size))
        if not m:
            continue
        it = _blank_item()
        raw_color = g(r, "color")
        resolved_color = g(r, "resolved_color")
        cname = str(raw_color or "").strip()
        it["색상원문"] = cname
        hold = holding_product_from_text(cname) if HOLDING_FEATURE_ENABLED else None
        # 믹스는 우측 수식의 대표 코드보다 색상표의 정확한 조합을 우선한다.
        # 홀딩도어는 별도 공식 H 품목표로 매칭한다.
        mix = None if hold else di_mix_info(cname)
        raw_mix_codes = None if hold or mix else (
            _di_mix_codes(raw_color) or _di_mix_codes(resolved_color))
        if hold:
            it.update({"제품군": "홀딩도어", "_product_group": "holding",
                       "_ledger_prefix": "H", "품목코드": hold.get("코드"),
                       "_holding_product_name": hold["품명"],
                       "_holding_label": hold.get("장부표시"),
                       "_holding_upper_roller": "+상하로라" in cname.replace(" ", ""),
                       "타입": "C자", "종류": "투코드"})
        elif mix:
            it["_mix_name"], it["_mix_codes"] = mix
            it["품목코드"] = mix[1]
        elif raw_mix_codes:
            # 이름 없이 코드가 2개 이상이면 MIX로 처리한다. 공식 조합이 코드만으로
            # 유일하게 식별되면 정식 품명, 아니면 B MIX/B 원코드 MIX를 사용한다.
            it["_raw_mix_codes"] = raw_mix_codes
            it["품목코드"] = raw_mix_codes
        else:
            normalized_name = normalize_di_color_name(cname)
            it["품목코드"] = (_di_direct_code(resolved_color)
                               or _di_direct_code(raw_color)
                               or DI_COLOR.get(normalized_name))
        if not hold and (cname in DI_COLOR_UNSURE
                or (cname and not it["품목코드"])):
            it["확신도"]["품목코드"] = 0.4      # 미등록 색상명 -> 담당자 확인
        # 종류/부속 판정 — 헤더 이름이 파일마다 달라 값으로 판단한다.
        # '원'/'투'/'셔터' 이면 종류, 그 밖의 문구면 커버·점보 등 부속.
        ks = ""
        if not hold:
            for f in ("kind", "cover"):
                v = str(g(r, f) or "").strip()
                if not v or v in ("0", "-"):
                    continue
                if v in DI_KIND:
                    ks = v
                else:
                    it["예외품목"] = "부속"
        if not hold:
            it["종류"] = DI_KIND.get(ks, "투코드")
            type_text = str(g(r, "type") or "").strip().upper()
            it["타입"] = "L자" if type_text.startswith("L") else "C자"
            # 일반 MIX 단가표는 L18/L21을 별도 품목으로 등록하므로
            # MIX일 때만 실제 L 규격을 보존한다.
            if type_text.startswith("L21"):
                it["_di_mix_family"] = "L21"
            elif type_text.startswith("L"):
                it["_di_mix_family"] = "L18"
            else:
                it["_di_mix_family"] = "C자"
            if it.get("_raw_mix_codes"):
                official = di_mix_from_codes(it["_raw_mix_codes"])
                if official:
                    it["_mix_name"], it["_mix_codes"] = official
                else:
                    it["_mix_name"] = "MIX"
                    it["_mix_codes"] = it["_raw_mix_codes"]
                    it["_generic_mix"] = True
                it["품목코드"] = it["_mix_codes"]
                it.pop("_raw_mix_codes", None)
        it["가로"], it["세로"] = float(m.group(1)), float(m.group(2))
        d = str(g(r, "line") or "").strip()
        if hold:
            it["_holding_operation"] = normalize_holding_operation(
                d if looks_like_holding_operation(d) else None)
            it["손잡이방향"] = None
        elif d.startswith(("좌", "우")):
            it["손잡이방향"] = d[0]
            n = re.search(r"(\d+)", d)
            if n:
                it["손잡이길이"] = int(n.group(1))
        recv = str(g(r, "recv") or "").strip()
        note = str(g(r, "note") or "").strip()
        piece = str(g(r, "piece") or "").strip()
        note_handle, note_special = _di_note_fields(note)
        # 비고에 손잡이 길이가 따로 적혀 있으면 그 값이 더 명시적이므로 반영한다.
        if note_handle is not None:
            it["손잡이길이"] = note_handle
        # DI 장부의 기재사항에는 수령인만 표기하고, 비고의 나머지(예: 긴급)는
        # 특이 칸으로 분리한다. 피스 원문은 확인용으로 보존한다.
        it["기재사항"] = recv if recv and recv != "None" else None
        it["_수동특이"] = note_special
        it["원문"] = f"{cname} {ks} {size} {d} {note} {piece}".strip()
        o["items"].append(it)
    o["배송"]["방식"] = None                      # DI/SP 출고
    if not o["items"]:
        raise ValueError("DI 발주서에서 유효한 주문 행을 찾지 못했습니다")
    _propagate_di_urgent(o)
    return o


# ─────────────────────────────────────────────
# 유앤아이티엔에스
# ─────────────────────────────────────────────
def _unit_directions(text, qty):
    """'좌2 우2', '좌/우', '오른쪽'을 창별 방향으로 펼친다."""
    s = str(text or "").strip()
    left = sum(int(x or 1) for x in re.findall(r"좌(?:측)?\s*(\d*)", s))
    right = sum(int(x or 1) for x in re.findall(r"우(?:른쪽|측)?\s*(\d*)", s))
    if "왼쪽" in s and not left:
        left = 1
    if "오른쪽" in s and not right:
        right = 1
    if not left and not right:
        return [None] * qty
    dirs = ["좌"] * left + ["우"] * right
    return (dirs + [None] * qty)[:qty]


def _unit_code(value):
    m = re.search(r"(?<!\d)(\d{3,4}[A-Za-z]*)(?!\d)", str(value or ""))
    return m.group(1) if m else None


def parse_unitns(path):
    """UNITNS 발주 엑셀 -> 고객명 블록별 주문."""
    wb = openpyxl.load_workbook(path, data_only=True)
    orders = []
    for ws in wb.worksheets:
        if "UNITNS" not in str(ws.cell(1, 1).value or "").upper():
            continue
        r = 1
        while r <= ws.max_row:
            if str(ws.cell(r, 1).value or "").strip() != "고객명":
                r += 1
                continue
            customer = str(ws.cell(r, 2).value or "").strip()
            address = str(ws.cell(r, 7).value or "").strip()
            phone = str(ws.cell(r + 1, 2).value or "").strip()
            delivery_note = str(ws.cell(r + 2, 2).value or "").strip()
            header = r + 3
            if str(ws.cell(header, 1).value or "").strip() != "NO":
                r += 1
                continue
            o = _shell("유앤")
            o["고객명"] = customer or None
            o["전체기재사항"] = "피스"
            o["_unit_accessories"] = []
            o["배송"].update({
                "방식": "택배", "주소": _short_addr(address),
                "수령인": customer or None, "연락처": phone or None,
                "선불착불": "선불", "전달사항": delivery_note or None,
            })
            rr = header + 1
            while rr <= ws.max_row:
                first = str(ws.cell(rr, 1).value or "").strip()
                if first == "고객명" or (rr > header + 1 and first == "NO"):
                    break
                code = _unit_code(ws.cell(rr, 3).value)
                w, h = ws.cell(rr, 4).value, ws.cell(rr, 5).value
                if code and isinstance(w, (int, float)) and isinstance(h, (int, float)):
                    raw_qty = ws.cell(rr, 7).value
                    qty = max(1, int(raw_qty)) if isinstance(raw_qty, (int, float)) else 1
                    dirs = _unit_directions(ws.cell(rr, 6).value, qty)
                    product = str(ws.cell(rr, 2).value or "")
                    item_note = str(ws.cell(rr, 8).value or "").strip()
                    row_text = " ".join(str(ws.cell(rr, c).value or "") for c in range(2, 9))
                    hm = re.search(r"(?:손잡이(?:길이)?|손|줄|봉|각도봉)\s*[:=]?\s*(\d{2,3})", row_text)
                    handle_len = int(hm.group(1)) if hm else None
                    screw = None
                    accessory_name = None
                    note_key = item_note.replace(" ", "")
                    if "노피스" in note_key:
                        nm = re.search(r"노피스[^0-9]*(2|3)", note_key)
                        n = nm.group(1) if nm else None
                        screw = f"노피스({n})" if n else "노피스"
                        accessory_name = f"노피스브라켓({n}EA)" if n in {"2", "3"} else "노피스"
                    elif "콘크리트" in item_note:
                        screw, accessory_name = "콘크리트", "피스(콘크리트)"
                    elif "석고날개" in note_key:
                        screw, accessory_name = "석고날개", "피스(석고날개)"
                    elif "석고앙카" in note_key or "석고용앙카" in note_key:
                        screw, accessory_name = "석고앙카", "피스(석고앙카)"
                    elif "석고" in item_note:
                        # '석고'만 적힌 경우에는 특정 앙카/날개를 임의 확정하지 않는다.
                        screw, accessory_name = "석고", "석고"
                    elif "직결" in item_note or "시공" in item_note:
                        screw, accessory_name = "직결", "피스(시공/직결)"
                    if screw:
                        screw_count = 2 if float(w) < 150 else (3 if float(w) < 200 else 4)
                        o["_unit_accessories"].append({
                            "품명": accessory_name, "세트": qty, "피스수": screw_count,
                        })
                    type_ = "L자" if re.search(r"L\s*(?:자|형|타입|18|21)", product, re.I) else "C자"
                    for direction in dirs:
                        it = _blank_item()
                        it.update({"품목코드": code, "색상원문": str(ws.cell(rr, 3).value),
                                   "타입": type_, "종류": "투코드",
                                   "가로": float(w), "세로": float(h),
                                   "손잡이방향": direction,
                                   "손잡이길이": handle_len,
                                   "기재사항": None,
                                   "_수동특이": screw,
                                   "원문": " ".join(str(ws.cell(rr, c).value or "")
                                                      for c in range(2, 9)).strip()})
                        o["items"].append(it)
                rr += 1
            if o["items"]:
                orders.append(o)
            r = max(rr, r + 1)
    wb.close()
    if not orders:
        raise ValueError("유앤아이티엔에스 발주서에서 유효한 주문을 찾지 못했습니다")
    return orders


# ─────────────────────────────────────────────
# 보노 · 미래가공
# ─────────────────────────────────────────────
def _head_map(ws, required):
    for r in range(1, min(ws.max_row, 12) + 1):
        vals = {str(ws.cell(r, c).value or "").replace("\n", "").strip(): c
                for c in range(1, ws.max_column + 1)}
        if all(k in vals for k in required):
            return r, vals
    return None, {}


def _product_parts(product, color, operation):
    text = " ".join(str(x or "") for x in (product, color, operation))
    code = _unit_code(color) or _unit_code(text)
    type_ = "L자" if re.search(r"L\s*(?:자|형|타입|18|21)|알루미늄L", text, re.I) else "C자"
    if "셔터" in text:
        kind = "셔터"
    elif "원코드" in text:
        kind = "원코드"
    else:
        kind = "투코드"
    return code, type_, kind


def _direction(value):
    s = str(value or "").strip()
    if s.startswith(("좌", "왼")):
        return "좌"
    if s.startswith(("우", "오른")):
        return "우"
    return None


def _freight_info(value):
    """'안산초지점/착불/받는사람:보노(010...)'을 배송 정보로 분리."""
    s = str(value or "").strip()
    if not s:
        return "화물", None, None
    if "택배" in s:
        method = "택배"
    elif "배달" in s:
        method = "배달"
    else:
        method = "화물"
    first = s.split("/")[0].strip().replace("점", "")
    m = re.search(r"받는\s*사람\s*:\s*([^/(]+)?\s*\(?\s*(01\d[-\d]+)", s)
    receiver = (m.group(1).strip() if m and m.group(1) else "")
    phone = m.group(2).strip() if m else ""
    address = "-".join(x for x in (first, receiver) if x)
    if phone:
        address = f"{address} {phone}".strip()
    pay = "선불" if "선불" in s else ("착불" if "착불" in s else None)
    return method, address or None, pay


def _bono_freight_info(value, explicit_address=None):
    """보노의 화물지점/받는사람을 장부 기재사항과 배송정보로 분리한다.

    받는사람이 보노면 기재사항1은 대신화물 지점명(끝의 '점'만 제거),
    보노가 아니면 실제 받는사람/상호를 사용한다.
    """
    s = str(value or "").strip()
    if "택배" in s:
        method = "택배"
    elif "배달" in s:
        method = "배달"
    else:
        method = "화물"
    parts = [x.strip() for x in s.split("/") if x.strip()]
    branch_raw = parts[0] if parts else ""
    branch = re.sub(r"점\s*$", "", branch_raw).strip()
    pay = "선불" if "선불" in s else ("착불" if "착불" in s else None)

    receiver = ""
    phone = ""
    m = re.search(r"받는\s*사람\s*:\s*(.*)", s)
    if m:
        tail = m.group(1).strip()
        pm = re.search(r"(01\d[-\d]{7,})", tail)
        if pm:
            phone = pm.group(1).strip()
            receiver = tail[:pm.start()].strip(" ()/-")
        else:
            receiver = tail.strip(" ()/-")
    receiver = re.sub(r"\s+", " ", receiver).strip()
    is_bono = "보노" in re.sub(r"\s+", "", receiver) if receiver else False
    note1 = branch if is_bono else (receiver or branch)

    address = str(explicit_address or "").strip()
    if not address:
        # 구형 보노 파일에는 별도 도로명 주소 열이 없어서 화물지점-수령인 형태가 주소 역할을 했다.
        address = "-".join(x for x in (branch, receiver) if x)
    return method, address or None, pay, receiver or None, phone or None, branch or None, note1 or None, is_bono


def _sender_text(value, default="보노 010-2478-9290"):
    s = str(value or "").strip()
    if not s:
        return default
    s = re.sub(r"^발신\s*:\s*", "", s)
    s = s.replace("(", " ").replace(")", "")
    return re.sub(r"\s+", " ", s).strip()


def _finish_excel_order(order):
    if not order:
        return None
    items = order.get("items") or []
    if len(items) > 1:
        for it in items:
            if it.get("종류") == "원코드":
                it["연창"] = True
    return order if items else None


def parse_bono(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    hr, col = _head_map(ws, ("오더명", "거래처(화물지점)", "품명", "원단명"))
    if not col:
        wb.close()
        raise ValueError("보노 발주서 헤더를 찾지 못했습니다")

    # 새 양식에 주소 열이 추가돼도 헤더명에 '주소'가 포함되면 그대로 사용한다.
    address_col = next((c for h, c in col.items() if "주소" in str(h)), None)
    orders, cur = [], None
    for r in range(hr + 1, ws.max_row + 1):
        product = ws.cell(r, col["품명"]).value
        color = ws.cell(r, col["원단명"]).value
        size_col = col.get("발주사이즈")
        w = ws.cell(r, size_col).value if size_col else None
        h = ws.cell(r, size_col + 1).value if size_col else None
        if not (product and isinstance(w, (int, float)) and isinstance(h, (int, float))):
            continue
        order_no = str(ws.cell(r, col["오더명"]).value or "").replace("\n", "").strip()
        if order_no:
            done = _finish_excel_order(cur)
            if done:
                orders.append(done)
            cur = _shell("보노")
            cur["주문번호"] = order_no
            d = ws.cell(r, col["출고일"]).value if "출고일" in col else None
            if hasattr(d, "date"):
                cur["_ship_date"] = d.date().isoformat()
            explicit_address = ws.cell(r, address_col).value if address_col else None
            (method, address, pay, receiver, phone, branch, note1, is_bono) = \
                _bono_freight_info(ws.cell(r, col["거래처(화물지점)"]).value, explicit_address)
            cur["_bono_note1"] = note1
            cur["_bono_branch"] = branch
            cur["_bono_receiver_is_bono"] = bool(is_bono)
            cur["배송"].update({"방식": method, "주소": address, "화물지점": branch,
                                "수령인": receiver, "연락처": phone,
                                "선불착불": pay, "발신": "보노 010-2478-9290"})
        if cur is None:
            continue
        operation = ws.cell(r, col.get("작동방식필증", 0)).value \
            if col.get("작동방식필증") else None
        code, type_, kind = _product_parts(product, color, operation)
        it = _blank_item()
        it.update({"품목코드": code, "색상원문": str(color or ""),
                   "타입": type_, "종류": kind, "가로": float(w),
                   "세로": float(h),
                   "손잡이방향": _direction(ws.cell(r, col.get("손잡이방향", 0)).value
                                             if col.get("손잡이방향") else None),
                   "손잡이길이": ws.cell(r, col.get("길이", 0)).value
                                  if col.get("길이") else None,
                   "설치장소": str(ws.cell(r, col.get("시공위치", 0)).value or "").strip() or None,
                   "원문": " ".join(str(ws.cell(r, c).value or "")
                                      for c in range(1, ws.max_column + 1)).strip()})
        cur["items"].append(it)
    done = _finish_excel_order(cur)
    if done:
        orders.append(done)
    wb.close()
    if not orders:
        raise ValueError("보노 발주서에서 유효한 주문을 찾지 못했습니다")
    return orders

def parse_future(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    hr, col = _head_map(ws, ("고객명", "태영품명", "색상", "작동방식"))
    if not col:
        wb.close()
        raise ValueError("미래가공 발주서 헤더를 찾지 못했습니다")
    cur = None
    for r in range(hr + 1, ws.max_row + 1):
        product = ws.cell(r, col["태영품명"]).value
        w = ws.cell(r, col.get("가로", 0)).value if col.get("가로") else None
        h = ws.cell(r, col.get("세로", 0)).value if col.get("세로") else None
        if product and isinstance(w, (int, float)) and isinstance(h, (int, float)):
            order_no = str(ws.cell(r, col["고객명"]).value or "").strip()
            if cur is None:
                cur = _shell("미래가공")
                cur["주문번호"] = order_no.replace("\n", "")
                d = ws.cell(r, col.get("출고일", 0)).value if col.get("출고일") else None
                if hasattr(d, "date"):
                    cur["_ship_date"] = d.date().isoformat()
            color = ws.cell(r, col["색상"]).value
            operation = ws.cell(r, col["작동방식"]).value
            code, type_, kind = _product_parts(product, color, operation)
            it = _blank_item()
            it.update({"품목코드": code, "색상원문": str(color or ""),
                       "타입": type_, "종류": kind, "가로": float(w), "세로": float(h),
                       "손잡이방향": _direction(ws.cell(r, col.get("방향", 0)).value
                                                 if col.get("방향") else None),
                       "손잡이길이": ws.cell(r, col.get("줄길이", 0)).value
                                      if col.get("줄길이") else None,
                       "설치장소": str(ws.cell(r, col.get("비고", 0)).value or "").strip() or None,
                       "원문": " ".join(str(ws.cell(r, c).value or "")
                                          for c in range(1, ws.max_column + 1)).strip()})
            cur["items"].append(it)
        elif cur and "발신" in str(ws.cell(r, col["고객명"]).value or ""):
            sender = _sender_text(ws.cell(r, col["고객명"]).value)
            method, address, pay = _freight_info(ws.cell(r, col["태영품명"]).value)
            cur["배송"].update({"방식": method, "주소": address,
                                "선불착불": pay, "발신": sender})
    wb.close()
    done = _finish_excel_order(cur)
    if not done:
        raise ValueError("미래가공 발주서에서 유효한 주문을 찾지 못했습니다")
    return [done]


# ─────────────────────────────────────────────
def detect_excel_client(path):
    """파일명이 아닌 실제 시트/헤더로 DI·휴안·유앤을 구분한다."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if any("UNITNS" in str(ws.cell(1, 1).value or "").upper()
               for ws in wb.worksheets):
            return "유앤"
        for ws in wb.worksheets:
            _, headers = _head_map(ws, ("오더명", "거래처(화물지점)", "품명", "원단명"))
            if headers:
                return "보노"
            _, headers = _head_map(ws, ("고객명", "태영품명", "색상", "작동방식"))
            if headers:
                return "미래가공"
        # DI는 '사이즈' + '원단/색상' 헤더가 있다.
        for ws in wb.worksheets:
            _, col = _di_header(ws)
            if col:
                return "DI"

        # 휴안은 고정 열 구조: C=수령인, G=상품명, I=규격.
        ws = wb.worksheets[0]
        for r in range(1, min(ws.max_row, 15) + 1):
            vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, 13)]
            joined = " ".join(vals)
            if "상품" in joined and any(x in joined for x in ("사이즈", "규격", "수령인")):
                return "휴안"
    finally:
        wb.close()
    raise ValueError("엑셀 발주서의 거래처를 자동 판별하지 못했습니다. 거래처를 직접 선택해 다시 시도해 주세요.")


def parse_excel(path, client=None):
    client = client or detect_excel_client(path)
    if client == "휴안":
        return parse_huan(path)
    if client == "DI":
        return [parse_di(path)]
    if client == "유앤":
        return parse_unitns(path)
    if client == "보노":
        return parse_bono(path)
    if client == "미래가공":
        return parse_future(path)
    raise ValueError(f"엑셀 파서 없음: {client}")


if __name__ == "__main__":
    import json, sys
    p = sys.argv[1]
    c = "DI" if "DI" in p.upper() else "휴안"
    for o in parse_excel(p, c):
        print(json.dumps(o, ensure_ascii=False, indent=2, default=str))
