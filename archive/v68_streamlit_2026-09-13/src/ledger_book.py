"""장부 / 작업지시서 엑셀 쓰기(서식, 부분 색상, 병합, 페이지·DI 시트 분할).

2026-09-13 뼈대 정리: v68 output.py 에서 코드 변경 없이 분리.
"""
import re
from datetime import date
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from ledger_rows import sort_di_items


# ── 서식 상수 (원본 장부.xls 에서 추출) ──
FONT_DATA = "THE행복열매"
FONT_UI = "맑은 고딕"
ROW_H = 20.1
LEDGER_ROWS = 35              # 3행 ~ 37행
SHEET_ROWS = 16               # 작업지시서 블록당 데이터 행 수
DI_TARGET_ROWS = 35           # DI: 한 장의 기본 목표
DI_MAX_ROWS = 40              # DI: 같은 수령인을 유지하는 최대 행 수
DI_SHEET_GROUPS = (
    ("C 원코드", ("C자", "원코드")),
    ("C 투코드", ("C자", "투코드")),
    ("C 셔터", ("C자", "셔터")),
    ("L 원코드", ("L자", "원코드")),
    ("L 투코드", ("L자", "투코드")),
)
COL_W = [3.5, 15.1, 16.6, 6.6, 2.6, 6.6, 6.1, 6.0, 6.0, 6.1, 6.1, 8.4]
#        A         B(상호)   C(색상)  D(가로)  E(X)      F(세로)
#        G(수량)   H(모형1)  I(모형2) J(기재)  K(기재2)  L(출고일)
ALIGN = ["center", "right", "left", "right", "center", "left",
         "center", "left", "right", "center", "center", "center"]
DETAIL_COL_W = [3.5, 14.125, 16.625, 6.625, 2.625, 6.625, 4.875,
                5.125, 4.875, 5.625, 5.625, 5.625, 8.375]
# 세부 장부: A / 상호 / 색상 / 가로 / X / 세로 / 수량 / 방향 / 길이 /
#           특이 / 기재사항1 / 기재사항2 / 출고일
DETAIL_ALIGN = ["center", "right", "left", "right", "center", "left",
                "left", "center", "center", "right", "center", "center",
                "center"]
# 그룹 안쪽에는 세로선이 없다: 규격(D,E,F) / 모형(H,I) / 기재사항(J,K)
NO_LEFT = {4, 5, 8, 10}      # 0-based col index
NO_RIGHT = {3, 4, 7, 9}
HAIR = Side(style="hair")
THIN = Side(style="thin")
GREEN, RED, BLUE = "FF008000", "FFFF0000", "FF0000FF"


def _al(i):
    h = ALIGN[i]
    return Alignment(horizontal=(None if h == "general" else h),
                     vertical="center", shrink_to_fit=True)


def _border(r, c, last_row, top=None):
    left = None if c in NO_LEFT else (THIN if c == 0 else HAIR)
    right = None if c in NO_RIGHT else (THIN if c == 11 else HAIR)
    return Border(top=(top if top is not None else (THIN if r == 2 else HAIR)),
                  bottom=HAIR if r < last_row else THIN,
                  left=left, right=right)


def _tb(text, color=None, sz=14):
    """글꼴을 명시한 텍스트 블록 (기본 글꼴로 떨어지는 것을 방지)"""
    return TextBlock(InlineFont(rFont=FONT_DATA, sz=sz, color=color), text)


def _rich(text, word, color):
    """text 안의 word 부분만 색을 넣은 리치텍스트"""
    if not text or word not in text:
        return text
    i = text.index(word)
    parts = []
    if i:
        parts.append(_tb(text[:i]))
    parts.append(_tb(word, color))
    tail = text[i + len(word):]
    if tail:
        parts.append(_tb(tail))
    return CellRichText(*parts)


def _client_cell(client, mark):
    """상호: 코드는 14pt, 내부표시 (K) 는 11pt"""
    if not client:
        return mark or None
    if not mark:
        return client
    return CellRichText(_tb(client), _tb(f"        {mark}", sz=11))


def _color_cell(text):
    """색상 열 부분색상.

    - 원코드: `원코드` 글자만 초록
    - 셔터: `셔터` 글자만 빨강
    - P/FP: 품명 전체가 아니라 `029FP`, `102P` 같은 실제 색상코드만 파랑

    예: `B 원코드 029FP` -> `원코드` 초록 + `029FP` 파랑.
    """
    if not text:
        return text
    raw = str(text)
    token_re = re.compile(r"원코드|셔터|\d{1,4}(?:FP|P)(?![A-Za-z0-9])", re.I)
    matches = list(token_re.finditer(raw))
    if not matches:
        return raw
    parts = []
    pos = 0
    for m in matches:
        if m.start() > pos:
            parts.append(_tb(raw[pos:m.start()]))
        token = m.group(0)
        if re.fullmatch(r"\d{1,4}(?:FP|P)", token, re.I):
            color = BLUE
        elif token == "원코드":
            color = GREEN
        else:
            color = RED
        parts.append(_tb(token, color))
        pos = m.end()
    if pos < len(raw):
        parts.append(_tb(raw[pos:]))
    return CellRichText(*parts)


def _note_cell(text):
    """기재사항: '틀안' 파랑"""
    return _rich(text, "틀안", BLUE) if text and "틀안" in text else text


def _special_cell(text):
    """특이 칸: '틀안'은 파랑, 연창 #는 일반 색상."""
    return _rich(text, "틀안", BLUE) if text and "틀안" in text else text


# ─────────────────────────────────────────────
# 2. 장부 워크북
# ─────────────────────────────────────────────
HEAD = {1: "상      호", 2: "색       상", 3: "규     격", 6: "수량",
        7: "모   형", 9: "기재사항", 11: "출고일"}
DETAIL_HEAD = {1: "상      호", 2: "색       상", 3: "규     격", 6: "수량",
               7: "방향", 8: "길이", 9: "특이", 10: "기재사항", 12: "출고일"}


def _init_sheet(ws, total_rows, detailed=False):
    widths = DETAIL_COL_W if detailed else COL_W
    for c, width in enumerate(widths):
        ws.column_dimensions[get_column_letter(c + 1)].width = width
    for r in range(1, total_rows + 1):
        ws.row_dimensions[r].height = ROW_H


def _block(ws, top, nrows, header_date=None, detailed=False):
    """top 행에 날짜, top+1 에 헤더, 그 아래 nrows 개 데이터 영역을 만든다.
       데이터 시작 행을 반환."""
    last_col = 13 if detailed else 12
    head = DETAIL_HEAD if detailed else HEAD
    aligns = DETAIL_ALIGN if detailed else ALIGN
    no_left = {4, 5, 11} if detailed else NO_LEFT
    no_right = {3, 4, 10} if detailed else NO_RIGHT
    ws.merge_cells(start_row=top, end_row=top, start_column=8, end_column=last_col)
    d = header_date or date.today()
    c = ws.cell(top, 8, f"{d.month} 월   {d.day} 일        번째")
    c.font = Font(name=FONT_UI, size=11, bold=True)
    c.alignment = Alignment(horizontal="right", vertical="center")

    hr = top + 1
    for i in range(last_col):
        cell = ws.cell(hr, i + 1, head.get(i))
        cell.font = Font(name=FONT_UI, size=8 if i == last_col - 1 else 11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=THIN, bottom=HAIR,
                             left=None if i in no_left else (THIN if i == 0 else HAIR),
                             right=None if i in no_right else (THIN if i == last_col - 1 else HAIR))
    merges = ((4, 6), (11, 12)) if detailed else ((4, 6), (8, 9), (10, 11))
    for a, b in merges:
        ws.merge_cells(start_row=hr, end_row=hr, start_column=a, end_column=b)

    first, last = hr + 1, hr + nrows
    for r in range(first, last + 1):
        for i in range(last_col):
            cell = ws.cell(r, i + 1)
            cell.font = Font(name=FONT_UI if i == 4 else FONT_DATA,
                             size=11 if i == 4 else (18 if i == 0 else 14))
            h_align = aligns[i]
            cell.alignment = Alignment(horizontal=h_align, vertical="center",
                                       shrink_to_fit=True)
            left = None if i in no_left else (THIN if i == 0 else HAIR)
            right = None if i in no_right else (THIN if i == last_col - 1 else HAIR)
            # 샘플 장부 규칙: 헤더 아래 첫 데이터 행도 내부선(hair)이다.
            cell.border = Border(top=HAIR,
                                 bottom=HAIR, left=left, right=right)
            if r == last:
                cell.border = Border(top=HAIR, bottom=THIN,
                                     left=cell.border.left, right=cell.border.right)
        ws.cell(r, 5, "X")
    return first


def _new_sheet(wb, title, nrows, header_date=None, detailed=False):
    ws = wb.create_sheet(title) if wb.sheetnames != ["Sheet"] else wb.active
    ws.title = title
    _init_sheet(ws, nrows + 2, detailed)
    _block(ws, 1, nrows, header_date, detailed)
    return ws


def _write_rows(ws, rows, start=3, center_head=False, detailed=False):
    """rows 를 start 행부터 기록. 모형_2·기재사항은 조건에 맞으면 세로 병합"""
    r = start
    span = []
    previous_order_client = None
    for row in rows:
        if row.get("_특수") is not None:
            if row.get("_특수") == "부속":
                ws.cell(r, 3, row.get("문구"))
                ws.cell(r, 3).alignment = Alignment(horizontal="center",
                                                     vertical="center",
                                                     shrink_to_fit=True)
                ws.cell(r, 7, row.get("수량"))
                ws.cell(r, 7).alignment = Alignment(horizontal="center",
                                                     vertical="center",
                                                     shrink_to_fit=True)
                r += 1
                continue
            ws.cell(r, 3, row["_특수"])
            ws.cell(r, 3).alignment = Alignment(horizontal="right", vertical="center")
            ws.cell(r, 4, row.get("문구"))
            ws.cell(r, 4).alignment = Alignment(horizontal="left", vertical="center",
                                                shrink_to_fit=True)
            if row.get("선불"):
                pay_col = 12 if detailed else 11
                ws.merge_cells(start_row=r, end_row=r, start_column=4,
                               end_column=pay_col - 1)
                ws.cell(r, pay_col, row["선불"]).font = Font(
                    name=FONT_DATA, size=14, color=RED)
                ws.cell(r, pay_col).alignment = Alignment(horizontal="center",
                                                          vertical="center")
            else:
                ws.merge_cells(start_row=r, end_row=r, start_column=4,
                               end_column=12 if detailed else 11)
        else:
            client_text = row["거래처"]
            mark_text = row["내부표시"]
            # 같은 업체 주문이 연속되면 첫 주문에만 업체명을 쓰고
            # 다음 주문부터는 내부표시만 남긴다. 내부표시가 없는 업체(DI 등)는
            # 공통 반복 표시 (K)를 사용한다. SP/JL처럼 주문별 표시가 있는 경우는 보존한다.
            if client_text:
                same_client = previous_order_client == client_text
                display_client = None if same_client else client_text
                if same_client and not mark_text:
                    mark_text = "(K)"
                previous_order_client = client_text
            else:
                display_client = None
            c2 = ws.cell(r, 2, _client_cell(display_client, mark_text))
            if center_head:
                c2.alignment = Alignment(horizontal="center",
                                         vertical="center", shrink_to_fit=True)
            c3 = ws.cell(r, 3, _color_cell(row["색상"]))
            if center_head:
                c3.alignment = Alignment(horizontal="center",
                                         vertical="center", shrink_to_fit=True)
            ws.cell(r, 4, row["가로"])
            ws.cell(r, 6, row["세로"])
            # 같은 세로의 따옴표 표시는 앞 공백 2칸을 유지하고 셀 중앙에 정렬한다.
            if str(row.get("세로") or "").strip() == '"':
                ws.cell(r, 6).alignment = Alignment(horizontal="center", vertical="center",
                                                    shrink_to_fit=True)
            ws.cell(r, 7, row["수량"])
            ws.cell(r, 8, row["모형1"])
            ws.cell(r, 9, row["모형2"])
            if detailed:
                ws.cell(r, 10, _special_cell(row.get("특이")))
                # 연창 표시 #을 포함한 특이 칸은 가운데 정렬한다.
                ws.cell(r, 10).alignment = Alignment(horizontal="center",
                                                      vertical="center",
                                                      shrink_to_fit=True)
                ws.cell(r, 11, _note_cell(row["기재사항"]))
                ws.cell(r, 12, row["기재사항2"])
                note_col, note2_col, date_col = 11, 12, 13
            else:
                ws.cell(r, 10, _note_cell(row["기재사항"]))
                ws.cell(r, 11, row["기재사항2"])
                note_col, note2_col, date_col = 10, 11, 12
            if (row["기재사항"] and not row["기재사항2"]
                    and not row.get("_DI수령인") and not row.get("_JL고객")
                    and not row.get("_JO주문번호")):
                span.append(r)          # 방 이름이 없으면 J:K 를 합친다
            ws.cell(r, date_col, row["출고일"])
        r += 1

    # DI 수령인은 색상이 달라도 같은 사람의 연속 행 전체를 하나로
    # 병합한다. 세부 양식에서는 기재사항 두 칸(K:L)을 직사각형으로 묶는다.
    i = 0
    while i < len(rows):
        recipient = rows[i].get("_DI수령인")
        if not recipient or rows[i].get("_특수") is not None:
            i += 1
            continue
        j = i + 1
        first_note = str(rows[i].get("기재사항") or "").strip()
        while (j < len(rows) and rows[j].get("_특수") is None
               and rows[j].get("_DI수령인") == recipient
               and str(rows[j].get("기재사항") or "").strip() == first_note):
            j += 1
        ws.merge_cells(start_row=start + i, end_row=start + j - 1,
                       start_column=11 if detailed else 10,
                       end_column=12 if detailed else 11)
        i = j

    # JL은 같은 주문/고객의 여러 제품행에서 `피스/이름` 같은 공통
    # 기재사항을 한 번만 보이게 한다. 기재사항2(설치장소)가 전부 비어 있으면
    # K:L 전체를 직사각형으로, 설치장소가 있으면 K열만 세로 병합한다.
    i = 0
    while i < len(rows):
        row0 = rows[i]
        customer = row0.get("_JL고객")
        order_id = row0.get("_oi")
        if not customer or row0.get("_특수") is not None:
            i += 1
            continue
        j = i + 1
        while (j < len(rows) and rows[j].get("_특수") is None
               and rows[j].get("_JL고객") == customer
               and rows[j].get("_oi") == order_id):
            j += 1
        if j - i > 1:
            # 병합할 공통 기재사항은 첫 행의 값을 기준으로 하고, 하위 행에서
            # 고객명/피스만 빠진 형태도 같은 공통값으로 간주한다.
            first_note = str(row0.get("기재사항") or "").strip()
            if first_note:
                all_note2_empty = all(not str(rows[k].get("기재사항2") or "").strip()
                                      for k in range(i, j))
                if all_note2_empty:
                    ws.merge_cells(start_row=start + i, end_row=start + j - 1,
                                   start_column=11 if detailed else 10,
                                   end_column=12 if detailed else 11)
                else:
                    ws.merge_cells(start_row=start + i, end_row=start + j - 1,
                                   start_column=11 if detailed else 10,
                                   end_column=11 if detailed else 10)
                ws.cell(start + i, 11 if detailed else 10).alignment = Alignment(
                    horizontal="center", vertical="center", shrink_to_fit=True)
        i = j

    # JO 주문번호는 기재사항1 열에만 두고, 같은 주문의 여러 행이면 세로 병합한다.
    i = 0
    while i < len(rows):
        row0 = rows[i]
        order_no = row0.get("_JO주문번호")
        order_id = row0.get("_oi")
        if not order_no or row0.get("_특수") is not None:
            i += 1
            continue
        j = i + 1
        while (j < len(rows) and rows[j].get("_특수") is None
               and rows[j].get("_oi") == order_id
               and rows[j].get("_JO주문번호") == order_no):
            j += 1
        # JO는 3창 이상일 때만 기재사항1을 주문번호 전용으로 세로 병합한다.
        # 1~2창은 각 행의 주문번호+개별 내용을 보존해야 하므로 병합하지 않는다.
        if j - i > 1 and row0.get("_JO기재2사용"):
            col = 11 if detailed else 10
            ws.merge_cells(start_row=start + i, end_row=start + j - 1,
                           start_column=col, end_column=col)
            ws.cell(start + i, col).alignment = Alignment(
                horizontal="center", vertical="center", shrink_to_fit=True)
        i = j

    # 전체 공통: 같은 주문 안에서 기재사항 값이 연속해서 완전히 같으면
    # 색상이 달라도 한 셀로 세로 병합한다. DI/JL/JO는 위 전용 병합 규칙을 우선한다.
    for note_col, note_key in ((11 if detailed else 10, "기재사항"),
                               (12 if detailed else 11, "기재사항2")):
        i = 0
        while i < len(rows):
            row0 = rows[i]
            value = str(row0.get(note_key) or "").strip()
            order_id = row0.get("_oi")
            if (row0.get("_특수") is not None or not value or order_id is None
                    or row0.get("_DI수령인") or row0.get("_JL고객")
                    or row0.get("_JO주문번호")):
                i += 1
                continue
            j = i + 1
            while (j < len(rows) and rows[j].get("_특수") is None
                   and rows[j].get("_oi") == order_id
                   and str(rows[j].get(note_key) or "").strip() == value
                   and not rows[j].get("_DI수령인")
                   and not rows[j].get("_JL고객")
                   and not rows[j].get("_JO주문번호")):
                j += 1
            if j - i > 1:
                try:
                    ws.merge_cells(start_row=start + i, end_row=start + j - 1,
                                   start_column=note_col, end_column=note_col)
                    ws.cell(start + i, note_col).alignment = Alignment(
                        horizontal="center", vertical="center", shrink_to_fit=True)
                except ValueError:
                    pass
            i = j

    # DI의 같은 수령인·같은 품목코드·같은 유형이 연속되면
    # 색상(C열) 칸을 세로로 병합한다. P/FP는 별도 코드이므로 서로 합치지 않는다.
    i = 0
    while i < len(rows):
        row0 = rows[i]
        recipient = row0.get("_DI수령인")
        code = row0.get("_DI품목코드")
        type_group = row0.get("_DI유형")
        if not recipient or not code or row0.get("_특수") is not None:
            i += 1
            continue
        j = i + 1
        while (j < len(rows) and rows[j].get("_특수") is None
               and rows[j].get("_DI수령인") == recipient
               and rows[j].get("_DI품목코드") == code
               and rows[j].get("_DI유형") == type_group):
            j += 1
        if j - i > 1:
            ws.merge_cells(start_row=start + i, end_row=start + j - 1,
                           start_column=3, end_column=3)
            ws.cell(start + i, 3).alignment = Alignment(
                horizontal="center" if center_head else "left",
                vertical="center", shrink_to_fit=True)
        i = j

    for rr in span:                 # 기재사항 가로 병합 (J:K)
        ws.merge_cells(start_row=rr, end_row=rr,
                       start_column=11 if detailed else 10,
                       end_column=12 if detailed else 11)

    # 출고일은 같은 주문(휴안·유앤은 한 사람)의 제품·포장·배송·주소·
    # 연락처·전달사항 전체를 세로로 병합한다.
    # 페이지가 나뉘면 각 페이지 안의 구간만 병합된다.
    dated = [(start + i, row) for i, row in enumerate(rows)
             if row.get("_oi") is not None]
    i = 0
    while i < len(dated):
        r0, row0 = dated[i]
        order_id = row0.get("_oi")
        j = i + 1
        while j < len(dated) and dated[j][1].get("_oi") == order_id:
            j += 1
        if j - i > 1 and row0.get("출고일"):
            ws.merge_cells(start_row=r0, end_row=dated[j - 1][0],
                           start_column=13 if detailed else 12,
                           end_column=13 if detailed else 12)
            ws.cell(r0, 13 if detailed else 12).alignment = Alignment(horizontal="center",
                                                   vertical="center")
        i = j

    # 세로 병합: 색상이 이어지는 구간에서 같은 값이 반복되면 묶는다
    prod = [(start + i, row) for i, row in enumerate(rows)
            if row.get("_특수") is None]
    spanset = set(span)
    merge_specs = ((9, "모형2"), (11, "기재사항")) if detailed \
                  else ((9, "모형2"), (10, "기재사항"))
    for col, key in merge_specs:
        i = 0
        while i < len(prod):
            r0, row0 = prod[i]
            if key == "기재사항" and (row0.get("_DI수령인") or row0.get("_JL고객")
                                          or row0.get("_JO주문번호")):
                i += 1
                continue
            j = i + 1
            while (j < len(prod) and prod[j][1].get("_같은색")
                   and prod[j][1].get(key) in (None, "", row0.get(key))):
                j += 1
            if col == (11 if detailed else 10) and any(rr in spanset for rr in
                                 (p[0] for p in prod[i:j])):
                i = j                # 가로 병합한 행은 세로 병합에서 제외
                continue
            if j - i > 1 and row0.get(key):
                ws.merge_cells(start_row=r0, end_row=prod[j - 1][0],
                               start_column=col, end_column=col)
            i = j
    return r


def build_ledger(all_rows, path, header_date=None):
    wb = openpyxl.Workbook()
    is_di = bool(all_rows) and all(r.get("_거래처코드") == "DI" for r in all_rows)
    if is_di:
        first_sheet = True
        for group_name, group_key in DI_SHEET_GROUPS:
            group_rows = [r for r in all_rows if r.get("_DI유형") == group_key]
            for page_no, chunk in enumerate(_di_paginate(group_rows) or [[]], 1):
                title = group_name if page_no == 1 else f"{group_name} ({page_no})"
                page_rows = max(DI_MAX_ROWS, len(chunk))
                if first_sheet:
                    ws = wb.active
                    ws.title = title
                    _init_sheet(ws, page_rows + 2, detailed=True)
                    _block(ws, 1, page_rows, header_date, detailed=True)
                    first_sheet = False
                else:
                    ws = _new_sheet(wb, title, page_rows, header_date,
                                    detailed=True)
                _write_rows(ws, chunk, detailed=True)
    else:
        pages = _paginate_order_rows(all_rows, LEDGER_ROWS)
        for n, chunk in enumerate(pages, 1):
            ws = _new_sheet(wb, "장부" if n == 1 else f"장부 ({n})",
                            LEDGER_ROWS, header_date, detailed=True)
            _write_rows(ws, chunk, detailed=True)
    wb.save(path)
    return path


def _paginate_order_rows(rows, limit=LEDGER_ROWS):
    """같은 주문/사람을 가능한 한 자르지 않고 장부 페이지를 나눈다."""
    if not rows:
        return [[]]
    blocks, i = [], 0
    while i < len(rows):
        order_id = rows[i].get("_oi")
        j = i + 1
        while j < len(rows) and rows[j].get("_oi") == order_id:
            j += 1
        blocks.append(rows[i:j])
        i = j

    pages, current = [], []
    for block in blocks:
        if current and len(current) + len(block) > limit:
            pages.append(current)
            current = []
        if len(block) > limit:
            if current:
                pages.append(current)
                current = []
            pages.extend(block[i:i + limit] for i in range(0, len(block), limit))
        else:
            current.extend(block)
    if current:
        pages.append(current)
    return pages or [[]]


# ─────────────────────────────────────────────
# 3. 장부 -> 작업지시서
# ─────────────────────────────────────────────
INNER = re.compile(r"\s*\((?:[A-Z]|\d+)\)\s*$")


def to_worksheet_rows(rows):
    """내부표시 제거, 색상 앞 'B ' 제거, 특수 행 제외.

    작업지시서도 세부 양식을 사용하므로 #·틀안은 길이/기재사항에
    중복하지 않고 특이 칸에만 유지한다.
    """
    out = []
    for row in rows:
        r = dict(row)
        if r.get("_특수") is None:
            r["내부표시"] = None                  # (K), (N), (1) 등 삭제
            if r.get("색상"):
                r["색상"] = re.sub(r"^\s*[BHRC]\s+", "", r["색상"], flags=re.I).strip()
        # 포장비/부속/주소/전달사항도 장부와 작업지시서가 서로 어긋나지 않게 보존한다.
        out.append(r)
    return out


def _group_len(rows):
    """거래처가 채워진 지점을 기준으로 업체별 묶음 길이"""
    idx = [i for i, r in enumerate(rows) if r.get("거래처")] + [len(rows)]
    return [(idx[i], idx[i + 1] - idx[i]) for i in range(len(idx) - 1)]


def _di_paginate(rows):
    """DI 행을 35행 목표/최대 40행으로 나누되 수령인은 자르지 않는다."""
    blocks = []
    for row in rows:
        recipient = row.get("_DI수령인") or "__unknown"
        if blocks and blocks[-1][0] == recipient:
            blocks[-1][1].append(row)
        else:
            blocks.append((recipient, [row]))

    if not blocks:
        return []

    # 108행처럼 큰 수령인 묶음이 뒤쪽에 있어도 마지막에 2~3행짜리
    # 추가 페이지가 생기지 않도록, 필요 페이지 수를 먼저 정하고 수령인
    # 블록을 균등 배치한다. 페이지 안에서는 원래 색상 순서를 유지한다.
    total = sum(len(block) for _, block in blocks)
    page_count = max(1, (total + DI_MAX_ROWS - 1) // DI_MAX_ROWS)
    indexed = [(i, recipient, block) for i, (recipient, block) in enumerate(blocks)]
    desired = (total + page_count - 1) // page_count
    # 최윤동 26행처럼 큰 묶음은 전용 페이지 틀을 먼저 잡아 둔다.
    large = [x for x in indexed if len(x[2]) > desired / 2]
    reserved = [{"load": len(block), "blocks": [(i, block)]}
                for i, _, block in large]
    large_ids = {i for i, _, _ in large}
    regular_count = max(0, page_count - len(reserved))
    regular = [{"load": 0, "blocks": []} for _ in range(regular_count)]

    leftovers = []
    bi = 0
    for index, recipient, block in indexed:
        if index in large_ids:
            continue
        if bi < regular_count:
            cur = regular[bi]
            if (cur["load"] >= DI_TARGET_ROWS
                    or (cur["load"] + len(block) > desired
                        and cur["load"] >= DI_TARGET_ROWS)):
                bi += 1
            if bi < regular_count and regular[bi]["load"] + len(block) <= DI_MAX_ROWS:
                regular[bi]["blocks"].append((index, block))
                regular[bi]["load"] += len(block)
                continue
        leftovers.append((index, block))

    # 앞장을 채우고 남은 작은 묶음은 큰 수령인 페이지에 붙인다.
    bins = regular + reserved
    for index, block in leftovers:
        candidates = [b for b in bins if b["load"] + len(block) <= DI_MAX_ROWS]
        chosen = min(candidates or bins, key=lambda b: b["load"])
        chosen["blocks"].append((index, block))
        chosen["load"] += len(block)

    # 최초 정렬 순서가 빠른 페이지부터 나오게 하고, 페이지 내부도 유지한다.
    bins.sort(key=lambda b: min((i for i, _ in b["blocks"]), default=999999))
    pages = []
    for b in bins:
        page = []
        for _, block in sorted(b["blocks"], key=lambda x: x[0]):
            page.extend(block)
        pages.append(page)
    return pages


def sort_di_items_in_ledger_order(items, M):
    """DI 품목을 최종 장부의 시트·페이지·행 순서로 반환한다.

    단순 색상 정렬뿐 아니라 종류별 시트 분리와 35~40행 페이지 배치까지
    장부 생성과 똑같이 적용한다. 경영박사 EDI가 발주서 원본 순서가 아닌
    실제 장부에 보이는 순서를 따라야 할 때 사용한다.
    """
    sorted_items = sort_di_items(list(items), M)
    ordered, included = [], set()
    for _, group_key in DI_SHEET_GROUPS:
        group = []
        for item in sorted_items:
            if (item.get("타입"), item.get("종류")) != group_key:
                continue
            group.append({
                "_DI수령인": str(item.get("기재사항") or "").strip(),
                "_EDI품목": item,
            })
            included.add(id(item))
        for page in _di_paginate(group):
            ordered.extend(row["_EDI품목"] for row in page)

    # 검증 전의 미확정 유형이 있더라도 EDI에서 행이 사라지지는 않게 한다.
    ordered.extend(item for item in sorted_items if id(item) not in included)
    return ordered


def build_worksheet(rows, path, header_date=None):
    """업체 묶음을 자르지 않고 SHEET_ROWS 씩 분할.
       한 시트 = 제작용 블록 + 확인용 블록 (내용 동일, 각각 날짜·헤더 포함).
       변경된 장부와 동일하게 방향·길이·특이·기재사항을 분리한 13열
       세부 양식을 사용한다."""
    ws_rows = to_worksheet_rows(rows)
    is_di = bool(ws_rows) and all(r.get("_거래처코드") == "DI" for r in ws_rows)
    if is_di:
        wb = openpyxl.Workbook()
        first_sheet = True
        for group_name, group_key in DI_SHEET_GROUPS:
            group_rows = [r for r in ws_rows if r.get("_DI유형") == group_key]
            for page_no, chunk in enumerate(_di_paginate(group_rows) or [[]], 1):
                title = group_name if page_no == 1 else f"{group_name} ({page_no})"
                ws = wb.active if first_sheet else wb.create_sheet()
                first_sheet = False
                ws.title = title
                page_rows = max(DI_MAX_ROWS, len(chunk))
                _init_sheet(ws, page_rows + 2, detailed=True)
                first = _block(ws, 1, page_rows, header_date, detailed=True)
                _write_rows(ws, chunk, start=first, center_head=True,
                            detailed=True)
        wb.save(path)
        return path

    pages, cur = [], []
    for start, n in _group_len(ws_rows):
        block = ws_rows[start:start + n]
        if cur and len(cur) + n > SHEET_ROWS:
            pages.append(cur)
            cur = []
        cur.extend(block)
    if cur:
        pages.append(cur)

    wb = openpyxl.Workbook()
    for i, chunk in enumerate(pages or [[]], 1):
        ws = wb.create_sheet("현장" if i == 1 else f"현장 ({i})") \
             if wb.sheetnames != ["Sheet"] else wb.active
        ws.title = "현장" if i == 1 else f"현장 ({i})"
        total = (SHEET_ROWS + 2) * 2
        _init_sheet(ws, total, detailed=True)
        f1 = _block(ws, 1, SHEET_ROWS, header_date, detailed=True)
        _write_rows(ws, chunk, start=f1, center_head=True, detailed=True)
        f2 = _block(ws, SHEET_ROWS + 3, SHEET_ROWS, header_date,
                    detailed=True)
        _write_rows(ws, chunk, start=f2, center_head=True,
                    detailed=True)      # 확인용
    wb.save(path)
    return path
