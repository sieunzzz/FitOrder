"""
21건 정답셋으로 추출 정확도 측정

사용:
    python evaluate.py ../data/samples          # 전체
    python evaluate.py ../data/samples DU       # 한 거래처만

정답 파일은 발주서와 같은 이름의 .xlsx (장부 시트) 이며,
같은 폴더에 있어야 한다.  DU_01.png <-> DU_01.xlsx
"""
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

from extract import extract_order, ledger_color

IMG_EXT = {".png", ".jpg", ".jpeg"}
CACHE = Path("_cache")
CACHE.mkdir(exist_ok=True)


def read_truth(xlsx_path):
    """정답 장부 -> 행 리스트 (색상/가로/세로/수량/모형1/모형2/기재사항)"""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["장부"] if "장부" in wb.sheetnames else wb.worksheets[0]
    rows, prev_color = [], None
    for r in range(3, ws.max_row + 1):
        g = lambda c: ws.cell(r, c).value
        color, w, h = g(3), g(4), g(6)
        if w in (None, "") and color in (None, ""):
            continue
        if color and str(color).strip().startswith(("#", "☆")):
            continue                                  # 특수 문구 행 제외
        color = str(color).strip() if color else prev_color   # 공란 = 위와 동일
        prev_color = color
        rows.append({
            "색상": color,
            "가로": w,
            "세로": h,                                # '"' 는 그대로 둠
            "수량": (str(g(7)).strip() if g(7) else None),
            "모형1": (str(g(8)).strip() if g(8) else None),
            "모형2": (str(g(9)).strip() if g(9) else None),
            "기재사항": " ".join(str(g(c)).strip() for c in (10, 11) if g(c)),
        })
    return rows


def to_ledger_rows(res):
    """추출 결과 -> 장부 행 형태 (비교용)"""
    out = []
    for it in res.get("items", []):
        if it.get("예외품목"):
            continue
        hl = it.get("손잡이길이")
        out.append({
            "색상": ledger_color(it),
            "가로": it.get("가로"),
            "세로": it.get("세로"),
            "수량": it.get("수량"),
            "모형1": it.get("손잡이방향"),
            "모형2": (f"손{hl}" if hl else None),
            "기재사항": " ".join(x for x in (it.get("기재사항"),
                                            it.get("설치장소")) if x),
        })
    return out


def num_eq(a, b):
    try:
        return abs(float(a) - float(b)) < 0.01
    except (TypeError, ValueError):
        return str(a or "").strip() == str(b or "").strip()


def resolve_ditto(truth):
    """'"' 를 위 행 세로값으로 치환 (비교 편의)"""
    prev = None
    for t in truth:
        v = str(t["세로"]).strip() if t["세로"] is not None else ""
        if v in ('"', "''", "”"):
            t["세로"] = prev
        else:
            prev = t["세로"]
    return truth


FIELDS = ["색상", "가로", "세로", "수량", "모형1", "모형2", "기재사항"]


def main():
    root = Path(sys.argv[1] if len(sys.argv) > 1 else "../data/samples")
    only = sys.argv[2] if len(sys.argv) > 2 else None

    pairs = []
    for d in sorted(root.iterdir()):
        if not d.is_dir() or (only and d.name != only):
            continue
        for img in sorted(d.iterdir()):
            if img.suffix.lower() not in IMG_EXT:
                continue
            truth = img.with_suffix(".xlsx")
            if truth.exists():
                pairs.append((d.name, img, truth))
        for xl in sorted(d.glob("*.xlsx")):        # 엑셀 발주(DI·휴안)
            if "정답" in xl.name or xl.with_suffix(".png").exists():
                continue

    print(f"평가 대상 {len(pairs)}건\n")
    stat = defaultdict(lambda: [0, 0])
    row_stat = [0, 0]
    per_client = defaultdict(lambda: [0, 0])

    for client, img, tpath in pairs:
        cache = CACHE / f"{img.stem}.json"
        if cache.exists():
            res = json.loads(cache.read_text(encoding="utf-8"))
        else:
            res = extract_order([img], client_hint=client)
            cache.write_text(json.dumps(res, ensure_ascii=False, indent=2),
                             encoding="utf-8")

        truth = resolve_ditto(read_truth(tpath))
        got = to_ledger_rows(res)

        ok_row = len(truth) == len(got)
        marks = []
        for i, t in enumerate(truth):
            g = got[i] if i < len(got) else {}
            bad = []
            for f in FIELDS:
                stat[f][1] += 1
                if num_eq(t[f], g.get(f)):
                    stat[f][0] += 1
                else:
                    bad.append(f)
            row_stat[1] += 1
            if not bad:
                row_stat[0] += 1
            else:
                ok_row = False
                marks.append(f"{i + 1}행: {','.join(bad)}")

        per_client[client][1] += 1
        per_client[client][0] += 1 if ok_row else 0
        flag = "O" if ok_row else "X"
        print(f"[{flag}] {img.name:16} 정답{len(truth)}행 / 추출{len(got)}행"
              + ("  " + " | ".join(marks[:3]) if marks else ""))

    print("\n── 필드별 정확도 ──")
    for f in FIELDS:
        ok, n = stat[f]
        if n:
            print(f"  {f:8} {ok}/{n}  {ok * 100 // n}%")
    ok, n = row_stat
    if n:
        print(f"\n행 완전일치 {ok}/{n}  {ok * 100 // n}%")
    print("\n── 거래처별 ──")
    for c, (ok, n) in sorted(per_client.items()):
        print(f"  {c:8} {ok}/{n}")


if __name__ == "__main__":
    main()
