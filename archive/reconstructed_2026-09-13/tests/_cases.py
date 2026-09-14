"""특성(characterization) 테스트용 케이스 정의.

현재 코드가 실제로 내는 결과를 golden 파일로 고정해, 구조 정리(리팩터링) 중
출력이 의도치 않게 바뀌지 않았는지 확인한다.

- golden 파일은 "정답"이 아니라 "현재 동작"이다.
  규칙을 의도적으로 바꾼 경우에만 `python tests/_generate_golden.py`로 다시 만들고,
  바뀐 부분을 사용자에게 확인받는다.
- 개인정보가 아닌 가상의 값만 사용한다.
"""
from __future__ import annotations

import copy
import csv
import json
import sys
import tempfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))

from fitorder.desktop_workflow import LEDGER_COLUMNS, build_preview_rows  # noqa: E402
from fitorder.output import build_erp, build_ledger  # noqa: E402
from fitorder.rules import CLIENTS, note_fields  # noqa: E402

GOLDEN_PATH = Path(__file__).resolve().parent / "golden" / "characterization.json"
SHIP_DATE = date(2026, 9, 14)  # 월요일
ALL_CLIENTS = CLIENTS + ["미확인"]


def _full_order(client):
    """택배·선불·주소·수령인·전달사항과 3개 품목을 모두 가진 주문."""
    return {
        "거래처": client,
        "주문번호": "ORD-1",
        "고객명": "홍길동",
        "전체기재사항": "피스",
        "배송": {
            "방식": "택배", "선불착불": "선불", "주소": "서울특별시 테스트구 테스트로 1, 101호",
            "수령인": "김수령", "연락처": "010-0000-0000", "전달사항": "문앞",
        },
        "items": [
            {"품목코드": "102", "종류": "원코드", "가로": 150.0, "세로": 124, "창개수": 1,
             "손잡이방향": "좌", "손잡이길이": 120, "연창": True, "설치장소": "거실"},
            {"품목코드": "029FP", "종류": "투코드", "가로": 29.5, "세로": 371,
             "기재사항": "안방", "특이": "특"},
            {"품목코드": "B 원코드 520", "종류": "원코드", "가로": 80, "세로": 200, "창개수": 2},
        ],
    }


def _minimal_order(client):
    """배송정보·주문번호 없이 품목 1개."""
    return {"거래처": client, "고객명": "이고객", "items": [
        {"품목코드": "330", "종류": "셔터", "가로": "100", "세로": "150.5"}]}


def _cargo_no_address(client):
    """화물·주소 없음·2개 품목."""
    return {"거래처": client, "주문번호": "C-7", "배송": {"방식": "화물", "선불착불": "착불"}, "items": [
        {"품목코드": "102", "종류": "원코드", "가로": 90, "세로": 120, "설치장소": "작은방"},
        {"품목코드": "", "종류": "", "가로": None, "세로": 130},
    ]}


def _bono_branch_order(client):
    """보노 분기(받는사람=보노 → 지점)와 수동 포장 해제."""
    return {"거래처": client, "주문번호": "부천67327013", "고객명": "보노",
            "지점": "안산초지", "_manual_packing_enabled": False,
            "배송": {"방식": "택배", "수령인": "보노 안산", "주소": "경기도 안산시 1"},
            "items": [{"품목코드": "102", "종류": "원코드", "가로": 100, "세로": 100, "설치장소": "거실1-1"}]}


ORDER_VARIANTS = {
    "full": _full_order,
    "minimal": _minimal_order,
    "cargo_no_address": _cargo_no_address,
    "bono_branch": _bono_branch_order,
}


def _jsonable(value):
    return json.loads(json.dumps(value, ensure_ascii=False, default=str))


def _visible(rows):
    return [{k: r.get(k, "") for k in LEDGER_COLUMNS + ["_특수", "_detail_kind"]} for r in rows]


def snapshot_note_fields():
    orders = {
        "with_common": {"주문번호": "N-1", "고객명": "고객", "전체기재사항": "공통",
                        "배송": {"수령인": "받는분"}},
        "no_common": {"고객명": "고객", "배송": {}},
    }
    items = {"설치장소": {"설치장소": "거실"}, "기재사항": {"기재사항": "안방"}, "없음": {}}
    out = []
    for client in ALL_CLIENTS:
        for oname, base in orders.items():
            for iname, item in items.items():
                for total in (1, 2, 3):
                    order = dict(copy.deepcopy(base), 거래처=client)
                    n1, n2 = note_fields(order, item, 0, total)
                    out.append([client, oname, iname, total, n1, n2])
    return out


def snapshot_preview():
    out = {}
    for client in ALL_CLIENTS:
        for vname, make in ORDER_VARIANTS.items():
            rows, _ = build_preview_rows([make(client)], SHIP_DATE)
            out[f"{client}/{vname}"] = _visible(rows)
    return _jsonable(out)


def all_orders():
    return [make(client) for client in ALL_CLIENTS for make in ORDER_VARIANTS.values()]


def _read_ledger(path):
    import openpyxl
    ws = openpyxl.load_workbook(path).active
    values = [[c.value for c in row] for row in ws.iter_rows()]
    colored = []
    for row in ws.iter_rows():
        for c in row:
            rgb = getattr(c.font.color, "rgb", None) if c.font and c.font.color else None
            if isinstance(rgb, str) and rgb not in ("FF000000",):
                colored.append([c.coordinate, rgb])
    return {"values": values, "colored": colored}


def snapshot_ledger():
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "ledger.xlsx"
        build_ledger(all_orders(), SHIP_DATE, path)
        return _jsonable(_read_ledger(path))


def snapshot_erp():
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "erp.csv"
        build_erp(all_orders(), SHIP_DATE, path)
        with path.open(encoding="utf-8-sig", newline="") as f:
            return list(csv.reader(f))


SNAPSHOTS = {
    "note_fields": snapshot_note_fields,
    "preview": snapshot_preview,
    "ledger": snapshot_ledger,
    "erp": snapshot_erp,
}


def snapshot_all():
    return {name: fn() for name, fn in SNAPSHOTS.items()}
