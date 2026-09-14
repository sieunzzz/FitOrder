"""장부 / 작업지시서 / 경영박사 EDI 파일 생성.

모든 출력은 `pipeline.build_rows` 결과(사용자 수정값 반영 완료)만 사용한다.
출력마다 규칙을 다시 계산하지 않아야 장부·작업지시서·ERP가 서로 일관된다.
"""
from __future__ import annotations
from datetime import date
from pathlib import Path
import csv, re

from .pipeline import LEDGER_COLUMNS, build_rows

HEADERS = LEDGER_COLUMNS

def _style_sheet(ws):
    from openpyxl.styles import Alignment, Border, Font, Side
    thin = Side(style="thin", color="FF444444")
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name="맑은 고딕", size=10)
            c.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
            c.border = Border(left=thin,right=thin,top=thin,bottom=thin)
    ws.freeze_panes = "A2"

def build_ledger(orders, ship_date, path, rows=None):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "장부"
    ws.append(HEADERS)
    rows = build_rows(orders, ship_date) if rows is None else rows
    for row in rows:
        ws.append([row.get(h,"") for h in HEADERS])
    _style_sheet(ws)
    # 전체 셀색 대신 확실한 단일 토큰 셀만 색을 적용. 부분 글자색(rich text)은 출력 정리 단계에서 구현.
    for r in range(2, ws.max_row+1):
        txt = str(ws.cell(r,2).value or "")
        if re.fullmatch(r"\d{3,4}(?:FP|P)", txt, re.I):
            ws.cell(r,2).font = Font(name="맑은 고딕", size=10, color="FF0000FF")
    widths = [16,24,18,4,10,8,8,8,10,25,25,10]
    for i,w in enumerate(widths,1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    wb.save(path)
    return Path(path)

def build_worksheet(orders, ship_date, path, rows=None):
    # 현재 정리본에서는 장부와 동일 데이터 구조로 만들되 파일명을 분리.
    return build_ledger(orders, ship_date, path, rows)

ERP_FIELDS = ["일자","전표번호","구분","거래처","품명","규격","수량","단가","공급가액","부가세","적요","기재사항1","기재사항2"]

def build_erp(orders, ship_date, path, rows=None):
    """경영박사 import용 개발 기준 CSV.
    실제 공장 EDI 열 순서/코드는 legacy output.py(v79)와 대조하여 보완한다.
    """
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    rows = build_rows(orders, ship_date) if rows is None else rows
    item_rows = [r for r in rows if r.get("_특수") is None]
    heads = {r["_oi"]: r.get("상호") for r in item_rows if r.get("_ii") == 0}
    with p.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=ERP_FIELDS)
        w.writeheader()
        for r in item_rows:
            oi = r["_oi"]
            w.writerow({
                "일자": str(ship_date),
                "전표번호": oi + 1,
                "구분": "외출",
                "거래처": heads.get(oi),
                "품명": r.get("_품명"),
                "규격": f"{r.get('가로','')}x{r.get('세로','')}",
                "수량": r.get("수량"),
                "단가": "",
                "공급가액": "",
                "부가세": "",
                "적요": orders[oi].get("주문번호") or "",
                "기재사항1": r.get("기재사항1",""),
                "기재사항2": r.get("기재사항2",""),
            })
    return p

def read_ledger(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in ws[1]]
    rows = []
    for vals in ws.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(headers, vals)))
    return rows

def build_all(orders, ship_date=None, out_dir="out"):
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    stamp = date.today().strftime("%Y%m%d")
    ship_date = ship_date or date.today()
    rows = build_rows(orders, ship_date)  # 세 출력이 같은 계산 결과를 공유
    ledger = build_ledger(orders, ship_date, out / f"장부_{stamp}.xlsx", rows)
    work = build_worksheet(orders, ship_date, out / f"작업지시서_{stamp}.xlsx", rows)
    erp = build_erp(orders, ship_date, out / f"경영박사_EDI_{stamp}.csv", rows)
    return {"ledger":str(ledger),"worksheet":str(work),"erp":str(erp)}
